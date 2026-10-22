# -*- coding: utf-8 -*-
"""
Из файла/архива услуг → Excel для fix_service (уникальные талоны с «было»/«стало»).

Примеры:
  python prepare_from_services.py --zip ..\\dn_fix_3_....zip --from B04.015.003 --to B04.047.003
  python prepare_from_services.py --services услуги.xlsx --from B04.015.003 --to B04.047.003
  python prepare_from_services.py --services услуги.xlsx --list-codes
"""
from __future__ import annotations

import argparse
import re
import shutil
import sys
import zipfile
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook

WORK_DIR = Path(__file__).resolve().parent
DEFAULT_OUT = WORK_DIR / "data" / "fix_service.xlsx"

_SERVICE_CODE_RE = re.compile(
    r"[A-Za-z]\d{2}\.\d{2,3}\.\d{3}(?:\.\d+)*",
    re.IGNORECASE,
)


def _zip_name(info: zipfile.ZipInfo) -> str:
    name = info.filename
    try:
        return name.encode("cp437").decode("cp866")
    except Exception:
        return name


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v).strip()


def normalize_code(value: object) -> str:
    text = _cell(value).upper().replace(" ", "")
    if not text:
        return ""
    m = _SERVICE_CODE_RE.search(text)
    return m.group(0).upper() if m else text


def extract_services_xlsx(zip_path: Path, dest_dir: Path) -> Path:
    dest_dir.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(zip_path) as zf:
        for info in zf.infolist():
            name = _zip_name(info)
            low = name.lower()
            if "услуг" not in low and "service" not in low:
                continue
            if not name.lower().endswith((".xlsx", ".xls")):
                continue
            out = dest_dir / Path(name).name
            with zf.open(info) as src, open(out, "wb") as dst:
                shutil.copyfileobj(src, dst)
            return out
    raise SystemExit(f"В архиве нет xlsx услуг: {zip_path}")


def read_service_rows(path: Path) -> tuple[list[str], list[tuple[str, str]]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    headers = [_cell(h).lower() for h in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    col = {h: i for i, h in enumerate(headers) if h}

    def idx(*names: str) -> int | None:
        for n in names:
            if n in col:
                return col[n]
        return None

    i_talon = idx("талон", "номер", "id", "claim", "web-oms", "web_oms")
    i_code = idx("код услуги", "код", "услуга", "service", "code", "service_code")
    if i_talon is None or i_code is None:
        wb.close()
        raise SystemExit(f"Нужны колонки «Талон» и «Код услуги», сейчас: {headers}")

    rows: list[tuple[str, str]] = []
    for raw in ws.iter_rows(min_row=2, values_only=True):
        if not raw:
            continue
        talon = _cell(raw[i_talon])
        if talon.endswith(".0"):
            talon = talon[:-2]
        code = normalize_code(raw[i_code])
        if talon and code:
            rows.append((talon, code))
    wb.close()
    return headers, rows


def write_fix_xlsx(path: Path, pairs: list[tuple[str, str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "замена"
    ws.append(["Талон", "было", "стало"])
    for talon, code_from, code_to in pairs:
        ws.append([talon, code_from, code_to])
    wb.save(path)


def main() -> int:
    p = argparse.ArgumentParser(description="Услуги → уникальные талоны для fix_service")
    p.add_argument("--zip", type=Path, help="Архив с услугами/талонами")
    p.add_argument("--services", type=Path, help="Excel услуг")
    p.add_argument("--from", dest="code_from", default="", help="Код услуги «было»")
    p.add_argument("--to", dest="code_to", default="", help="Код услуги «стало»")
    p.add_argument("--out", type=Path, default=DEFAULT_OUT, help="Куда писать fix_service.xlsx")
    p.add_argument("--list-codes", action="store_true", help="Только топ кодов, без записи")
    p.add_argument("--top", type=int, default=30)
    args = p.parse_args()

    services = args.services
    if args.zip:
        extract_dir = WORK_DIR / "data" / "_extract"
        services = extract_services_xlsx(args.zip.resolve(), extract_dir)
        print(f"Услуги из zip: {services}")
    if not services or not services.is_file():
        print("Укажите --zip или --services", file=sys.stderr)
        return 2

    _headers, rows = read_service_rows(services)
    codes = Counter(c for _, c in rows)
    talons_all = {t for t, _ in rows}
    print(f"Строк услуг: {len(rows)} | уникальных талонов: {len(talons_all)} | кодов: {len(codes)}")

    if args.list_codes or not args.code_from:
        print("Топ кодов:")
        for code, n in codes.most_common(args.top):
            talons_n = len({t for t, c in rows if c == code})
            print(f"  {n:5d} строк | {talons_n:4d} талонов | {code}")
        if not args.code_from:
            print("\nДальше: --from КОД --to КОД")
            return 0

    code_from = normalize_code(args.code_from)
    code_to = normalize_code(args.code_to)
    if not code_to:
        print("Нужен --to", file=sys.stderr)
        return 2

    talons = sorted({t for t, c in rows if c == code_from})
    print(f"Талонов с «{code_from}»: {len(talons)} → ставим «{code_to}»")
    if not talons:
        print("Нечего писать.")
        return 1

    pairs = [(t, code_from, code_to) for t in talons]
    write_fix_xlsx(args.out, pairs)
    print(f"Wrote {args.out} ({len(pairs)} строк)")
    print(
        r'Запуск: .\.venv\Scripts\python run_fix_service.py '
        rf'--file {args.out} --limit 1 --dry-run --keep-open 15'
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
