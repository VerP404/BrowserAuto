#!/usr/bin/env python
"""
Правка услуг на амбулаторных талонах по выгрузке портала
(zip: талоны*.xlsx + услуги*.xlsx).

Логика (как приведение к целевому набору):
  1) прочитать коды на форме;
  2) сравнить с файлом (мультимножество);
  3) лишние строки → заменить на недостающие коды;
  4) если недостающих нет → удалить лишние (кнопка «Удалить сведение»);
  5) если лишних нет, а коды ещё нужны → #add-service и заполнить новую строку
     (как cel_3: новая строка сверху → combobox-service-0);
  6) Save только если были правки.

Пример:
  python run_sync_services.py --zip robots/oms/dn_fix_3_….zip --limit 5
  .\\robots\\oms\\fix_service\\run_sync_parallel.ps1 -Workers 4 -Zip ...
"""

from __future__ import annotations

import argparse
import csv
import re
import shutil
import sys
import zipfile
from collections import Counter
from dataclasses import asdict, dataclass, field
from datetime import datetime
from pathlib import Path
from time import sleep, time

from loguru import logger
from openpyxl import load_workbook
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.remote.webdriver import WebDriver

WORK_DIR = Path(__file__).resolve().parent
OMS_DIR = WORK_DIR.parent
PROJECT_ROOT = WORK_DIR.parents[2]
for p in (PROJECT_ROOT, WORK_DIR, OMS_DIR / "dv_opv", OMS_DIR / "cel_307", OMS_DIR / "cel_3"):
    if str(p) not in sys.path:
        sys.path.insert(0, str(p))

from browser_auto.config import load_credentials
from browser_auto.auth import login_oms
from browser_auto.driver import GeckoBrowser

import add_307 as base  # noqa: E402
import add_medical_exam as dv  # noqa: E402
import fix_service as fs  # noqa: E402

load_credentials(WORK_DIR / "credentials.env", PROJECT_ROOT / ".env")

RESULTS_DIR = WORK_DIR / "data"
_SERVICE_CODE_RE = re.compile(
    r"[A-Za-z]\d{2}\.\d{2,3}\.\d{3}(?:\.\d+)*",
    re.IGNORECASE,
)


@dataclass
class ServiceWanted:
    code: str
    begin: str = ""
    end: str = ""
    amount: str = "1"
    doctor: str = ""


@dataclass
class TalonBundle:
    talon: str
    doctor: str = ""
    begin: str = ""
    end: str = ""
    services: list[ServiceWanted] = field(default_factory=list)


@dataclass
class RowResult:
    index: int
    talon: str
    services_wanted: int
    services_touched: int
    status: str
    message: str
    elapsed_sec: float


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v).strip()


def normalize_code(value: object) -> str:
    text = _cell(value).upper().replace(" ", "")
    if not text:
        return ""
    m = _SERVICE_CODE_RE.search(text)
    return m.group(0).upper() if m else text


def _zip_member_name(info: zipfile.ZipInfo) -> str:
    try:
        return info.filename.encode("cp437").decode("cp866")
    except Exception:
        return info.filename


def _find_bundle_xlsx(dest: Path) -> tuple[Path | None, Path | None]:
    tal = svc = None
    for out in dest.glob("*.xlsx"):
        low = out.name.lower()
        if "талон" in low:
            tal = out
        if "услуг" in low:
            svc = out
    return tal, svc


def extract_bundle(zip_path: Path, dest: Path) -> tuple[Path, Path]:
    """Распаковать zip. У каждого воркера свой dest — иначе WinError 32 на rmtree."""
    if dest.exists():
        try:
            shutil.rmtree(dest)
        except PermissionError:
            tal, svc = _find_bundle_xlsx(dest)
            if tal and svc:
                logger.warning("extract занят ({}), читаем уже распакованные xlsx", dest)
                return tal, svc
            raise
    dest.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(zip_path) as zf:
        for info in zf.infolist():
            name = _zip_member_name(info)
            low = name.lower()
            if not low.endswith(".xlsx"):
                continue
            out = dest / Path(name).name
            with zf.open(info) as src, open(out, "wb") as dst:
                shutil.copyfileobj(src, dst)
    tal, svc = _find_bundle_xlsx(dest)
    if not tal or not svc:
        raise SystemExit(f"В архиве нужны талоны*.xlsx и услуги*.xlsx, есть: {list(dest.glob('*'))}")
    return tal, svc


def _xlsx_rows(path: Path) -> tuple[list[str], list[tuple]]:
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    headers = [_cell(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = []
    for raw in ws.iter_rows(min_row=2, values_only=True):
        if not raw or all(v is None or str(v).strip() == "" for v in raw):
            continue
        rows.append(tuple(raw))
    wb.close()
    return headers, rows


def _col(header: list[str], *names: str) -> int | None:
    idx = {h.lower().strip(): i for i, h in enumerate(header) if h}
    for n in names:
        if n.lower() in idx:
            return idx[n.lower()]
    return None


def load_bundle(
    *,
    talons_path: Path,
    services_path: Path,
) -> list[TalonBundle]:
    th, trows = _xlsx_rows(talons_path)
    sh, srows = _xlsx_rows(services_path)

    i_id = _col(th, "талон", "талон_омс")
    i_doc = _col(th, "врач")
    i_beg = _col(th, "начало")
    i_end = _col(th, "окончание")
    if i_id is None:
        raise SystemExit(f"В талонах нужна колонка «Талон», есть: {th}")

    si_id = _col(sh, "талон")
    si_code = _col(sh, "код услуги", "услуга")
    si_beg = _col(sh, "дата начала", "начало")
    si_end = _col(sh, "дата окончания", "окончание")
    si_amt = _col(sh, "кол-во", "количество")
    si_doc = _col(sh, "врач")
    if si_id is None or si_code is None:
        raise SystemExit(f"В услугах нужны «Талон» и «Код услуги», есть: {sh}")

    by_svc: dict[str, list[ServiceWanted]] = {}
    for raw in srows:
        tid = _cell(raw[si_id])
        if tid.endswith(".0"):
            tid = tid[:-2]
        code = normalize_code(raw[si_code])
        if not tid or not code:
            continue
        by_svc.setdefault(tid, []).append(
            ServiceWanted(
                code=code,
                begin=_cell(raw[si_beg]) if si_beg is not None else "",
                end=_cell(raw[si_end]) if si_end is not None else "",
                amount=(_cell(raw[si_amt]) if si_amt is not None else "") or "1",
                doctor=_cell(raw[si_doc]) if si_doc is not None else "",
            )
        )

    meta: dict[str, tuple[str, str, str]] = {}
    for raw in trows:
        tid = _cell(raw[i_id])
        if tid.endswith(".0"):
            tid = tid[:-2]
        if not tid:
            continue
        meta[tid] = (
            _cell(raw[i_doc]) if i_doc is not None else "",
            _cell(raw[i_beg]) if i_beg is not None else "",
            _cell(raw[i_end]) if i_end is not None else "",
        )

    out: list[TalonBundle] = []
    for tid, services in by_svc.items():
        doc, beg, end = meta.get(tid, ("", "", ""))
        for s in services:
            if not s.doctor:
                s.doctor = doc
            if not s.begin:
                s.begin = beg
            if not s.end:
                s.end = end
        out.append(TalonBundle(talon=tid, doctor=doc, begin=beg, end=end, services=services))

    out.sort(key=lambda t: t.talon)
    return out


def _force_id(driver: WebDriver, element_id: str, value: str) -> bool:
    if not value or not driver.find_elements(By.ID, element_id):
        return False
    try:
        return bool(base._force_input(driver, element_id, value))
    except Exception:
        try:
            dv.fill_date_field(driver, element_id, value)
            return True
        except Exception as exc:
            logger.debug("{}: {}", element_id, exc)
            return False


def fill_service_row(driver: WebDriver, index: int, svc: ServiceWanted, *, talon_doctor: str) -> None:
    code = normalize_code(svc.code)
    fs.set_service_code(driver, index, code)

    begin = dv.normalize_service_date(svc.begin)
    end = dv.normalize_service_date(svc.end)
    if begin:
        _force_id(driver, f"begin-date-service-{index}", begin)
    if end:
        _force_id(driver, f"end-date-service-{index}", end)

    amount = (svc.amount or "1").strip() or "1"
    if driver.find_elements(By.ID, f"amount-service-{index}"):
        if not _force_id(driver, f"amount-service-{index}", amount):
            el = driver.find_element(By.ID, f"amount-service-{index}")
            try:
                el.clear()
                el.send_keys(amount, Keys.ENTER)
            except Exception:
                pass

    doctor = (svc.doctor or talon_doctor or "").strip()
    if doctor:
        for sid in (f"doctor-service-{index}", f"combobox-doctor-{index}"):
            els = driver.find_elements(By.ID, sid)
            if not els:
                continue
            target = els[-1]
            try:
                if target.is_displayed():
                    base._fill_select_fast(driver, target, doctor)
                    break
            except Exception as exc:
                logger.debug("doctor {}: {}", sid, exc)

    logger.info(
        "sync service[{}] {} amt={} {}…{} doctor={}",
        index,
        code,
        amount,
        begin or "—",
        end or "—",
        doctor or "—",
    )


def _wanted_by_code(wanted: list[ServiceWanted]) -> dict[str, list[ServiceWanted]]:
    by: dict[str, list[ServiceWanted]] = {}
    for s in wanted:
        code = normalize_code(s.code)
        if code:
            by.setdefault(code, []).append(s)
    return by


@dataclass
class DiffState:
    """Снимок расхождения формы и файла."""

    cur_codes: list[str]
    want_codes: list[str]
    keep_indices: set[int]
    free_indices: list[int]  # лишние + пустые — кандидаты на замену/удаление
    need_queue: list[ServiceWanted]
    surplus_codes: list[str]


def compute_diff(
    current: list[tuple[int, str]],
    wanted: list[ServiceWanted],
) -> DiffState:
    want_codes = [normalize_code(s.code) for s in wanted if normalize_code(s.code)]
    want_cnt = Counter(want_codes)
    cur_codes = [c for _, c in current if c]
    cur_cnt = Counter(cur_codes)
    surplus_codes = list((cur_cnt - want_cnt).elements())

    by_want = _wanted_by_code(wanted)
    kept_cnt: Counter[str] = Counter()
    keep_indices: set[int] = set()
    for idx, code in current:
        if not code:
            continue
        if kept_cnt[code] < want_cnt[code]:
            keep_indices.add(idx)
            kept_cnt[code] += 1

    free_indices = [idx for idx, _ in current if idx not in keep_indices]

    need_queue: list[ServiceWanted] = []
    remain = want_cnt - kept_cnt
    for code, n in remain.items():
        pool = by_want.get(code) or []
        start = kept_cnt[code]
        for k in range(n):
            pi = start + k
            need_queue.append(pool[pi] if pi < len(pool) else ServiceWanted(code=code))

    return DiffState(
        cur_codes=cur_codes,
        want_codes=want_codes,
        keep_indices=keep_indices,
        free_indices=free_indices,
        need_queue=need_queue,
        surplus_codes=surplus_codes,
    )


def _read_current_services(driver: WebDriver) -> list[tuple[int, str]]:
    indices = fs.list_service_indices(driver)
    return [(i, fs.service_code_at(driver, i) or "") for i in indices]


def _add_and_fill_top(
    driver: WebDriver,
    svc: ServiceWanted,
    *,
    talon_doctor: str,
) -> int:
    """Как cel_3/service_notebook: Add → заполнить combobox-service-0."""
    before = set(fs.list_service_indices(driver))
    fs.click_add_service(driver)
    after = fs.list_service_indices(driver)
    # новая строка обычно сверху (index 0); иначе — первый новый индекс
    new_idxs = [i for i in after if i not in before]
    target = 0 if 0 in after else (new_idxs[0] if new_idxs else (after[0] if after else None))
    if target is None:
        raise RuntimeError("После add-service нет строк услуг")
    fill_service_row(driver, target, svc, talon_doctor=talon_doctor)
    return target


def sync_services(driver: WebDriver, row: TalonBundle) -> tuple[str, int, str]:
    """Привести услуги формы к файлу. Returns (action, touched, detail)."""
    wanted = [s for s in row.services if normalize_code(s.code)]
    if not wanted:
        raise RuntimeError(f"В файле нет услуг для талона {row.talon}")

    fs.open_services_section(driver)
    fs.wait_service_rows(driver, timeout=20)

    changes: list[str] = []
    touched = 0
    guard = 0
    max_steps = max(40, len(wanted) * 3 + 10)

    while guard < max_steps:
        guard += 1
        current = _read_current_services(driver)
        diff = compute_diff(current, wanted)

        if not diff.free_indices and not diff.need_queue:
            detail = (
                f"уже ок ({len(diff.cur_codes)}): {', '.join(diff.cur_codes) or '—'}"
                if not changes
                else (
                    f"готово | форма[{len(diff.cur_codes)}]: {', '.join(diff.cur_codes) or '—'} | "
                    f"файл[{len(diff.want_codes)}]: {', '.join(diff.want_codes)}"
                )
            )
            if not changes:
                logger.info("талон {}: {}", row.talon, detail)
                return "skip", 0, detail
            msg = f"правок {touched}: {', '.join(changes)} | {detail}"
            logger.info("талон {}: {}", row.talon, msg)
            return "patch", touched, msg

        snapshot = (
            f"форма[{len(diff.cur_codes)}]: {', '.join(diff.cur_codes) or '—'} | "
            f"файл[{len(diff.want_codes)}]: {', '.join(diff.want_codes)} | "
            f"лишние: {', '.join(diff.surplus_codes) or '—'} | "
            f"нет: {', '.join(normalize_code(s.code) for s in diff.need_queue) or '—'}"
        )
        if guard == 1:
            logger.info("талон {}: {}", row.talon, snapshot)

        # 1) есть лишняя/пустая строка и есть недостающий код → заменить
        if diff.free_indices and diff.need_queue:
            idx = diff.free_indices[0]
            svc = diff.need_queue[0]
            old = fs.service_code_at(driver, idx) or "—"
            fill_service_row(driver, idx, svc, talon_doctor=row.doctor)
            new = normalize_code(svc.code)
            changes.append(f"{old}→{new}")
            touched += 1
            continue

        # 2) лишние есть, недостающих нет → удалить (только по индексу — дубли опасны)
        if diff.free_indices and not diff.need_queue:
            idx = max(diff.free_indices)  # с хвоста — меньше сдвиг индексов
            old = fs.service_code_at(driver, idx) or "—"
            if not fs.delete_service_row(driver, idx):
                raise RuntimeError(f"Не удалось удалить лишнюю услугу [{idx}] {old}")
            changes.append(f"del:{old}")
            touched += 1
            continue

        # 3) нужны коды, свободных строк нет → добавить (cel_3: add → service-0)
        if diff.need_queue and not diff.free_indices:
            svc = diff.need_queue[0]
            target = _add_and_fill_top(driver, svc, talon_doctor=row.doctor)
            new = normalize_code(svc.code)
            changes.append(f"add[{target}]:{new}")
            touched += 1
            continue

        raise RuntimeError(f"Неожиданное состояние sync: {snapshot}")

    raise RuntimeError(
        f"Слишком много шагов sync ({guard}) для талона {row.talon}: {', '.join(changes[-8:])}"
    )


def process_row(
    driver: WebDriver,
    row: TalonBundle,
    *,
    index: int,
    total: int,
    dry_run: bool,
) -> RowResult:
    t0 = time()
    pos = f"{index}/{total}"
    logger.info("——— [{}] talon={} файл={} услуг ———", pos, row.talon, len(row.services))
    try:
        fs.open_claim(driver, row.talon)
        action, n, detail = sync_services(driver, row)

        if action == "skip":
            logger.success("[{}] SKIP — {}", pos, detail)
            return RowResult(
                index, row.talon, len(row.services), 0, "skip", detail, round(time() - t0, 2)
            )

        if dry_run:
            msg = f"dry-run: {detail}, без Save"
            logger.success("[{}] DRY — {}", pos, msg)
            return RowResult(
                index, row.talon, len(row.services), n, "dry_run", msg, round(time() - t0, 2)
            )

        ok, msg = fs.save_existing_claim(driver, row.talon)
        status = "ok" if ok else "fail"
        full = f"{detail} | {msg}"
        log = logger.success if ok else logger.error
        log("[{}] {} (правок={}) — {}", pos, status.upper(), n, full[:350])
        return RowResult(index, row.talon, len(row.services), n, status, full, round(time() - t0, 2))
    except Exception as exc:
        msg = str(exc).split("Stacktrace:")[0].strip() or str(exc)
        logger.error("[{}] ERROR — {}", pos, msg[:500])
        return RowResult(index, row.talon, len(row.services), 0, "error", msg, round(time() - t0, 2))


def write_results(path: Path, results: list[RowResult]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = list(asdict(results[0]).keys()) if results else [
        "index", "talon", "services_wanted", "services_touched", "status", "message", "elapsed_sec"
    ]
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        for r in results:
            w.writerow(asdict(r))


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Синхронизация всех услуг талона по zip портала")
    p.add_argument("--zip", type=Path, default=None, help="Архив талоны+услуги из портала")
    p.add_argument("--talons", type=Path, default=None)
    p.add_argument("--services", type=Path, default=None)
    p.add_argument("--limit", type=int, default=0)
    p.add_argument("--offset", type=int, default=0)
    p.add_argument("--dry-run", action="store_true")
    p.add_argument("--keep-open", type=int, default=0)
    p.add_argument("--worker", default="")
    p.add_argument("--window-x", type=int, default=None)
    p.add_argument("--window-y", type=int, default=None)
    return p.parse_args()


def main() -> int:
    args = parse_args()
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)

    if args.zip:
        zpath = args.zip if args.zip.is_file() else PROJECT_ROOT / args.zip
        if not zpath.is_file():
            zpath = OMS_DIR / args.zip.name
        if not zpath.is_file():
            logger.error("Нет zip {}", args.zip)
            return 1
        worker_tag = (args.worker or "").strip() or "main"
        extract_dir = RESULTS_DIR / f"_sync_extract_{worker_tag}"
        talons_path, services_path = extract_bundle(zpath, extract_dir)
    else:
        if not args.talons or not args.services:
            logger.error("Нужен --zip или пара --talons/--services")
            return 1
        talons_path, services_path = args.talons, args.services

    rows = load_bundle(talons_path=talons_path, services_path=services_path)
    logger.info(
        "Загружено талонов с услугами: {} (из {})",
        len(rows),
        services_path.name,
    )
    if args.offset:
        rows = rows[args.offset :]
    if args.limit and args.limit > 0:
        rows = rows[: args.limit]

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    worker_tag = (args.worker or "").strip() or f"o{args.offset}"
    result_path = RESULTS_DIR / f"sync_services_result_{stamp}_{worker_tag}.csv"
    logger.info(
        "[{}] к обработке={} | dry_run={} | offset={}",
        worker_tag,
        len(rows),
        args.dry_run,
        args.offset,
    )
    if not rows:
        return 0

    win_kwargs: dict = {"implicit_wait": 0, "maximize": False}
    if args.window_x is not None and args.window_y is not None:
        win_kwargs["window_position"] = (args.window_x, args.window_y)
        win_kwargs["window_size"] = (1100, 850)
    else:
        win_kwargs["maximize"] = True

    results: list[RowResult] = []
    t_all = time()
    batch_total = len(rows)
    with GeckoBrowser(**win_kwargs) as session:
        driver = session.driver
        driver.implicitly_wait(0)
        login_oms(driver)
        sleep(0.8)
        dv.wait_busy_gone(driver, timeout=40)

        for n_done, row in enumerate(rows, start=1):
            results.append(
                process_row(
                    driver,
                    row,
                    index=n_done,
                    total=batch_total,
                    dry_run=args.dry_run,
                )
            )
            write_results(result_path, results)
            if n_done % 10 == 0 or n_done == batch_total:
                skip_n = sum(1 for r in results if r.status == "skip")
                ok_n = sum(1 for r in results if r.status in ("ok", "dry_run"))
                fail_n = sum(1 for r in results if r.status in ("fail", "error"))
                logger.info(
                    "[{}] прогресс {}/{} | skip={} patch_ok={} fail={}",
                    worker_tag,
                    n_done,
                    batch_total,
                    skip_n,
                    ok_n,
                    fail_n,
                )

        if args.keep_open > 0:
            logger.info("keep-open {}s", args.keep_open)
            sleep(args.keep_open)

    skip_n = sum(1 for r in results if r.status == "skip")
    ok_n = sum(1 for r in results if r.status in ("ok", "dry_run"))
    fail_n = sum(1 for r in results if r.status in ("fail", "error"))
    logger.info(
        "——— [{}] ИТОГ ({:.0f}s) skip={} ok={} fail={} → {} ———",
        worker_tag,
        time() - t_all,
        skip_n,
        ok_n,
        fail_n,
        result_path,
    )
    return 0 if fail_n == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
