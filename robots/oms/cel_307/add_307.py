#!/usr/bin/env python
"""
Ввод талонов цели 307 в Web.ОМС (амбулаторный талон).

URL: /claim/ambulatory
Логика полей — как в талоны2_0.ipynb; соцстатус/занятость — как в oms_dv_opv
(из Excel или по ДР в карте пациента).

Одна услуга на талон — всё в одном Excel.

Пример:
  python add_307.py --make-template
  python add_307.py --limit 3 --building "ГП №11"
  python add_307.py --file data/talons_307.xlsx --dry-run --keep-open 20
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
from dataclasses import asdict, dataclass
from datetime import date, datetime
from pathlib import Path
from time import sleep, time

from loguru import logger
from openpyxl import Workbook, load_workbook
from selenium.common.exceptions import NoSuchElementException, TimeoutException
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.remote.webdriver import WebDriver
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait

WORK_DIR = Path(__file__).resolve().parent
OMS_DIR = WORK_DIR.parent  # robots/oms
PROJECT_ROOT = WORK_DIR.parents[2]  # BrowserAuto
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))
if str(WORK_DIR) not in sys.path:
    sys.path.insert(0, str(WORK_DIR))
DV_DIR = OMS_DIR / "dv_opv"
if str(DV_DIR) not in sys.path:
    sys.path.insert(0, str(DV_DIR))

from browser_auto.config import OMS_BASE_URL, load_credentials
from browser_auto.auth import login_oms
from browser_auto.driver import GeckoBrowser

import add_medical_exam as dv  # noqa: E402 — общие хелперы ОМС из oms_dv_opv

load_credentials(WORK_DIR / "credentials.env")

DEFAULT_FILE = WORK_DIR / "data" / "talons_307.xlsx"
RESULTS_DIR = WORK_DIR / "data"
AMBULATORY_PATH = "/claim/ambulatory"

DEFAULTS_307 = {
    "цель": "307",
    "случай": "Первичный",
    "случай_2": "Законченный",
    "ДИСП. НАБЛ": "1",
    "тип": "3",
    "результат": "301",
    "исход": "304",
    "характер": "3",
}

# Направление (как в талоны2_0)
DEFAULT_REF_NUMBER = "бн"
DEFAULT_REF_ORG = "360025"
DEFAULT_REF_COND = "Амбулаторно"
DEFAULT_PLACE = "1"
DEFAULT_HOME = "0"

TEMPLATE_COLUMNS = [
    "ЕНП",
    "начало",
    "окончание",
    "врач",
    "Корпус",
    "цель",
    "случай",
    "случай_2",
    "посещений в МО",
    "Диагноз",
    "Диагноз 2",
    "ДИСП. НАБЛ",
    "тип",
    "результат",
    "исход",
    "характер",
    "Код услуги",
    "Кол-во",
    "Соцстатус",
    "Вид занятости",
]


@dataclass
class Talon307:
    enp: str
    begin: str
    end: str
    doctor: str
    building: str
    goal: str
    occurrence: str
    occurrence2: str
    mo_visits: str
    diagnosis: str
    diagnosis2: str
    dn: str
    claim_type: str
    result: str
    outcome: str
    character: str
    service_code: str
    service_amount: str = "1"  # кол-во услуги (всегда 1 строка/1 ед.)
    social_status: str = ""
    occupation: str = ""
    place: str = DEFAULT_PLACE
    home: str = DEFAULT_HOME
    ref_number: str = DEFAULT_REF_NUMBER
    ref_org: str = DEFAULT_REF_ORG
    ref_cond: str = DEFAULT_REF_COND
    raw_index: int = 0


@dataclass
class RowResult:
    index: int
    enp: str
    service_code: str
    status: str
    message: str
    elapsed_sec: float = 0.0


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d-%m-%y")
    if isinstance(value, date):
        return value.strftime("%d-%m-%y")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    m = re.match(r"^(\d{1,2})[./](\d{1,2})[./](\d{2,4})$", text)
    if m:
        d, mo, y = m.groups()
        if len(y) == 4:
            y = y[-2:]
        return f"{int(d):02d}-{int(mo):02d}-{y}"
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", text)
    if m:
        y, mo, d = m.groups()
        return f"{int(d):02d}-{int(mo):02d}-{y[-2:]}"
    return text


def to_oms_date(value: object) -> str:
    return _cell(value) if not isinstance(value, str) or re.search(r"\d", value or "") else _cell(value)


def _wait(driver: WebDriver, timeout: float = 12) -> WebDriverWait:
    return WebDriverWait(driver, timeout, poll_frequency=0.15)


def ambulatory_url() -> str:
    base = (OMS_BASE_URL or "http://10.36.0.142:9000").rstrip("/")
    return f"{base}{AMBULATORY_PATH}"


def open_ambulatory(driver: WebDriver) -> None:
    url = ambulatory_url()
    logger.info("Открываем {}", url)
    driver.get(url)
    dv.wait_busy_gone(driver, timeout=20)
    dv.dismiss_overlays(driver)
    _wait(driver, 12).until(EC.presence_of_element_located((By.ID, "enp")))
    sleep(0.15)


def _click_xpath(driver: WebDriver, xpath: str) -> bool:
    els = driver.find_elements(By.XPATH, xpath)
    if not els:
        return False
    try:
        els[0].click()
    except Exception:
        driver.execute_script("arguments[0].click();", els[0])
    return True


def fill_referral(driver: WebDriver, row: Talon307) -> None:
    """talon_napr из ноутбука."""
    ref_date = dv.normalize_service_date(row.begin)
    dv.try_input_enter_id(driver, "reference-number", row.ref_number or DEFAULT_REF_NUMBER)
    if driver.find_elements(By.ID, "reference-date"):
        if not _force_input(driver, "reference-date", ref_date):
            try:
                dv.fill_date_field(driver, "reference-date", row.begin)
            except Exception:
                dv.try_input_enter_id(driver, "reference-date", ref_date)
    dv.try_input_enter_id(driver, "reference-org", row.ref_org or DEFAULT_REF_ORG)
    dv.try_input_enter_id(
        driver,
        "referenceMedicalCareConditionsId",
        row.ref_cond or DEFAULT_REF_COND,
    )


def _enter_select(driver: WebDriver, element_id: str, text: str) -> None:
    """react-select / обычный input с Enter."""
    if not text:
        return
    els = driver.find_elements(By.ID, element_id)
    if not els:
        logger.warning("Нет поля #{}", element_id)
        return
    el = els[0]
    cls = (el.get_attribute("class") or "") + " " + (el.get_attribute("role") or "")
    parent_cls = ""
    try:
        parent_cls = el.find_element(By.XPATH, "../..").get_attribute("class") or ""
    except Exception:
        pass
    if "Select" in cls or "Select" in parent_cls or "react-select" in cls.lower():
        _fill_select_fast(driver, el, text)
    else:
        try:
            el.clear()
        except Exception:
            pass
        el.send_keys(text, Keys.ENTER)
    sleep(0.06)
    dv.wait_busy_gone(driver, timeout=6)


def _fill_select_fast(driver: WebDriver, input_el, text: str) -> None:
    """Быстрый ввод в react-select (короче пауз, чем oms_dv_opv)."""
    driver.execute_script("arguments[0].scrollIntoView({block:'center'});", input_el)
    try:
        input_el.click()
    except Exception:
        driver.execute_script("arguments[0].focus();", input_el)
    sleep(0.05)
    try:
        input_el.send_keys(Keys.CONTROL, "a")
        input_el.send_keys(Keys.BACKSPACE)
    except Exception:
        pass
    input_el.send_keys(str(text))
    sleep(0.25)
    options = driver.find_elements(
        By.CSS_SELECTOR,
        ".Select-option, .Select-menu .Select-option, .Select-menu-outer .Select-option, div[role='option']",
    )
    needle = str(text).strip().lower()
    for opt in options:
        label = (opt.text or "").strip().lower()
        if not label:
            continue
        if needle == label or label.startswith(needle) or needle in label:
            try:
                driver.execute_script("arguments[0].click();", opt)
                sleep(0.08)
                return
            except Exception:
                continue
    if options:
        try:
            driver.execute_script("arguments[0].click();", options[0])
            sleep(0.08)
            return
        except Exception:
            pass
    input_el.send_keys(Keys.ENTER)
    sleep(0.1)


def _force_input(driver: WebDriver, element_id: str, value: str) -> bool:
    """Поставить value даже если input не reachable by keyboard (даты услуг)."""
    els = driver.find_elements(By.ID, element_id)
    if not els or not value:
        return False
    el = els[0]
    driver.execute_script(
        """
        const el = arguments[0], val = arguments[1];
        el.removeAttribute('readonly');
        el.removeAttribute('disabled');
        el.disabled = false;
        el.readOnly = false;
        const setter = Object.getOwnPropertyDescriptor(
            window.HTMLInputElement.prototype, 'value'
        ).set;
        el.focus();
        setter.call(el, '');
        el.dispatchEvent(new Event('input', {bubbles:true}));
        setter.call(el, val);
        el.dispatchEvent(new Event('input', {bubbles:true}));
        el.dispatchEvent(new Event('change', {bubbles:true}));
        el.dispatchEvent(new Event('blur', {bubbles:true}));
        """,
        el,
        value,
    )
    actual = (el.get_attribute("value") or "").strip()
    ok = bool(actual)
    logger.info("{} = {} (force{})", element_id, actual or "пусто", "" if ok else " FAIL")
    return ok


def fill_main(driver: WebDriver, row: Talon307) -> None:
    """talon_main из ноутбука."""
    if not (row.building or "").strip():
        raise RuntimeError("Пустой Корпус — укажите в Excel или --building")

    _enter_select(driver, "subdivision", row.building)
    _enter_select(driver, "doctor", row.doctor)
    _enter_select(driver, "requestPurpose", row.goal or "307")
    sleep(0.35)  # цель 307 подтягивает шаблон услуг
    _enter_select(driver, "service-place", row.place or DEFAULT_PLACE)
    _enter_select(driver, "occurrence", row.occurrence or DEFAULTS_307["случай"])
    _enter_select(driver, "second-occurrence", row.occurrence2 or DEFAULTS_307["случай_2"])
    _enter_select(driver, "mo", row.mo_visits)
    _enter_select(driver, "home", row.home or DEFAULT_HOME)
    _enter_select(driver, "mainDiagnosis", dv.normalize_mkb(row.diagnosis))
    _fill_accompanying(driver, row.diagnosis2)
    _enter_select(driver, "dispensaryObservation", row.dn or DEFAULTS_307["ДИСП. НАБЛ"])
    _enter_select(driver, "type", row.claim_type or DEFAULTS_307["тип"])
    _enter_select(driver, "request-result", row.result or DEFAULTS_307["результат"])
    _enter_select(driver, "outcome", row.outcome or DEFAULTS_307["исход"])
    _enter_select(driver, "characterMainDisease", row.character or DEFAULTS_307["характер"])
    _enter_select(driver, "doctor", row.doctor)
    if row.social_status:
        _enter_select(driver, "socialStatus", row.social_status)
    if row.occupation:
        _enter_select(driver, "occupation", row.occupation)



def _fill_accompanying(driver: WebDriver, raw: str) -> None:
    codes = []
    text = (raw or "").strip()
    if text:
        for sep in (",", ";", "|"):
            if sep in text:
                codes = [p.strip() for p in text.split(sep) if p.strip()]
                break
        else:
            codes = [text]
    if not codes:
        return
    xpath = '//*[@id="accompanying-diagnosis"]'
    for i, code in enumerate(codes):
        if i > 0:
            els = driver.find_elements(By.XPATH, xpath)
            if els:
                try:
                    els[0].click()
                except Exception:
                    pass
            sleep(0.2)
        el = driver.find_element(By.ID, "accompanying-diagnosis")
        el.send_keys(dv.normalize_mkb(code), Keys.ENTER)
        sleep(0.3)


def dump_claim_errors(driver: WebDriver, *, click_button: bool = True) -> str:
    """Текст валидации / snackbar; опционально кнопка ОШИБКИ."""
    parts: list[str] = []
    form_errs = dv.collect_form_errors(driver)
    if form_errs:
        parts.append(form_errs)
    snack = dv.snackbar_text(driver)
    if snack:
        parts.append(snack)
    if not click_button:
        return " | ".join(p for p in parts if p)

    btns = driver.find_elements(By.ID, "claim-error")
    if not btns:
        btns = [
            e
            for e in driver.find_elements(By.XPATH, "//button[contains(.,'ОШИБК') or contains(.,'Ошибк')]")
            if e.is_displayed()
        ]
    if btns:
        try:
            btns[0].click()
        except Exception:
            driver.execute_script("arguments[0].click();", btns[0])
        sleep(0.5)
        dlg = dv.warning_dialog_text(driver)
        if dlg:
            parts.append(dlg)
        else:
            for sel in (
                "[role='dialog']",
                ".MuiDialog-paper",
                ".MuiPopover-paper",
                ".MuiDrawer-paper",
            ):
                for el in driver.find_elements(By.CSS_SELECTOR, sel):
                    try:
                        if el.is_displayed() and (el.text or "").strip():
                            parts.append((el.text or "").strip()[:800])
                    except Exception:
                        continue
        for xp in (
            "//button[contains(.,'Закрыть') or contains(.,'OK') or contains(.,'Ок')]",
            "//button[@aria-label='Close']",
        ):
            for b in driver.find_elements(By.XPATH, xp):
                try:
                    if b.is_displayed():
                        b.click()
                        break
                except Exception:
                    continue
    text = " | ".join(p for p in parts if p)
    if text:
        logger.warning("Ошибки талона: {}", text[:1000])
    return text


def read_snackbar(driver: WebDriver) -> str:
    """Всплывающее уведомление #client-snackbar (как snackbar_page в ноутбуке)."""
    # 1) DOM даже если анимация / display
    try:
        text = driver.execute_script(
            """
            const el = document.getElementById('client-snackbar');
            if (!el) return '';
            const p = el.querySelector('p');
            const t = ((p && p.innerText) || el.innerText || el.textContent || '').trim();
            return t;
            """
        )
        if text:
            return str(text).strip()
    except Exception:
        pass
    # 2) selenium helpers из oms_dv_opv
    text = dv.snackbar_text(driver)
    if text:
        return text
    # 3) MUI / notistack
    try:
        text = driver.execute_script(
            """
            const sels = [
              '.MuiSnackbar-root .MuiAlert-message',
              '.MuiSnackbarContent-message',
              '.notistack-Snackbar',
              '[class*="Snackbar"] [class*="message"]',
            ];
            for (const s of sels) {
              for (const el of document.querySelectorAll(s)) {
                const t = (el.innerText || '').trim();
                if (t) return t;
              }
            }
            return '';
            """
        )
        if text:
            return str(text).strip()
    except Exception:
        pass
    return ""


def wait_snackbar(driver: WebDriver, *, timeout: float = 20, ignore: str = "") -> str:
    """Ждём появления уведомления после Save (игнор мигающего Error)."""
    deadline = time() + timeout
    last = ""
    while time() < deadline:
        text = read_snackbar(driver)
        if text and text != ignore:
            # краткий «Error» часто мелькает до реального текста / диалога
            if text.strip().lower() == "error":
                dlg = dv.warning_dialog_text(driver)
                if dlg:
                    logger.info("Диалог вместо snackbar Error: {}", dlg[:300])
                    return f"[dialog] {dlg}"
                last = text
                sleep(0.35)
                continue
            logger.info("Snackbar: {}", text[:300])
            return text
        dlg = dv.warning_dialog_text(driver)
        if dlg and (
            "предупрежд" in dlg.lower()
            or "внимание" in dlg.lower()
            or "проверк" in dlg.lower()
            or "данные пациента были изменены" in dlg.lower()
            or "сохранить изменения" in dlg.lower()
            or "для талона" in dlg.lower()
        ):
            logger.info("Диалог вместо snackbar: {}", dlg[:300])
            return f"[dialog] {dlg}"
        last = text or last
        sleep(0.25)
    return last


def classify_snackbar(text: str) -> tuple[bool | None, str]:
    """True=OK, False=FAIL, None=неясно (смотреть диалоги)."""
    raw = (text or "").strip()
    if not raw:
        return None, ""
    # диалог передаём дальше
    if raw.startswith("[dialog]"):
        return None, raw
    low = raw.lower()
    fail_markers = (
        "заполните обязательн",
        "ошибка",
        "онкология:",
        "дубликат",
        "не найден",
        "не сохран",
        "отсутствует пациент",
        "необходимо",
        "не указан",
        "превышает текущую дату",
        "дата окончания лечения превышает",
    )
    if any(x in low for x in fail_markers):
        return False, raw
    ok_markers = (
        "сохранен",
        "сохранён",
        "сохранено",
        "успешн",
        "создан",
        "записан",
        "добавлен",
        "талон сохран",
    )
    if any(x in low for x in ok_markers):
        return True, raw
    return None, raw


def _click_save_button(driver: WebDriver) -> str:
    """Как в ноутбуке: //*[@id='save-button']."""
    dv.dismiss_overlays(driver)
    if driver.find_elements(By.ID, "main-tab"):
        try:
            dv.click_id(driver, "main-tab")
            dv.wait_busy_gone(driver, timeout=20)
            sleep(0.4)
        except Exception:
            pass

    deadline = time() + 12
    while time() < deadline:
        els = driver.find_elements(By.XPATH, '//*[@id="save-button"]')
        for el in els:
            try:
                if not el.is_displayed():
                    continue
                label = ((el.text or "") + " " + (el.get_attribute("aria-label") or "")).upper()
                if "ПАЦИЕНТ" in label and "СОХРАНИТЬ (F2)" not in label and "F2" not in label:
                    # пропускаем только явные кнопки карты пациента
                    if "СОХРАНИТЬ" not in label:
                        continue
                driver.execute_script("arguments[0].scrollIntoView({block:'center'});", el)
                sleep(0.1)
                try:
                    el.click()
                except Exception:
                    driver.execute_script("arguments[0].click();", el)
                logger.info("Клик save-button: «{}»", (el.text or "").strip()[:50])
                return "save-button"
            except Exception:
                continue
        sleep(0.3)
    raise RuntimeError("Кнопка #save-button не найдена")


def save_ambulatory(driver: WebDriver) -> tuple[bool, str]:
    """Save → читаем #client-snackbar / диалог (логика талоны2_0.save)."""
    snack_before = read_snackbar(driver)

    try:
        _click_save_button(driver)
    except RuntimeError as exc:
        return False, str(exc)

    dv.wait_busy_gone(driver, timeout=60)
    snack = wait_snackbar(driver, timeout=18, ignore=snack_before)
    verdict, detail = classify_snackbar(snack)

    if verdict is False:
        return False, detail or snack or "ошибка по уведомлению"

    # предупреждения (Проверка врачей / Пересечение) — подтвердить
    dlg = dv.warning_dialog_text(driver)
    if (not dlg) and detail.startswith("[dialog]"):
        dlg = detail[len("[dialog]") :].strip()
    if dlg or dv.find_visible_dialogs(driver):
        low = (dlg or "").lower()
        if "дубликат" in low:
            return False, dlg or "Дубликат талона"
        if "результат поиска в црп" in low or "отсутствует пациент" in low:
            return False, dlg
        # «Данные пациента были изменены» → ДЛЯ ТАЛОНА И ПАЦИЕНТА
        # (confirm_warning_dialog / handle_post_save_prompts)
        ok, wmsg = dv.handle_post_save_prompts(driver, timeout=30)
        if ok:
            # после подтверждения ещё раз смотрим snackbar
            snack2 = wait_snackbar(driver, timeout=8, ignore=snack)
            v2, d2 = classify_snackbar(snack2)
            if v2 is False and "данные пациента были изменены" not in (d2 or "").lower():
                return False, d2 or wmsg
            if v2 is True:
                return True, d2 or wmsg or snack or "сохранено (предупреждение подтверждено)"
            return True, d2 or wmsg or snack or "сохранено (предупреждение подтверждено)"
        return False, wmsg or dlg or "диалог не подтверждён"

    if verdict is True:
        return True, detail

    # пустой/неясный snackbar — ещё одна попытка прочитать
    snack3 = read_snackbar(driver)
    v3, d3 = classify_snackbar(snack3)
    if v3 is True:
        return True, d3
    if v3 is False:
        return False, d3
    if snack3:
        logger.warning("Неясное уведомление: {}", snack3[:300])
        return False, f"неясное уведомление: {snack3}"
    return False, "нет уведомления после Save — талон не подтверждён"


def process_row(driver: WebDriver, row: Talon307, *, index: int, dry_run: bool) -> RowResult:
    t0 = time()
    logger.info("=== [{}] 307 {} {} ===", index, row.enp, row.service_code)
    try:
        open_ambulatory(driver)
        fill_patient_search(driver, row)
        if (row.social_status or "").strip() and (row.occupation or "").strip():
            logger.info("Соцстатус/занятость из Excel: {} / {}", row.social_status, row.occupation)
        else:
            ensure_social_status(driver, row)

        fill_referral(driver, row)
        fill_main(driver, row)
        fill_single_service(driver, row)

        if dry_run:
            soft = dump_claim_errors(driver, click_button=False)
            msg = f"dry-run: форма заполнена, услуга {row.service_code}"
            if soft:
                msg += f" | {soft[:200]}"
            logger.success("[{}] {}", index, msg)
            return RowResult(index, row.enp, row.service_code, "dry_run", msg, round(time() - t0, 2))

        ok, msg = save_ambulatory(driver)
        logger.info("[{}] уведомление/итог Save: {}", index, msg)
        if (not ok) and dv._is_social_status_error(msg):
            logger.warning("[{}] ошибка соцстатуса — повторяем", index)
            ensure_social_status(driver, row)
            fill_main(driver, row)
            fill_single_service(driver, row)
            ok, msg = save_ambulatory(driver)
            logger.info("[{}] уведомление/итог Save (повтор): {}", index, msg)

        status = "ok" if ok else "fail"
        log = logger.success if ok else logger.error
        log("[{}] {} → {} ({:.1f}s)", index, status.upper(), msg, time() - t0)
        return RowResult(index, row.enp, row.service_code, status, msg, round(time() - t0, 2))
    except Exception as exc:
        msg = str(exc).split("Stacktrace:")[0].strip() or str(exc)
        logger.error("[{}] error ({:.1f}s): {}", index, time() - t0, msg)
        try:
            driver.refresh()
            dv.wait_busy_gone(driver)
        except Exception:
            pass
        return RowResult(index, row.enp, row.service_code, "error", msg, round(time() - t0, 2))


def fill_single_service(driver: WebDriver, row: Talon307) -> None:
    """Одна услуга: код + даты + кол-во (+ врач при наличии)."""
    code = (row.service_code or "").strip().upper()
    if not code:
        raise RuntimeError("Пустой Код услуги")

    if driver.find_elements(By.ID, "services-tab"):
        try:
            dv.click_id(driver, "services-tab")
            dv.wait_busy_gone(driver, timeout=8)
            sleep(0.15)
        except Exception:
            pass

    svc = driver.find_elements(By.ID, "combobox-service-0")
    if not svc:
        add = driver.find_elements(By.ID, "add-service")
        if add:
            try:
                add[0].click()
            except Exception:
                driver.execute_script("arguments[0].click();", add[0])
            sleep(0.25)
            svc = driver.find_elements(By.ID, "combobox-service-0")
    if not svc:
        raise RuntimeError("Нет поля combobox-service-0")

    _fill_select_fast(driver, svc[0], code)
    sleep(0.15)

    begin = dv.normalize_service_date(row.begin)
    end = dv.normalize_service_date(row.end)
    # даты услуги: input часто «not reachable by keyboard» → JS
    if driver.find_elements(By.ID, "begin-date-service-0"):
        if not _force_input(driver, "begin-date-service-0", begin):
            try:
                dv.fill_date_field(driver, "begin-date-service-0", row.begin)
            except Exception as exc:
                logger.warning("begin-date-service-0: {}", exc)
    if driver.find_elements(By.ID, "end-date-service-0"):
        if not _force_input(driver, "end-date-service-0", end):
            try:
                dv.fill_date_field(driver, "end-date-service-0", row.end)
            except Exception as exc:
                logger.warning("end-date-service-0: {}", exc)

    amount = (row.service_amount or "1").strip() or "1"
    if driver.find_elements(By.ID, "amount-service-0"):
        if not _force_input(driver, "amount-service-0", amount):
            el = driver.find_element(By.ID, "amount-service-0")
            try:
                el.clear()
                el.send_keys(amount, Keys.ENTER)
            except Exception:
                driver.execute_script(
                    "arguments[0].value=arguments[1];"
                    "arguments[0].dispatchEvent(new Event('input',{bubbles:true}));"
                    "arguments[0].dispatchEvent(new Event('change',{bubbles:true}));",
                    el,
                    amount,
                )
        sleep(0.08)

    for sid in ("doctor-service-0", "combobox-doctor-0"):
        els = driver.find_elements(By.ID, sid)
        if not els:
            continue
        target = els[-1]
        try:
            if not target.is_displayed():
                continue
            _fill_select_fast(driver, target, row.doctor)
            break
        except Exception as exc:
            logger.warning("Врач услуги ({}): {}", sid, exc)

    logger.info(
        "Услуга {} × {} | {}…{} (явки МО={})",
        code,
        amount,
        begin,
        end,
        row.mo_visits,
    )


def _read_birth_anywhere(driver: WebDriver) -> str:
    birth = dv._read_birth_from_main_tab(driver)
    if birth:
        return birth
    for element_id in ("birth-date", "birthDate", "patient-birth-date"):
        els = driver.find_elements(By.ID, element_id)
        if els:
            val = (els[0].get_attribute("value") or els[0].text or "").strip()
            if val:
                return val
    # подписи «Дата рождения» рядом
    for el in driver.find_elements(By.XPATH, "//*[contains(translate(.,'ДАТАРОЖДЕНИЯ','датарождения'),'дата рождения')]"):
        try:
            blob = (el.text or "").strip()
            m = re.search(r"(\d{1,2}[-./]\d{1,2}[-./]\d{2,4})", blob)
            if m:
                return m.group(1)
        except Exception:
            continue
    return ""


def ensure_social_status(driver: WebDriver, row: Talon307) -> tuple[str, str]:
    """Соцстатус/занятость: Excel → ДР на форме → 22/3 (логика возраста как в ДВ4)."""
    soc = (row.social_status or "").strip()
    occ = (row.occupation or "").strip()

    birth_raw = _read_birth_anywhere(driver)
    if (not soc or not occ) and birth_raw:
        year = dv._birth_year(birth_raw)
        if year is not None:
            age = datetime.now().year - year
            soc, occ = dv.social_occupation_by_age(age)
            logger.info(
                "ДР={} → возраст≈{} → socialStatus={}, occupation={}",
                birth_raw,
                age,
                soc,
                occ,
            )

    if not soc or not occ:
        soc, occ = "22", "3"
        logger.info("Соцстатус по умолчанию (как ДВ4 ≥60): {} / {}", soc, occ)

    if driver.find_elements(By.ID, "socialStatus"):
        dv.try_input_enter_id(driver, "socialStatus", soc)
    if driver.find_elements(By.ID, "occupation"):
        dv.try_input_enter_id(driver, "occupation", occ)

    row.social_status, row.occupation = soc, occ
    logger.info("Соцстатус для талона: {} / {}", soc, occ)
    return soc, occ


def _patient_loaded(driver: WebDriver) -> bool:
    """Признак, что после поиска подтянулась карта (не пустая форма)."""
    for element_id in ("subdivision", "doctor", "requestPurpose"):
        els = driver.find_elements(By.ID, element_id)
        if not els:
            continue
        try:
            val = (els[0].get_attribute("value") or "").strip()
            # react-select: смотрим текст обёртки
            if val:
                return True
            wrap = dv._closest_react_select(els[0])
            if wrap is not None and (wrap.text or "").strip():
                return True
        except Exception:
            continue
    # ФИО / таблица пациента
    for xp in (
        "//*[contains(@class,'patient')]",
        "//table//tbody/tr",
        "//*[contains(text(),'Пол') or contains(text(),'пол')]",
    ):
        for el in driver.find_elements(By.XPATH, xp):
            try:
                if el.is_displayed() and (el.text or "").strip():
                    return True
            except Exception:
                continue
    return False


def fill_patient_search(driver: WebDriver, row: Talon307) -> None:
    """Даты + ЕНП + поиск (талоны2_0.talon_pacient)."""
    dv.dismiss_overlays(driver)
    begin = dv.normalize_service_date(row.begin)
    end = dv.normalize_service_date(row.end)
    if not _force_input(driver, "begin-date", begin):
        dv.fill_date_field(driver, "begin-date", row.begin)
    if not _force_input(driver, "end-date", end):
        dv.fill_date_field(driver, "end-date", row.end)
    dv.click_id(driver, "enp")
    el_enp = driver.find_element(By.ID, "enp")
    el_enp.clear()
    el_enp.send_keys(row.enp)

    _click_xpath(
        driver,
        "/html/body/div[1]/div/div[2]/div[3]/div[1]/div/div[1]/div[2]/div/div[2]/div/div/div/label/span[1]",
    )
    sleep(0.08)
    searched = _click_xpath(
        driver,
        "/html/body/div[1]/div/div[2]/div[3]/div[1]/div/div[1]/div[2]/div/div[4]/button",
    )
    if not searched:
        for xp in (
            '//button[.//span[contains(.,"Поиск")] or contains(.,"Найти")]',
            '//*[@id="enp"]/../../..//button',
        ):
            if _click_xpath(driver, xp):
                searched = True
                break
    if not searched:
        el_enp.send_keys(Keys.ENTER)

    dv.wait_busy_gone(driver, timeout=40)
    sleep(0.35)
    err = dv.snackbar_text(driver)
    low = (err or "").lower()
    if err and any(
        x in low
        for x in ("пациент не найден", "некорректн", "ошибка поиска", "отсутствует пациент", "результат поиска в црп")
    ):
        raise RuntimeError(f"Поиск пациента: {err}")
    if err:
        logger.info("После поиска snackbar: {}", err)

    for xp in (
        "//div[contains(@class,'Select-menu')]//div[contains(@class,'Select-option')]",
        "//div[@role='option']",
        "//table//tbody/tr[1]",
    ):
        opts = [e for e in driver.find_elements(By.XPATH, xp) if e.is_displayed()]
        if opts:
            try:
                opts[0].click()
                dv.wait_busy_gone(driver, timeout=15)
                sleep(0.15)
            except Exception:
                pass
            break

    try:
        _wait(driver, 10).until(EC.presence_of_element_located((By.ID, "doctor")))
    except TimeoutException as exc:
        raise RuntimeError("После поиска нет поля doctor") from exc

    deadline = time() + 5
    while time() < deadline:
        if _patient_loaded(driver) or _read_birth_anywhere(driver):
            break
        sleep(0.2)
    else:
        logger.warning("После поиска признаки карты пациента слабые — продолжаем")


def load_talons(path: Path, *, default_building: str = "") -> list[Talon307]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["талоны"] if "талоны" in wb.sheetnames else wb.active
    rows_iter = ws.iter_rows(values_only=True)
    header = [_cell(h).lower() for h in next(rows_iter)]
    idx = {h: i for i, h in enumerate(header) if h}

    def col(*names: str) -> int | None:
        for n in names:
            if n.lower() in idx:
                return idx[n.lower()]
        return None

    i_enp = col("енп", "enp")
    i_beg = col("начало", "begin", "дата начала")
    i_end = col("окончание", "end", "дата окончания")
    i_doc = col("врач", "doctor")
    i_bldg = col("корпус", "building", "subdivision")
    i_goal = col("цель", "goal", "requestpurpose")
    i_sl1 = col("случай", "occurrence")
    i_sl2 = col("случай_2", "случай 2", "second-occurrence", "second_occurrence")
    i_mo = col("посещений в мо", "посещений", "явки", "mo")
    i_amt = col("кол-во", "количество", "amount", "колво")
    i_dx = col("диагноз", "мкб", "diagnosis", "ds")
    i_dx2 = col("диагноз 2", "диагноз2", "accompanying")
    i_dn = col("дисп. набл", "дн", "dispensary")
    i_type = col("тип", "type")
    i_res = col("результат", "result")
    i_out = col("исход", "outcome")
    i_har = col("характер", "character")
    i_svc = col("код услуги", "услуга", "service_code", "service")
    i_soc = col("соцстатус", "социальный статус", "social_status")
    i_occ = col("вид занятости", "занятость", "occupation")

    if i_enp is None or i_beg is None or i_end is None or i_doc is None or i_svc is None or i_dx is None:
        raise SystemExit("В Excel нужны колонки: ЕНП, начало, окончание, врач, Диагноз, Код услуги")

    out: list[Talon307] = []
    for n, raw in enumerate(rows_iter, start=2):
        if not raw or all(v is None or str(v).strip() == "" for v in raw):
            continue
        enp = _cell(raw[i_enp])
        if not enp:
            continue
        building = _cell(raw[i_bldg]) if i_bldg is not None else ""
        building = building or default_building
        out.append(
            Talon307(
                enp=enp,
                begin=_cell(raw[i_beg]),
                end=_cell(raw[i_end]),
                doctor=_cell(raw[i_doc]),
                building=building,
                goal=_cell(raw[i_goal]) if i_goal is not None else DEFAULTS_307["цель"],
                occurrence=_cell(raw[i_sl1]) if i_sl1 is not None else DEFAULTS_307["случай"],
                occurrence2=_cell(raw[i_sl2]) if i_sl2 is not None else DEFAULTS_307["случай_2"],
                mo_visits=_cell(raw[i_mo]) if i_mo is not None else "4",
                diagnosis=_cell(raw[i_dx]),
                diagnosis2=_cell(raw[i_dx2]) if i_dx2 is not None else "",
                dn=_cell(raw[i_dn]) if i_dn is not None else DEFAULTS_307["ДИСП. НАБЛ"],
                claim_type=_cell(raw[i_type]) if i_type is not None else DEFAULTS_307["тип"],
                result=_cell(raw[i_res]) if i_res is not None else DEFAULTS_307["результат"],
                outcome=_cell(raw[i_out]) if i_out is not None else DEFAULTS_307["исход"],
                character=_cell(raw[i_har]) if i_har is not None else DEFAULTS_307["характер"],
                service_code=_cell(raw[i_svc]).upper(),
                service_amount=_cell(raw[i_amt]) if i_amt is not None else "1",
                social_status=_cell(raw[i_soc]) if i_soc is not None else "",
                occupation=_cell(raw[i_occ]) if i_occ is not None else "",
                raw_index=n,
            )
        )
    wb.close()
    return out


def build_from_short_source(
    source: Path,
    *,
    out_path: Path,
    default_building: str = "",
    clear_social: bool = True,
) -> int:
    """Короткий файл (ЕНП/МКБ/услуга/явки/даты/врач) → один шаблон talons_307.xlsx."""
    wb = load_workbook(source, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        raise SystemExit(f"Пустой файл: {source}")

    header = [_cell(h).lower() for h in rows[0]]
    idx = {h: i for i, h in enumerate(header) if h}

    def col(*names: str) -> int | None:
        for n in names:
            if n in idx:
                return idx[n]
        return None

    i_enp = col("енп", "enp")
    i_mkb = col("мкб", "диагноз")
    i_svc = col("код услуги", "услуга")
    i_vis = col("явки", "посещений в мо", "кол-во")
    i_beg = col("дата начала", "начало")
    i_end = col("дата окончания", "окончание")
    i_doc = col("врач")
    i_bldg = col("корпус")
    if None in (i_enp, i_mkb, i_svc, i_beg, i_end):
        raise SystemExit("В источнике нужны ЕНП, МКБ, Код услуги, Дата начала, Дата окончания")

    out_rows: list[list[object]] = []
    for raw in rows[1:]:
        if not raw:
            continue
        enp = _cell(raw[i_enp])
        mkb = _cell(raw[i_mkb])
        code = _cell(raw[i_svc]).upper()
        begin = _cell(raw[i_beg])
        end = _cell(raw[i_end])
        if not enp or not mkb or not code or not begin or not end:
            continue
        visits = _cell(raw[i_vis]) if i_vis is not None else "4"
        doctor = _cell(raw[i_doc]) if i_doc is not None else ""
        building = (_cell(raw[i_bldg]) if i_bldg is not None else "") or default_building
        out_rows.append(
            [
                enp,
                begin,
                end,
                doctor,
                building,
                DEFAULTS_307["цель"],
                DEFAULTS_307["случай"],
                DEFAULTS_307["случай_2"],
                visits,  # посещений в МО = явки школы
                mkb,
                "",
                DEFAULTS_307["ДИСП. НАБЛ"],
                DEFAULTS_307["тип"],
                DEFAULTS_307["результат"],
                DEFAULTS_307["исход"],
                DEFAULTS_307["характер"],
                code,
                "1",  # кол-во услуги = 1 (одна услуга на талон)
                "" if clear_social else "22",
                "" if clear_social else "3",
            ]
        )

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb_out = Workbook()
    ws_out = wb_out.active
    ws_out.title = "талоны"
    ws_out.append(TEMPLATE_COLUMNS)
    for row in out_rows:
        ws_out.append(row)
    wb_out.save(out_path)
    return len(out_rows)


def ensure_template(path: Path) -> None:
    if path.is_file():
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "талоны"
    ws.append(TEMPLATE_COLUMNS)
    ws.append(
        [
            "3153999734000161",
            "08-09-26",
            "11-09-26",
            "11097116",
            "ГП №11",
            "307",
            "Первичный",
            "Законченный",
            "4",
            "K29.5",
            "",
            "1",
            "3",
            "301",
            "304",
            "3",
            "B04.004.002",
            "1",
            "",
            "",
        ]
    )
    wb.save(path)
    logger.info("Создан шаблон: {}", path)


def write_results(path: Path, results: list[RowResult]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = list(asdict(results[0]).keys()) if results else [
        "index", "enp", "service_code", "status", "message", "elapsed_sec"
    ]
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        for r in results:
            w.writerow(asdict(r))


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Ввод талонов цели 307 (Web.ОМС ambulatory)")
    p.add_argument("--file", type=Path, default=DEFAULT_FILE)
    p.add_argument("--limit", type=int, default=0)
    p.add_argument("--offset", type=int, default=0)
    p.add_argument("--building", default="", help="Корпус по умолчанию, если в Excel пусто")
    p.add_argument("--dry-run", action="store_true")
    p.add_argument("--keep-open", type=int, default=0)
    p.add_argument("--make-template", action="store_true")
    p.add_argument("--worker", default="", help="Метка воркера для логов/результатов (w0…)")
    p.add_argument("--window-x", type=int, default=None)
    p.add_argument("--window-y", type=int, default=None)
    p.add_argument(
        "--from-source",
        type=Path,
        default=None,
        help="Короткий Excel → один talons_307.xlsx",
    )
    return p.parse_args()


def main() -> int:
    args = parse_args()
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)

    if args.make_template:
        ensure_template(args.file)
        return 0

    if args.from_source:
        src = args.from_source
        if not src.is_file():
            src = WORK_DIR / "data" / src.name
        n = build_from_short_source(src, out_path=args.file, default_building=args.building)
        logger.info("Собран шаблон {} строк → {}", n, args.file)
        return 0

    if not args.file.is_file():
        ensure_template(args.file)
        logger.error("Нет файла {}. Заполните шаблон и запустите снова.", args.file)
        return 1

    rows = load_talons(args.file, default_building=args.building)
    if args.offset:
        rows = rows[args.offset :]
    if args.limit and args.limit > 0:
        rows = rows[: args.limit]

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    worker_tag = (args.worker or "").strip() or f"o{args.offset}"
    result_path = RESULTS_DIR / f"oms_307_result_{stamp}_{worker_tag}.csv"
    logger.info(
        "[{}] Строк: {} | dry_run={} | offset={} | building={!r}",
        worker_tag,
        len(rows),
        args.dry_run,
        args.offset,
        args.building,
    )

    results: list[RowResult] = []
    t_all = time()
    win_kwargs: dict = {"implicit_wait": 0, "maximize": False}
    if args.window_x is not None and args.window_y is not None:
        win_kwargs["window_position"] = (args.window_x, args.window_y)
        win_kwargs["window_size"] = (1100, 850)
    elif args.window_x is not None or args.window_y is not None:
        win_kwargs["window_position"] = (args.window_x or 20, args.window_y or 20)
        win_kwargs["window_size"] = (1100, 850)
    else:
        win_kwargs["maximize"] = True

    with GeckoBrowser(**win_kwargs) as session:
        driver = session.driver
        driver.implicitly_wait(0)
        login_oms(driver)
        sleep(0.5)
        dv.wait_busy_gone(driver, timeout=30)

        for i, row in enumerate(rows, start=1 + args.offset):
            res = process_row(driver, row, index=i, dry_run=args.dry_run)
            results.append(res)
            # промежуточный CSV — чтобы при обрыве воркера был прогресс
            if len(results) % 10 == 0:
                write_results(result_path, results)

        if args.keep_open > 0:
            sleep(args.keep_open)

    write_results(result_path, results)
    ok_n = sum(1 for r in results if r.status in ("ok", "dry_run"))
    fail_n = len(results) - ok_n
    logger.info(
        "——— [{}] ИТОГ ({:.0f}s) OK={}, FAIL={} → {} ———",
        worker_tag,
        time() - t_all,
        ok_n,
        fail_n,
        result_path,
    )
    return 0 if fail_n == 0 and results else 1


if __name__ == "__main__":
    sys.exit(main())
