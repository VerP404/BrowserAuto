#!/usr/bin/env python
"""
Шаблон и сборка Excel для ввода талонов цели 307 (амбулаторный талон Web.ОМС).

Два файла, как в талоны2_0.ipynb / списки:
  - talons_307.xlsx  — карта талона
  - services_307.xlsx — услуги (связь по колонке «Талон»)

Источник «короткого» экспорта портала (ЕНП, МКБ, Код услуги, Явки, даты, Врач)
дополняется константами цели 307 и датами в формате дд-мм-гг.

Примеры:
  python template_307.py --make-template
  python template_307.py --from-source "data/талоны для 307.xlsx"
"""

from __future__ import annotations

import argparse
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

WORK_DIR = Path(__file__).resolve().parent
DATA_DIR = WORK_DIR / "data"

TALON_COLUMNS = [
    "Талон",
    "начало",
    "окончание",
    "ЕНП",
    "врач",
    "Корпус",
    "цель",
    "случай",
    "случай_2",
    "посещений в МО",
    "Диагноз",
    "Диагноз 2",
    "ДИСП. НАБЛ",
    "тип",
    "результат",
    "исход",
    "характер",
    "Соцстатус",
    "Вид занятости",
]

SERVICE_COLUMNS = [
    "Талон",
    "Код услуги",
    "Дата начала",
    "Дата окончания",
    "Кол-во",
    "Врач",
]

# Константы цели 307 (как в списки/талоны.xlsx и MosaicMed HEADER_307)
DEFAULTS_307 = {
    "цель": "307",
    "случай": "Первичный",
    "случай_2": "Законченный",
    "ДИСП. НАБЛ": "1",
    "тип": "3",
    "результат": "301",
    "исход": "304",
    "характер": "3",
    "Соцстатус": "22",
    "Вид занятости": "3",
}

SAMPLE_TALON = {
    "Талон": "1",
    "начало": "08-09-26",
    "окончание": "11-09-26",
    "ЕНП": "3153999734000161",
    "врач": "11097116",
    "Корпус": "ГП №11",
    **DEFAULTS_307,
    "посещений в МО": "4",
    "Диагноз": "K29.5",
    "Диагноз 2": "",
}

SAMPLE_SERVICE = {
    "Талон": "1",
    "Код услуги": "B04.004.002",
    "Дата начала": "08-09-26",
    "Дата окончания": "11-09-26",
    "Кол-во": "4",
    "Врач": "11097116",
}


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def to_oms_date(value: object) -> str:
    """Любая дата → дд-мм-гг (формат ноутбуков / списки/талоны.xlsx)."""
    if value is None or value == "":
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d-%m-%y")
    if isinstance(value, date):
        return value.strftime("%d-%m-%y")
    text = str(value).strip()
    if not text or text.lower() in {"nan", "none", "-", "—"}:
        return ""
    # Excel serial
    if re.fullmatch(r"\d+(\.0)?", text):
        try:
            from openpyxl.utils.datetime import from_excel

            return from_excel(float(text)).strftime("%d-%m-%y")
        except Exception:
            pass
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", text)
    if m:
        y, mo, d = m.groups()
        return f"{int(d):02d}-{int(mo):02d}-{y[-2:]}"
    m = re.match(r"^(\d{1,2})[./-](\d{1,2})[./-](\d{2,4})$", text)
    if m:
        d, mo, y = m.groups()
        if len(y) == 4:
            y = y[-2:]
        return f"{int(d):02d}-{int(mo):02d}-{y}"
    return text


def _norm_header(name: object) -> str:
    return re.sub(r"\s+", " ", _cell(name)).strip().lower()


# Короткие имена колонок исходного файла портала / due-экспорта
_SOURCE_ALIASES = {
    "енп": "enp",
    "enp": "enp",
    "мкб": "mkb",
    "диагноз": "mkb",
    "диагноз 1": "mkb",
    "код услуги": "service_code",
    "услуга": "service_code",
    "явки": "visits",
    "посещений в мо": "visits",
    "кол-во": "visits",
    "дата начала": "date_start",
    "начало": "date_start",
    "дата окончания": "date_end",
    "окончание": "date_end",
    "врач": "doctor",
    "корпус": "building",
    "диагноз 2": "mkb2",
    "соцстатус": "social",
    "вид занятости": "occupation",
}


def _map_source_headers(header_row: tuple) -> dict[str, int]:
    out: dict[str, int] = {}
    for i, raw in enumerate(header_row):
        key = _SOURCE_ALIASES.get(_norm_header(raw))
        if key and key not in out:
            out[key] = i
    return out


def _write_sheet(path: Path, title: str, columns: list[str], rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append(columns)
    for row in rows:
        ws.append([row.get(col, "") for col in columns])
    wb.save(path)


def write_empty_templates(
    *,
    talons_path: Path | None = None,
    services_path: Path | None = None,
    with_sample: bool = True,
) -> tuple[Path, Path]:
    talons_path = talons_path or (DATA_DIR / "talons_307.xlsx")
    services_path = services_path or (DATA_DIR / "services_307.xlsx")
    talon_rows = [SAMPLE_TALON] if with_sample else []
    service_rows = [SAMPLE_SERVICE] if with_sample else []
    _write_sheet(talons_path, "талоны", TALON_COLUMNS, talon_rows)
    _write_sheet(services_path, "услуги", SERVICE_COLUMNS, service_rows)
    return talons_path, services_path


def build_from_source(
    source: Path,
    *,
    default_building: str = "",
    talons_path: Path | None = None,
    services_path: Path | None = None,
) -> tuple[Path, Path, int]:
    """Короткий Excel → пара файлов шаблона 307."""
    wb = load_workbook(source, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        raise SystemExit(f"Пустой файл: {source}")

    col = _map_source_headers(tuple(rows[0]))
    required = ("enp", "mkb", "service_code", "date_start", "date_end")
    missing = [k for k in required if k not in col]
    if missing:
        raise SystemExit(f"В источнике нет колонок: {', '.join(missing)}")

    talons: list[dict[str, str]] = []
    services: list[dict[str, str]] = []
    plan_id = 0
    for raw in rows[1:]:
        if raw is None:
            continue
        enp = _cell(raw[col["enp"]])
        mkb = _cell(raw[col["mkb"]])
        code = _cell(raw[col["service_code"]]).upper()
        begin = to_oms_date(raw[col["date_start"]])
        end = to_oms_date(raw[col["date_end"]])
        if not enp or not mkb or not code or not begin or not end:
            continue
        visits_raw = raw[col["visits"]] if "visits" in col else None
        visits = _cell(visits_raw) or "1"
        doctor = _cell(raw[col["doctor"]]) if "doctor" in col else ""
        building = _cell(raw[col["building"]]) if "building" in col else ""
        building = building or default_building
        mkb2 = _cell(raw[col["mkb2"]]) if "mkb2" in col else ""
        social = _cell(raw[col["social"]]) if "social" in col else DEFAULTS_307["Соцстатус"]
        occupation = (
            _cell(raw[col["occupation"]]) if "occupation" in col else DEFAULTS_307["Вид занятости"]
        )

        plan_id += 1
        tid = str(plan_id)
        talons.append(
            {
                "Талон": tid,
                "начало": begin,
                "окончание": end,
                "ЕНП": enp,
                "врач": doctor,
                "Корпус": building,
                **DEFAULTS_307,
                "посещений в МО": visits,
                "Диагноз": mkb,
                "Диагноз 2": mkb2,
                "Соцстатус": social or DEFAULTS_307["Соцстатус"],
                "Вид занятости": occupation or DEFAULTS_307["Вид занятости"],
            }
        )
        services.append(
            {
                "Талон": tid,
                "Код услуги": code,
                "Дата начала": begin,
                "Дата окончания": end,
                "Кол-во": visits,
                "Врач": doctor,
            }
        )

    talons_path = talons_path or (DATA_DIR / "talons_307.xlsx")
    services_path = services_path or (DATA_DIR / "services_307.xlsx")
    _write_sheet(talons_path, "талоны", TALON_COLUMNS, talons)
    _write_sheet(services_path, "услуги", SERVICE_COLUMNS, services)
    return talons_path, services_path, plan_id


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description="Шаблон Excel для цели 307 (Web.ОМС ambulatory)")
    p.add_argument("--make-template", action="store_true", help="Пустой шаблон с 1 примером")
    p.add_argument(
        "--from-source",
        type=Path,
        help="Короткий Excel (ЕНП/МКБ/услуга/явки/даты/врач) → полный шаблон",
    )
    p.add_argument("--building", default="", help="Корпус по умолчанию, если в источнике пусто")
    p.add_argument("--talons-out", type=Path, default=None)
    p.add_argument("--services-out", type=Path, default=None)
    args = p.parse_args(argv)

    if args.make_template:
        t, s = write_empty_templates(
            talons_path=args.talons_out,
            services_path=args.services_out,
            with_sample=True,
        )
        print(f"Шаблон: {t}")
        print(f"Услуги: {s}")
        return 0

    if args.from_source:
        src = args.from_source
        if not src.is_file():
            src = DATA_DIR / src
        if not src.is_file():
            raise SystemExit(f"Файл не найден: {args.from_source}")
        t, s, n = build_from_source(
            src,
            default_building=args.building,
            talons_path=args.talons_out,
            services_path=args.services_out,
        )
        print(f"Строк: {n}")
        print(f"Талоны: {t}")
        print(f"Услуги: {s}")
        return 0

    p.print_help()
    return 1


if __name__ == "__main__":
    raise SystemExit(main())
