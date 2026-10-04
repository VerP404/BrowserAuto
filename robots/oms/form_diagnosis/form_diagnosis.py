#!/usr/bin/env python
"""
Исправление основного диагноза (ds1) в уже созданных амбулаторных талонах Web.ОМС.

Excel (колонки):
  Талон | ds1
  39586490 | R54

Если ds1 заполнено → открываем
  {OMS_BASE_URL}/claim/ambulatory/{Талон}
и ставим #mainDiagnosis = ds1, затем Save.

Пример:
  python run_fix_diagnosis.py --limit 1 --dry-run --keep-open 20
  python run_fix_diagnosis.py --file Талоны.xlsx --limit 10
  python run_fix_diagnosis.py --offset 0 --limit 50 --worker w0
"""

from __future__ import annotations

import argparse
import csv
import sys
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from time import sleep, time

from loguru import logger
from openpyxl import load_workbook
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.remote.webdriver import WebDriver
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait

WORK_DIR = Path(__file__).resolve().parent
OMS_DIR = WORK_DIR.parent
PROJECT_ROOT = WORK_DIR.parents[2]
for p in (PROJECT_ROOT, WORK_DIR, OMS_DIR / "dv_opv", OMS_DIR / "cel_307"):
    if str(p) not in sys.path:
        sys.path.insert(0, str(p))

from browser_auto.config import OMS_BASE_URL, load_credentials
from browser_auto.auth import login_oms
from browser_auto.driver import GeckoBrowser

import add_medical_exam as dv  # noqa: E402
import add_307 as base  # noqa: E402

load_credentials(WORK_DIR / "credentials.env", PROJECT_ROOT / ".env")

RESULTS_DIR = WORK_DIR / "data"
DEFAULT_FILE = PROJECT_ROOT / "Талоны.xlsx"


@dataclass
class FixRow:
    talon: str
    ds1: str


@dataclass
class RowResult:
    index: int
    talon: str
    ds1: str
    status: str
    message: str
    elapsed_sec: float


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v).strip()


def load_rows(path: Path, *, default_ds1: str = "") -> list[FixRow]:
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    headers = [(_cell(c.value).lower() if c.value is not None else "") for c in next(ws.iter_rows(min_row=1, max_row=1))]
    col = {h: i for i, h in enumerate(headers) if h}

    def idx(*names: str) -> int | None:
        for n in names:
            if n in col:
                return col[n]
        return None

    i_talon = idx("талон", "номер", "id", "claim", "n")
    i_ds1 = idx("ds1", "диагноз", "мкб", "diagnosis", "ds")
    if i_talon is None:
        raise SystemExit(f"Нужна колонка «Талон», сейчас: {headers}")
    if i_ds1 is None and not default_ds1:
        raise SystemExit(
            f"Нужна колонка «ds1» или флаг --diagnosis. Сейчас: {headers}"
        )

    out: list[FixRow] = []
    for raw in ws.iter_rows(min_row=2, values_only=True):
        if not raw or not any(v is not None and str(v).strip() != "" for v in raw):
            continue
        talon = _cell(raw[i_talon])
        ds1 = _cell(raw[i_ds1]) if i_ds1 is not None else ""
        if not ds1:
            ds1 = default_ds1
        if not talon:
            continue
        if not ds1:
            continue  # пустой ds1 — пропускаем (нечего исправлять)
        # номер талона без .0
        if talon.endswith(".0"):
            talon = talon[:-2]
        out.append(FixRow(talon=talon, ds1=dv.normalize_mkb(ds1)))
    return out


def open_claim(driver: WebDriver, talon: str) -> str:
    url = f"{OMS_BASE_URL.rstrip('/')}/claim/ambulatory/{talon}"
    logger.info("Открываем {}", url)
    driver.get(url)
    dv.wait_busy_gone(driver, timeout=60)
    sleep(0.6)
    # дождаться поля диагноза
    WebDriverWait(driver, 25).until(EC.presence_of_element_located((By.ID, "mainDiagnosis")))
    return url


def set_main_diagnosis(driver: WebDriver, mkb: str) -> None:
    mkb = dv.normalize_mkb(mkb)
    if not mkb:
        raise RuntimeError("Пустой ds1 / МКБ")
    if driver.find_elements(By.ID, "main-tab"):
        try:
            dv.click_id(driver, "main-tab")
            dv.wait_busy_gone(driver, timeout=20)
            sleep(0.3)
        except Exception:
            pass

    el = driver.find_element(By.ID, "mainDiagnosis")
    driver.execute_script("arguments[0].scrollIntoView({block:'center'});", el)
    sleep(0.1)
    # переиспользуем быстрый ввод react-select из 307
    base._enter_select(driver, "mainDiagnosis", mkb)
    sleep(0.2)
    dv.wait_busy_gone(driver, timeout=15)

    # проверка, что значение попало
    wrap = dv._closest_react_select(el)
    shown = (wrap.text if wrap is not None else "") or ""
    if mkb.lower() not in shown.lower():
        # вторая попытка жёстче
        try:
            el.click()
        except Exception:
            pass
        el.send_keys(Keys.CONTROL, "a")
        el.send_keys(Keys.BACKSPACE)
        el.send_keys(mkb, Keys.ENTER)
        sleep(0.5)
        dv.wait_busy_gone(driver)
        wrap = dv._closest_react_select(el)
        shown = (wrap.text if wrap is not None else "") or ""
    if mkb.lower() not in shown.lower():
        raise RuntimeError(f"mainDiagnosis не принял «{mkb}», сейчас «{shown[:100]}»")
    logger.info("mainDiagnosis → {}", shown.split("\n")[0][:100])


def save_existing_claim(driver: WebDriver, *, expected_mkb: str) -> tuple[bool, str]:
    """Save уже открытого амбулаторного талона (правка диагноза).

    На edit часто нет snackbar «сохранено» — считаем успехом, если нет ошибки
    и #mainDiagnosis по-прежнему содержит ожидаемый МКБ (или ушли со страницы).
    """
    mkb = dv.normalize_mkb(expected_mkb)
    snack_before = base.read_snackbar(driver)
    try:
        base._click_save_button(driver)
    except RuntimeError as exc:
        return False, str(exc)

    dv.wait_busy_gone(driver, timeout=60)
    sleep(0.4)

    # диалоги «данные пациента…» / предупреждения
    if dv.find_visible_dialogs(driver) or dv.warning_dialog_text(driver):
        ok_dlg, dlg_msg = dv.handle_post_save_prompts(driver, timeout=35)
        if not ok_dlg:
            low = (dlg_msg or "").lower()
            if "пересечен" in low or "дубликат" in low:
                return False, dlg_msg
            return False, dlg_msg or "диалог не подтверждён"

    snack = base.wait_snackbar(driver, timeout=6, ignore=snack_before)
    verdict, detail = base.classify_snackbar(snack)
    if verdict is False:
        return False, detail or snack or "ошибка по уведомлению"
    if verdict is True:
        return True, detail or snack

    # тихий Save: проверяем диагноз на форме / уход со страницы
    url = driver.current_url or ""
    if "/claim/ambulatory/" not in url or not driver.find_elements(By.ID, "mainDiagnosis"):
        return True, detail or snack or "ушли со страницы талона после Save"

    shown = ""
    try:
        el = driver.find_element(By.ID, "mainDiagnosis")
        wrap = dv._closest_react_select(el)
        shown = ((wrap.text if wrap is not None else "") or el.get_attribute("value") or "").strip()
    except Exception:
        shown = ""
    if mkb and mkb.lower() in shown.lower():
        return True, f"сохранено (тихо), диагноз={shown.split(chr(10))[0][:80]}"
    if snack:
        return False, f"неясное уведомление: {snack}"
    return False, f"Save без подтверждения, диагноз сейчас «{shown[:80]}»"


def process_row(driver: WebDriver, row: FixRow, *, index: int, dry_run: bool) -> RowResult:
    t0 = time()
    logger.info("=== [{}] talon={} ds1={} ===", index, row.talon, row.ds1)
    try:
        open_claim(driver, row.talon)
        set_main_diagnosis(driver, row.ds1)
        if dry_run:
            msg = f"dry-run: диагноз {row.ds1} выставлен, без Save"
            logger.success("[{}] {}", index, msg)
            return RowResult(index, row.talon, row.ds1, "dry_run", msg, round(time() - t0, 2))

        ok, msg = save_existing_claim(driver, expected_mkb=row.ds1)
        status = "ok" if ok else "fail"
        log = logger.success if ok else logger.error
        log("[{}] {} → {}", index, status.upper(), (msg or "")[:300])
        return RowResult(index, row.talon, row.ds1, status, msg, round(time() - t0, 2))
    except Exception as exc:
        msg = str(exc).split("Stacktrace:")[0].strip() or str(exc)
        logger.error("[{}] error: {}", index, msg[:400])
        try:
            driver.get(f"{OMS_BASE_URL.rstrip('/')}/claim/ambulatory")
            dv.wait_busy_gone(driver, timeout=30)
        except Exception:
            pass
        return RowResult(index, row.talon, row.ds1, "error", msg, round(time() - t0, 2))


def write_results(path: Path, results: list[RowResult]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = list(asdict(results[0]).keys()) if results else [
        "index", "talon", "ds1", "status", "message", "elapsed_sec"
    ]
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        for r in results:
            w.writerow(asdict(r))


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Исправление диагноза ds1 в амбулаторных талонах")
    p.add_argument("--file", type=Path, default=DEFAULT_FILE, help="Excel: Талон + ds1")
    p.add_argument(
        "--diagnosis",
        default="",
        help="МКБ по умолчанию, если колонки ds1 нет или ячейка пустая (напр. Z72.4)",
    )
    p.add_argument("--limit", type=int, default=0)
    p.add_argument("--offset", type=int, default=0)
    p.add_argument("--dry-run", action="store_true")
    p.add_argument("--keep-open", type=int, default=0)
    p.add_argument("--worker", default="")
    p.add_argument("--window-x", type=int, default=None)
    p.add_argument("--window-y", type=int, default=None)
    return p.parse_args()


def main() -> int:
    args = parse_args()
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)

    path = args.file
    if not path.is_file():
        alt = PROJECT_ROOT / path.name
        if alt.is_file():
            path = alt
        else:
            logger.error("Нет файла {}", args.file)
            return 1

    rows = load_rows(path, default_ds1=(args.diagnosis or "").strip())

    if args.offset:
        rows = rows[args.offset :]
    if args.limit and args.limit > 0:
        rows = rows[: args.limit]

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    worker_tag = (args.worker or "").strip() or f"o{args.offset}"
    result_path = RESULTS_DIR / f"fix_diagnosis_result_{stamp}_{worker_tag}.csv"
    logger.info(
        "[{}] строк с ds1={} | dry_run={} | offset={} | file={}",
        worker_tag,
        len(rows),
        args.dry_run,
        args.offset,
        path,
    )
    if not rows:
        logger.warning("Нечего обрабатывать (нет строк с заполненным ds1)")
        return 0

    win_kwargs: dict = {"implicit_wait": 0, "maximize": False}
    if args.window_x is not None and args.window_y is not None:
        win_kwargs["window_position"] = (args.window_x, args.window_y)
        win_kwargs["window_size"] = (1100, 850)
    else:
        win_kwargs["maximize"] = True

    results: list[RowResult] = []
    t_all = time()
    with GeckoBrowser(**win_kwargs) as session:
        driver = session.driver
        driver.implicitly_wait(0)
        login_oms(driver)
        sleep(0.8)
        dv.wait_busy_gone(driver, timeout=40)

        for i, row in enumerate(rows, start=1 + args.offset):
            results.append(process_row(driver, row, index=i, dry_run=args.dry_run))
            write_results(result_path, results)

        if args.keep_open > 0:
            logger.info("keep-open {}s", args.keep_open)
            sleep(args.keep_open)

    ok_n = sum(1 for r in results if r.status in ("ok", "dry_run"))
    fail_n = sum(1 for r in results if r.status not in ("ok", "dry_run"))
    logger.info(
        "——— [{}] ИТОГ ({:.0f}s) OK/DRY={}, FAIL={} → {} ———",
        worker_tag,
        time() - t_all,
        ok_n,
        fail_n,
        result_path,
    )
    return 0 if fail_n == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
