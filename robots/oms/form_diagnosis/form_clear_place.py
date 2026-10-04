#!/usr/bin/env python
"""
Убрать «Место прохождения диспансеризации» (#medicalExaminationPlace) с уже
сохранённых талонов ДВ4/ОПВ и сохранить.

Excel: колонки «Талон» (id) + «Цель» (ДВ4/ОПВ, информационно).
URL: {OMS_BASE}/claim/medicalExamination/{talon_id}

Пример:
  python form_diagnosis.py --limit 1 --dry-run --keep-open 20
  python form_diagnosis.py --file data/ubrat_mesto.xlsx
  python run_form_diagnosis.py --offset 0 --limit 50
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from time import sleep, time

from loguru import logger
from openpyxl import load_workbook
from selenium.common.exceptions import TimeoutException
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.remote.webdriver import WebDriver
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait

WORK_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = WORK_DIR.parents[2]  # BrowserAuto
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))
# переиспользуем save/dialogs из dv_opv
DV_DIR = PROJECT_ROOT / "robots" / "oms" / "dv_opv"
if str(DV_DIR) not in sys.path:
    sys.path.insert(0, str(DV_DIR))

from browser_auto.auth import login_oms
from browser_auto.config import OMS_BASE_URL, load_credentials
from browser_auto.driver import GeckoBrowser

import add_medical_exam as dv

load_credentials(WORK_DIR / "credentials.env")

DEFAULT_FILE = WORK_DIR / "data" / "ubrat_mesto.xlsx"
RESULTS_DIR = WORK_DIR / "data"
PLACE_INPUT_ID = "medicalExaminationPlace"


@dataclass
class TalonRef:
    talon_id: str
    exam_type: str = ""
    excel_row: int = 0


@dataclass
class RowResult:
    index: int
    talon_id: str
    exam_type: str
    status: str
    message: str
    had_place: str = ""
    elapsed_sec: float = 0.0


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]
    if text.lower() in {"nan", "none", "-"}:
        return ""
    return text


def claim_url(talon_id: str) -> str:
    base = (OMS_BASE_URL or "http://10.36.0.142:9000").rstrip("/")
    return f"{base}/claim/medicalExamination/{talon_id}"


def load_talons(path: Path) -> list[TalonRef]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []
    header = [_cell(h).lower().replace("ё", "е") for h in rows[0]]

    def col(*names: str) -> int | None:
        for n in names:
            if n.lower() in header:
                return header.index(n.lower())
        return None

    i_id = col("талон", "id", "claim", "номер", "talon")
    i_type = col("цель", "тип", "type", "exam_type")
    if i_id is None:
        raise ValueError(f"Нужна колонка «Талон». Сейчас: {rows[0]}")

    out: list[TalonRef] = []
    for n, raw in enumerate(rows[1:], start=2):
        if not raw or all(v is None or str(v).strip() == "" for v in raw):
            continue
        tid = re.sub(r"\D", "", _cell(raw[i_id]))
        if len(tid) < 5:
            continue
        typ = _cell(raw[i_type]) if i_type is not None else ""
        out.append(TalonRef(talon_id=tid, exam_type=typ, excel_row=n))
    return out


def _wait(driver: WebDriver, timeout: float = 20) -> WebDriverWait:
    return WebDriverWait(driver, timeout)


def place_display_text(driver: WebDriver) -> str:
    els = driver.find_elements(By.ID, PLACE_INPUT_ID)
    if not els:
        return ""
    el = els[0]
    wrap = dv._closest_react_select(el)
    chunks: list[str] = []
    if wrap is not None:
        for sel in (
            ".Select-value-label",
            ".Select-value",
            ".Selected-item",
            "[class*='singleValue']",
        ):
            for node in wrap.find_elements(By.CSS_SELECTOR, sel):
                t = (node.text or node.get_attribute("title") or "").strip()
                if t and t not in chunks:
                    chunks.append(t)
    val = (el.get_attribute("value") or "").strip()
    if val and val not in chunks:
        chunks.append(val)
    return " | ".join(chunks)


def clear_medical_examination_place(
    driver: WebDriver, *, _attempt: int = 0
) -> tuple[bool, str]:
    """Очистить #medicalExaminationPlace. Только UI clear — DOM не ломаем."""
    try:
        _wait(driver, 40).until(
            lambda d: bool(
                d.find_elements(By.ID, PLACE_INPUT_ID)
                or d.find_elements(By.ID, "enp")
                or d.find_elements(By.ID, "begin-date")
            )
        )
    except TimeoutException:
        u = ""
        try:
            u = driver.current_url
        except Exception:
            pass
        logger.error("Форма не загрузилась. url={}", u)
        return False, f"форма талона не загрузилась. url={u}"

    if "oauth" in (driver.current_url or "").lower():
        return False, f"сессия потеряна (oauth): {driver.current_url}"

    els = driver.find_elements(By.ID, PLACE_INPUT_ID)
    if not els:
        return True, "поля medicalExaminationPlace нет на форме — ок"

    el = els[0]
    before = place_display_text(driver)
    if not (before or "").strip():
        logger.info("medicalExaminationPlace: уже пусто")
        return True, "было пусто"

    wrap = dv._closest_react_select(el)
    cleared = False

    # 1) крестик очистки react-select
    if wrap is not None:
        # навести — иногда clear появляется только на hover
        try:
            driver.execute_script(
                "arguments[0].scrollIntoView({block:'center'});"
                "arguments[0].dispatchEvent(new MouseEvent('mouseover',{bubbles:true}));",
                wrap,
            )
            sleep(0.15)
        except Exception:
            pass
        for sel in (
            ".Select-clear-zone",
            ".Select-clear",
            ".Select-value-icon",
            "[class*='clearIndicator']",
            "[aria-label='Clear']",
        ):
            icons = wrap.find_elements(By.CSS_SELECTOR, sel)
            if not icons:
                continue
            try:
                driver.execute_script("arguments[0].click();", icons[0])
                cleared = True
                logger.info("place: клик {}", sel)
                sleep(0.35)
                break
            except Exception:
                continue

    # 2) если ещё не пусто — Backspace без native click по input
    if (place_display_text(driver) or "").strip():
        try:
            driver.execute_script("arguments[0].scrollIntoView({block:'center'});", el)
            if wrap is not None:
                # ещё раз clear / клик по выбранному значению
                for sel in (
                    ".Select-clear-zone",
                    ".Select-clear",
                    ".Selected-item",
                    ".Select-value",
                    ".Select-control",
                ):
                    nodes = wrap.find_elements(By.CSS_SELECTOR, sel)
                    if not nodes:
                        continue
                    try:
                        driver.execute_script("arguments[0].click();", nodes[0])
                        sleep(0.12)
                    except Exception:
                        continue
            driver.execute_script("arguments[0].focus();", el)
            for _ in range(12):
                el.send_keys(Keys.BACKSPACE)
                el.send_keys(Keys.DELETE)
            el.send_keys(Keys.ESCAPE)
            sleep(0.3)
        except Exception as exc:
            return False, f"не удалось очистить place: {exc}"

    # закрыть меню
    try:
        el.send_keys(Keys.ESCAPE)
        driver.execute_script("document.activeElement && document.activeElement.blur();")
        driver.find_element(By.TAG_NAME, "body").send_keys(Keys.ESCAPE)
    except Exception:
        pass
    sleep(0.3)

    after = place_display_text(driver)
    # placeholder «Выберите...» не считаем значением
    after_clean = after
    for noise in ("выберите", "select", "placeholder"):
        if noise in (after or "").lower():
            after_clean = ""
            break

    if after_clean and after_clean == before:
        return False, f"place не очистилось (было={before!r}, стало={after!r})"

    # Save должен остаться в DOM
    if not driver.find_elements(By.ID, "id-save"):
        if _attempt >= 1:
            return False, "после очистки пропал #id-save"
        logger.warning("После очистки пропал #id-save — обновляем страницу и чистим мягко ещё раз")
        driver.refresh()
        dv.wait_busy_gone(driver, timeout=40)
        sleep(1.0)
        return clear_medical_examination_place(driver, _attempt=_attempt + 1)

    msg = f"очищено (было={before!r})"
    logger.info("medicalExaminationPlace: {}", msg)
    return True, msg


def save_claim_force(driver: WebDriver) -> tuple[bool, str]:
    """Save для уже открытого талона: JS-клик #id-save (без F2)."""
    dv.dismiss_overlays(driver)
    try:
        driver.find_element(By.TAG_NAME, "body").send_keys(Keys.ESCAPE)
    except Exception:
        pass
    sleep(0.25)

    def _js_click_id_save() -> bool:
        els = driver.find_elements(By.ID, "id-save")
        if not els:
            return False
        try:
            driver.execute_script(
                "arguments[0].scrollIntoView({block:'center'});"
                "arguments[0].removeAttribute('disabled');"
                "arguments[0].disabled=false;"
                "arguments[0].click();",
                els[0],
            )
            logger.info("Save: JS-клик #id-save")
            return True
        except Exception as exc:
            logger.warning("JS id-save fail: {}", exc)
            return False

    clicked = _js_click_id_save()
    if not clicked:
        for tab in ("main-tab", "services-tab"):
            if not driver.find_elements(By.ID, tab):
                continue
            try:
                dv.click_id(driver, tab)
                dv.wait_busy_gone(driver, timeout=15)
                sleep(0.35)
            except Exception:
                continue
            if _js_click_id_save():
                clicked = True
                break

    if not clicked:
        for btn in driver.find_elements(
            By.XPATH, "//button[contains(.,'Сохранить') or contains(.,'СОХРАНИТЬ')]"
        ):
            try:
                lab = (btn.text or "").lower()
                if any(x in lab for x in ("пациент", "предупред", "вернуться", "карт")):
                    continue
                driver.execute_script(
                    "arguments[0].scrollIntoView({block:'center'}); arguments[0].click();",
                    btn,
                )
                clicked = True
                logger.info("Save: кнопка «{}»", (btn.text or "")[:40])
                break
            except Exception:
                continue

    if not clicked:
        return False, "кнопка Save не найдена"

    dv.wait_busy_gone(driver, timeout=60)
    sleep(0.4)
    ok, msg = dv.handle_post_save_prompts(driver, timeout=35)
    if ok:
        return True, msg
    if dv.find_visible_dialogs(driver) or dv.warning_dialog_text(driver):
        ok2, w = dv.confirm_warning_dialog(driver)
        if ok2:
            return True, w or msg
        # ещё раз — «Сохранить с предупреждением»
        try:
            bb = " | ".join(
                (b.text or "").lower()
                for b in driver.find_elements(By.XPATH, "//button")
                if b.is_displayed()
            )
        except Exception:
            bb = ""
        if "предупрежден" in bb:
            ok3, w3 = dv.confirm_warning_dialog(driver)
            if ok3:
                return True, w3 or msg
    snack = dv.snackbar_text(driver)
    if snack and dv._is_save_ok_message(snack):
        return True, snack
    if "medicalExamination/" not in (driver.current_url or "") or not driver.find_elements(
        By.ID, "enp"
    ):
        return True, msg or snack or "ушли со страницы талона после Save"
    return False, msg or snack or "Save без подтверждения"


def open_claim(driver: WebDriver, talon_id: str) -> None:
    url = claim_url(talon_id)
    logger.info("Открываем {}", url)
    # если сессия умерла — перелогин
    if "oauth" in (driver.current_url or "").lower():
        logger.warning("Перед открытием талона — снова oauth, логинимся")
        login_oms(driver)
    driver.get(url)
    dv.wait_busy_gone(driver, timeout=60)
    # SPA: дождаться контента
    try:
        _wait(driver, 40).until(
            lambda d: bool(
                d.find_elements(By.ID, PLACE_INPUT_ID)
                or d.find_elements(By.ID, "enp")
                or d.find_elements(By.ID, "id-save")
            )
            or "oauth" in (d.current_url or "").lower()
        )
    except TimeoutException:
        pass
    if "oauth" in (driver.current_url or "").lower():
        logger.warning("После get талона — oauth, логин + повтор")
        login_oms(driver)
        driver.get(url)
        dv.wait_busy_gone(driver, timeout=60)
        sleep(1.0)
    dv.dismiss_overlays(driver)
    sleep(0.3)


def process_row(
    driver: WebDriver,
    row: TalonRef,
    *,
    index: int,
    dry_run: bool,
) -> RowResult:
    t0 = time()
    logger.info("=== [{}] talon={} type={} ===", index, row.talon_id, row.exam_type or "—")
    try:
        open_claim(driver, row.talon_id)
        had = place_display_text(driver)
        ok_clear, clear_msg = clear_medical_examination_place(driver)
        if not ok_clear:
            logger.error("[{}] clear fail: {}", index, clear_msg)
            return RowResult(
                index,
                row.talon_id,
                row.exam_type,
                "error",
                clear_msg,
                had,
                round(time() - t0, 2),
            )

        if dry_run:
            msg = f"dry-run: {clear_msg}, save пропущен"
            logger.success("[{}] {}", index, msg)
            return RowResult(
                index, row.talon_id, row.exam_type, "dry_run", msg, had, round(time() - t0, 2)
            )

        ok, msg = save_claim_force(driver)
        status = "ok" if ok else "fail"
        log = logger.success if ok else logger.error
        full = f"{clear_msg} | {msg}"
        log("[{}] {} → {}", index, status.upper(), full)
        return RowResult(
            index, row.talon_id, row.exam_type, status, full, had, round(time() - t0, 2)
        )
    except Exception as exc:
        msg = str(exc).split("Stacktrace:")[0].strip() or str(exc)
        logger.error("[{}] error: {}", index, msg)
        return RowResult(
            index, row.talon_id, row.exam_type, "error", msg, "", round(time() - t0, 2)
        )


def write_results(path: Path, results: list[RowResult]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = list(asdict(results[0]).keys()) if results else [
        "index", "talon_id", "exam_type", "status", "message", "had_place", "elapsed_sec"
    ]
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        for r in results:
            w.writerow(asdict(r))


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(
        description="Убрать medicalExaminationPlace с талонов и сохранить"
    )
    p.add_argument("--file", type=Path, default=DEFAULT_FILE)
    p.add_argument("--offset", type=int, default=0)
    p.add_argument("--limit", type=int, default=0, help="0 = все")
    p.add_argument("--dry-run", action="store_true")
    p.add_argument("--keep-open", type=int, default=0, help="секунд держать браузер в конце")
    return p.parse_args()


def main() -> int:
    args = parse_args()
    path = args.file if args.file.is_file() else DEFAULT_FILE
    if not path.is_file():
        logger.error("Нет файла: {}", path)
        return 2

    rows = load_talons(path)
    if args.offset:
        rows = rows[args.offset :]
    if args.limit and args.limit > 0:
        rows = rows[: args.limit]
    if not rows:
        logger.error("Нет строк для обработки")
        return 2

    logger.info(
        "Строк: {} | dry_run={} | file={}",
        len(rows),
        args.dry_run,
        path.name,
    )

    results: list[RowResult] = []
    with GeckoBrowser(implicit_wait=0) as session:
        driver = session.driver
        for attempt in range(1, 4):
            try:
                login_oms(driver)
                logger.info("OMS login OK (attempt {})", attempt)
                break
            except Exception as exc:
                logger.warning("login fail {}: {}", attempt, exc)
                if attempt == 3:
                    raise
                sleep(2)

        for i, row in enumerate(rows, start=1 + args.offset):
            results.append(process_row(driver, row, index=i, dry_run=args.dry_run))

        if args.keep_open > 0:
            logger.info("keep-open {}s…", args.keep_open)
            sleep(args.keep_open)

    out = RESULTS_DIR / f"clear_place_result_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    write_results(out, results)
    ok_n = sum(1 for r in results if r.status in ("ok", "dry_run"))
    fail_n = len(results) - ok_n
    logger.info("Готово: ok/dry={} fail/error={} → {}", ok_n, fail_n, out)
    return 0 if fail_n == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
