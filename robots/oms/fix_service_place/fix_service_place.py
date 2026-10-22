#!/usr/bin/env python
"""
Замена «Место обслуживания» (#service-place) в амбулаторных талонах Web.ОМС.

Excel (колонки):
  Талон | Место обслуживания
  39927692 | 9

Опционально: «было» — проверка текущего значения перед заменой.
CLI: --place 9 — общее место для всего файла (если колонки нет).

URL: {OMS}/claim/ambulatory/{Талон}

Примеры:
  python run_fix_service_place.py --limit 1 --dry-run --keep-open 20
  python run_fix_service_place.py --file "robots/oms/Место обсл.xlsx"
  python run_fix_service_place.py --file talons.xlsx --place 9
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from time import sleep, time

from loguru import logger
from openpyxl import Workbook, load_workbook
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.remote.webdriver import WebDriver
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait

WORK_DIR = Path(__file__).resolve().parent
OMS_DIR = WORK_DIR.parent
PROJECT_ROOT = WORK_DIR.parents[2]
for p in (PROJECT_ROOT, WORK_DIR, OMS_DIR / "dv_opv", OMS_DIR / "cel_307"):
    if str(p) not in sys.path:
        sys.path.insert(0, str(p))

from browser_auto.config import OMS_BASE_URL, load_credentials
from browser_auto.auth import login_oms
from browser_auto.driver import GeckoBrowser

import add_medical_exam as dv  # noqa: E402
import add_307 as base  # noqa: E402

load_credentials(WORK_DIR / "credentials.env", PROJECT_ROOT / ".env")

RESULTS_DIR = WORK_DIR / "data"
DEFAULT_FILE = RESULTS_DIR / "fix_service_place.xlsx"
PLACE_INPUT_ID = "service-place"


@dataclass
class FixRow:
    talon: str
    place: str
    place_from: str = ""


@dataclass
class RowResult:
    index: int
    talon: str
    place: str
    place_before: str
    status: str
    message: str
    elapsed_sec: float


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    text = str(v).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]
    if text.lower() in {"nan", "none", "-"}:
        return ""
    return text


def _norm_header(h: object) -> str:
    return _cell(h).lower().replace("ё", "е")


def load_rows(
    path: Path,
    *,
    default_place: str = "",
    skip_done: bool = True,
) -> list[FixRow]:
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    headers = [
        _norm_header(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))
    ]
    col = {h: i for i, h in enumerate(headers) if h}

    def idx(*names: str) -> int | None:
        for n in names:
            if n in col:
                return col[n]
        return None

    i_talon = idx("талон", "номер", "id", "claim", "n", "web-oms", "web_oms")
    i_place = idx(
        "место обслуживания",
        "место",
        "место обращения",
        "place",
        "service_place",
        "service-place",
        "стало",
        "to",
        "new",
    )
    i_from = idx("было", "from", "old", "place_from")
    i_status = idx("status", "статус")

    if i_talon is None:
        raise SystemExit(f"Нужна колонка «Талон», сейчас: {headers}")

    def_place = _cell(default_place)
    if i_place is None and not def_place:
        raise SystemExit(
            f"Нужна колонка «Место обслуживания» или флаг --place. Сейчас: {headers}"
        )

    out: list[FixRow] = []
    for raw in ws.iter_rows(min_row=2, values_only=True):
        if not raw or not any(v is not None and str(v).strip() != "" for v in raw):
            continue
        talon = _cell(raw[i_talon])
        if not talon:
            continue
        if talon.endswith(".0"):
            talon = talon[:-2]
        place = def_place or (
            _cell(raw[i_place]) if i_place is not None else ""
        )
        if not place:
            continue
        place_from = _cell(raw[i_from]) if i_from is not None else ""
        if skip_done and i_status is not None:
            st = _cell(raw[i_status]).lower()
            if st in ("ok", "done", "готово", "обработан", "обработано"):
                continue
        out.append(FixRow(talon=talon, place=place, place_from=place_from))
    return out


def make_template(path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "место"
    ws.append(["Талон", "Место обслуживания"])
    ws.append(["39927692", "9"])
    wb.save(path)
    return path


def open_claim(driver: WebDriver, talon: str) -> str:
    url = f"{OMS_BASE_URL.rstrip('/')}/claim/ambulatory/{talon}"
    logger.info("Открываем {}", url)
    driver.get(url)
    dv.wait_busy_gone(driver, timeout=45)
    sleep(0.5)
    WebDriverWait(driver, 25).until(
        EC.any_of(
            EC.presence_of_element_located((By.ID, PLACE_INPUT_ID)),
            EC.presence_of_element_located((By.ID, "add-service")),
            EC.presence_of_element_located((By.ID, "begin-date")),
            EC.presence_of_element_located((By.ID, "mainDiagnosis")),
        )
    )
    if "oauth" in (driver.current_url or "").lower():
        raise RuntimeError(f"сессия потеряна (oauth): {driver.current_url}")
    if not driver.find_elements(By.ID, PLACE_INPUT_ID):
        raise RuntimeError(f"На форме нет #{PLACE_INPUT_ID}")
    return url


def place_display_text(driver: WebDriver) -> str:
    els = driver.find_elements(By.ID, PLACE_INPUT_ID)
    if not els:
        return ""
    el = els[0]
    wrap = dv._closest_react_select(el)
    chunks: list[str] = []
    if wrap is not None:
        for sel in (
            ".Select-value-label",
            ".Select-value",
            ".Selected-item",
            "[class*='singleValue']",
        ):
            for node in wrap.find_elements(By.CSS_SELECTOR, sel):
                t = (node.text or node.get_attribute("title") or "").strip()
                if t and t not in chunks:
                    chunks.append(t)
    val = (el.get_attribute("value") or "").strip()
    if val and val not in chunks:
        chunks.append(val)
    return " | ".join(chunks)


def _place_matches(current: str, expected: str) -> bool:
    """Считаем совпадением, если код места уже выбран (начало строки / отдельный токен)."""
    cur = (current or "").strip().lower()
    exp = (expected or "").strip().lower()
    if not exp:
        return False
    if not cur:
        return False
    if cur == exp:
        return True
    # «9 — Поликлиника» / «9. ...»
    if cur.startswith(exp) and (
        len(cur) == len(exp) or cur[len(exp)] in " .-\t|—"
    ):
        return True
    # код в начале первого токена
    first = re.split(r"[\s|—\-]+", cur, maxsplit=1)[0]
    return first == exp


def set_service_place(driver: WebDriver, place: str) -> tuple[str, str]:
    """Поставить #service-place. Возвращает (было, стало_после)."""
    before = place_display_text(driver)
    if _place_matches(before, place):
        logger.info("service-place уже = {!r} (цель {})", before, place)
        return before, before

    ok = dv.try_input_enter_id(driver, PLACE_INPUT_ID, place)
    if not ok:
        els = driver.find_elements(By.ID, PLACE_INPUT_ID)
        if not els:
            raise RuntimeError(f"#{PLACE_INPUT_ID} не найден")
        try:
            dv.fill_react_select_input(driver, els[0], place)
        except Exception as exc:
            raise RuntimeError(f"не удалось заполнить service-place={place}: {exc}") from exc

    sleep(0.35)
    dv.wait_busy_gone(driver, timeout=15)
    after = place_display_text(driver)
    if not _place_matches(after, place):
        # повтор через fill_react_select
        els = driver.find_elements(By.ID, PLACE_INPUT_ID)
        if els:
            try:
                dv.fill_react_select_input(driver, els[0], place)
                sleep(0.35)
                dv.wait_busy_gone(driver, timeout=15)
                after = place_display_text(driver)
            except Exception:
                pass
    if not _place_matches(after, place):
        raise RuntimeError(
            f"service-place не сменилось: было={before!r}, цель={place!r}, стало={after!r}"
        )
    logger.info("service-place: {!r} → {!r}", before, after)
    return before, after


def _left_claim_form(driver: WebDriver, talon: str) -> bool:
    url = (driver.current_url or "").lower()
    if f"/ambulatory/{talon}" not in url:
        return True
    if not driver.find_elements(By.ID, "add-service") and not driver.find_elements(
        By.ID, "mainDiagnosis"
    ):
        return True
    return False


def save_existing_claim(driver: WebDriver, talon: str) -> tuple[bool, str]:
    """Save уже открытого талона (как в fix_service)."""
    url_before = driver.current_url or ""
    snack_before = base.read_snackbar(driver)

    try:
        base._click_save_button(driver)
    except RuntimeError as exc:
        return False, str(exc)

    dv.wait_busy_gone(driver, timeout=40)
    sleep(0.35)

    snack = base.wait_snackbar(driver, timeout=4, ignore=snack_before)
    verdict, detail = base.classify_snackbar(snack)

    dlg = dv.warning_dialog_text(driver)
    if (not dlg) and detail.startswith("[dialog]"):
        dlg = detail[len("[dialog]") :].strip()

    if dlg or dv.find_visible_dialogs(driver):
        low = (dlg or "").lower()
        if "дубликат" in low or dv._is_intersection_warning(dlg or ""):
            dv.dismiss_warning_dialog(driver)
            return False, dlg or "блокирующее предупреждение"
        if "результат поиска в црп" in low or "отсутствует пациент" in low:
            dv.dismiss_warning_dialog(driver)
            return False, dlg
        ok, wmsg = dv.handle_post_save_prompts(driver, timeout=20)
        if not ok:
            return False, wmsg or dlg or "диалог не подтверждён"
        dv.wait_busy_gone(driver, timeout=20)
        sleep(0.3)
        if _left_claim_form(driver, talon):
            return True, wmsg or "сохранено (уход в журнал)"
        snack2 = base.wait_snackbar(driver, timeout=3, ignore=snack)
        v2, d2 = base.classify_snackbar(snack2)
        if v2 is False:
            return False, d2 or wmsg
        return True, d2 or wmsg or "сохранено"

    if verdict is True:
        return True, detail

    if verdict is False:
        if _left_claim_form(driver, talon) and (detail or "").strip().lower() in (
            "error",
            "",
        ):
            logger.info("Snackbar Error, но ушли с формы талона {} — считаем OK", talon)
            return True, "сохранено (журнал после Error)"
        return False, detail or snack or "ошибка по уведомлению"

    if _left_claim_form(driver, talon) or (driver.current_url or "") != url_before:
        logger.info("Нет snackbar, уход с формы {} → OK (журнал)", talon)
        return True, "сохранено (журнал, без snackbar)"

    for _ in range(8):
        sleep(0.35)
        if _left_claim_form(driver, talon):
            return True, "сохранено (журнал)"
        sn = base.read_snackbar(driver)
        v, d = base.classify_snackbar(sn)
        if v is True:
            return True, d
        if v is False and (d or "").strip().lower() != "error":
            return False, d

    return False, "нет подтверждения Save — талон остался на форме"


def process_row(
    driver: WebDriver,
    row: FixRow,
    *,
    index: int,
    dry_run: bool,
) -> RowResult:
    t0 = time()
    logger.info("=== [{}] talon={} place={} ===", index, row.talon, row.place)
    try:
        open_claim(driver, row.talon)
        before = place_display_text(driver)

        if row.place_from and not _place_matches(before, row.place_from):
            msg = f"пропуск: ожидалось было={row.place_from!r}, сейчас={before!r}"
            logger.warning("[{}] {}", index, msg)
            return RowResult(
                index,
                row.talon,
                row.place,
                before,
                "skip",
                msg,
                round(time() - t0, 2),
            )

        if _place_matches(before, row.place):
            msg = f"уже установлено: {before!r}"
            logger.success("[{}] skip — {}", index, msg)
            return RowResult(
                index,
                row.talon,
                row.place,
                before,
                "skip",
                msg,
                round(time() - t0, 2),
            )

        before, after = set_service_place(driver, row.place)
        detail = f"{before!r} → {after!r}"

        if dry_run:
            msg = f"dry-run: {detail}, без Save"
            logger.success("[{}] {}", index, msg)
            return RowResult(
                index,
                row.talon,
                row.place,
                before,
                "dry_run",
                msg,
                round(time() - t0, 2),
            )

        ok, msg = save_existing_claim(driver, row.talon)
        status = "ok" if ok else "fail"
        full = f"{detail} | {msg}"
        log = logger.success if ok else logger.error
        log("[{}] {} → {}", index, status.upper(), full[:300])
        return RowResult(
            index,
            row.talon,
            row.place,
            before,
            status,
            full,
            round(time() - t0, 2),
        )
    except Exception as exc:
        msg = str(exc).split("Stacktrace:")[0].strip() or str(exc)
        logger.error("[{}] error: {}", index, msg[:400])
        return RowResult(
            index,
            row.talon,
            row.place,
            "",
            "error",
            msg,
            round(time() - t0, 2),
        )


def write_results(path: Path, results: list[RowResult]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = list(asdict(results[0]).keys()) if results else [
        "index",
        "talon",
        "place",
        "place_before",
        "status",
        "message",
        "elapsed_sec",
    ]
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        for r in results:
            w.writerow(asdict(r))


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(
        description="Замена «Место обслуживания» (#service-place) в амбулаторных талонах"
    )
    p.add_argument(
        "--file",
        type=Path,
        default=DEFAULT_FILE,
        help="Excel: Талон + Место обслуживания",
    )
    p.add_argument(
        "--place",
        default="",
        help="Место обслуживания для всего файла (если нет колонки в Excel)",
    )
    p.add_argument("--make-template", action="store_true", help="Создать шаблон Excel и выйти")
    p.add_argument("--limit", type=int, default=0)
    p.add_argument("--offset", type=int, default=0)
    p.add_argument("--dry-run", action="store_true")
    p.add_argument("--keep-open", type=int, default=0)
    p.add_argument("--worker", default="")
    p.add_argument("--window-x", type=int, default=None)
    p.add_argument("--window-y", type=int, default=None)
    return p.parse_args()


def main() -> int:
    args = parse_args()
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)

    if args.make_template:
        path = make_template(args.file if args.file != DEFAULT_FILE else DEFAULT_FILE)
        logger.info("Шаблон: {}", path)
        return 0

    path = args.file
    if not path.is_file():
        # поиск рядом: robots/oms/*.xlsx с кириллическим именем
        candidates = [
            PROJECT_ROOT / path.name,
            OMS_DIR / path.name,
            OMS_DIR / "Место обсл.xlsx",
            RESULTS_DIR / "fix_service_place.xlsx",
        ]
        for alt in candidates:
            if alt.is_file():
                path = alt
                break
        else:
            logger.error("Нет файла {}", args.file)
            return 1

    rows = load_rows(path, default_place=args.place)
    if args.offset:
        rows = rows[args.offset :]
    if args.limit and args.limit > 0:
        rows = rows[: args.limit]

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    worker_tag = (args.worker or "").strip() or f"o{args.offset}"
    result_path = RESULTS_DIR / f"fix_service_place_result_{stamp}_{worker_tag}.csv"
    logger.info(
        "[{}] строк={} | dry_run={} | place={} | offset={} | file={}",
        worker_tag,
        len(rows),
        args.dry_run,
        args.place or "(из Excel)",
        args.offset,
        path,
    )
    if not rows:
        logger.warning("Нечего обрабатывать")
        return 0

    win_kwargs: dict = {"implicit_wait": 0, "maximize": False}
    if args.window_x is not None and args.window_y is not None:
        win_kwargs["window_position"] = (args.window_x, args.window_y)
        win_kwargs["window_size"] = (1100, 850)
    else:
        win_kwargs["maximize"] = True

    results: list[RowResult] = []
    t_all = time()
    with GeckoBrowser(**win_kwargs) as session:
        driver = session.driver
        driver.implicitly_wait(0)
        login_oms(driver)
        sleep(0.8)
        dv.wait_busy_gone(driver, timeout=40)

        for i, row in enumerate(rows, start=1 + args.offset):
            results.append(
                process_row(
                    driver,
                    row,
                    index=i,
                    dry_run=args.dry_run,
                )
            )
            write_results(result_path, results)

        if args.keep_open > 0:
            logger.info("keep-open {}s", args.keep_open)
            sleep(args.keep_open)

    ok_n = sum(1 for r in results if r.status in ("ok", "dry_run", "skip"))
    fail_n = sum(1 for r in results if r.status not in ("ok", "dry_run", "skip"))
    logger.info(
        "——— [{}] ИТОГ ({:.0f}s) OK/DRY/SKIP={}, FAIL={} → {} ———",
        worker_tag,
        time() - t_all,
        ok_n,
        fail_n,
        result_path,
    )
    return 0 if fail_n == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
