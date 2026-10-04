# -*- coding: utf-8 -*-
"""
Сборка Excel для цели 541 (КТ) из talons_541_kt.xlsx.

Группировка: ЕНП + календарная дата исследования.
  • несколько услуг → лист «услуги» (несколько строк на один «Талон»)
  • несколько МКБ → Диагноз (основной) + Диагноз 2 (сопутствующий);
    остальные — в «Диагнозы_все»
  • даты → дд-мм-гг; явка (посещений в МО) = 1
  • кол-во услуги = число одинаковых кодов в группе

Примеры:
  python template_541.py
  python template_541.py --source data/fixed/talons_541_kt.xlsx
"""
from __future__ import annotations

import argparse
import re
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

WORK_DIR = Path(__file__).resolve().parent
DATA_DIR = WORK_DIR / "data"
DEFAULT_SOURCE = DATA_DIR / "fixed" / "talons_541_kt.xlsx"

DEFAULTS_541 = {
    "цель": "541",
    "случай": "Первичный",
    "случай_2": "Законченный",
    "ДИСП. НАБЛ": "1",
    "тип": "3",
    "результат": "301",
    "исход": "304",
    "характер": "3",
    "посещений в МО": "1",
}

TALON_COLUMNS = [
    "Талон",
    "начало",
    "окончание",
    "ЕНП",
    "Пациент",
    "врач",
    "Корпус",
    "цель",
    "случай",
    "случай_2",
    "посещений в МО",
    "Диагноз",
    "Диагноз 2",
    "Диагнозы_все",
    "ДИСП. НАБЛ",
    "тип",
    "результат",
    "исход",
    "характер",
    "услуг",
    "источник_строк",
]

SERVICE_COLUMNS = [
    "Талон",
    "Код услуги",
    "Название",
    "Дата начала",
    "Дата окончания",
    "Кол-во",
    "Врач",
]

HDR_FILL = PatternFill("solid", fgColor="1F4E79")
HDR_FONT = Font(bold=True, color="FFFFFF")


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def to_oms_date(value: object) -> str:
    """Любая дата → дд-мм-гг."""
    if value is None or value == "":
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d-%m-%y")
    if isinstance(value, date):
        return value.strftime("%d-%m-%y")
    text = str(value).strip()
    if not text or text.lower() in {"nan", "none", "-", "—"}:
        return ""
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", text)
    if m:
        y, mo, d = m.groups()
        return f"{int(d):02d}-{int(mo):02d}-{y[-2:]}"
    m = re.match(r"^(\d{1,2})[./](\d{1,2})[./](\d{2,4})", text)
    if m:
        d, mo, y = m.groups()
        if len(y) == 4:
            y = y[-2:]
        return f"{int(d):02d}-{int(mo):02d}-{y}"
    m = re.match(r"^(\d{1,2})-(\d{1,2})-(\d{2,4})$", text)
    if m:
        d, mo, y = m.groups()
        if len(y) == 4:
            y = y[-2:]
        return f"{int(d):02d}-{int(mo):02d}-{y}"
    return text


def parse_dt(value: object) -> datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    text = str(value).strip().replace(",", ".")
    for fmt in (
        "%Y-%m-%d %H:%M:%S.%f",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%dT%H:%M:%S.%f",
        "%Y-%m-%dT%H:%M:%S",
        "%Y-%m-%d",
        "%d.%m.%Y",
        "%d-%m-%Y",
        "%d-%m-%y",
    ):
        try:
            return datetime.strptime(text[:26], fmt)
        except ValueError:
            continue
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", text)
    if m:
        return datetime(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    return None


def _style_header(ws, n: int) -> None:
    for c in range(1, n + 1):
        cell = ws.cell(1, c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def load_doctor_codes(wb) -> dict[str, str]:
    """ФИО → код; берём код_врача, иначе примечание (если похоже на код)."""
    if "справочник_врачей" not in wb.sheetnames:
        return {}
    ws = wb["справочник_врачей"]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {}
    header = [_cell(h).lower() for h in rows[0]]
    i_fio = next((i for i, h in enumerate(header) if "фио" in h or h == "врач"), 0)
    i_code = next((i for i, h in enumerate(header) if "код" in h and "врач" in h), None)
    i_note = next((i for i, h in enumerate(header) if "примечан" in h), None)
    out: dict[str, str] = {}
    for r in rows[1:]:
        fio = _cell(r[i_fio] if i_fio < len(r) else "")
        if not fio:
            continue
        code = ""
        if i_code is not None and i_code < len(r):
            code = _cell(r[i_code])
        if not code and i_note is not None and i_note < len(r):
            note = _cell(r[i_note])
            if re.fullmatch(r"\d{5,}", note):
                code = note
        if code:
            out[fio] = code
    return out


def load_source_rows(path: Path) -> tuple[list[dict], dict[str, str]]:
    wb = load_workbook(path, data_only=True)
    doc_codes = load_doctor_codes(wb)
    ws = wb["талоны"] if "талоны" in wb.sheetnames else wb.active
    rows_raw = list(ws.iter_rows(values_only=True))
    wb.close()
    header = [_cell(h) for h in rows_raw[0]]
    idx = {h.lower(): i for i, h in enumerate(header)}

    def col(*names: str) -> int | None:
        for n in names:
            if n.lower() in idx:
                return idx[n.lower()]
        return None

    i_enp = col("енп")
    i_fio = col("пациент", "ф.и.о", "фио")
    i_created = col("дата создания", "дата")
    i_research = col("исследование")
    i_svc = col("услуга", "код услуги")
    i_svc_name = col("услуга_название", "название")
    i_doc = col("врач")
    i_doc_code = col("код_врача", "код врача")
    i_diag = col("диагноз", "код")
    if i_enp is None or i_svc is None or i_created is None:
        raise SystemExit(f"Нужны колонки ЕНП / услуга / Дата создания. Есть: {header}")

    rows: list[dict] = []
    for r in rows_raw[1:]:
        if not r or not any(v is not None and str(v).strip() for v in r):
            continue
        enp = _cell(r[i_enp])
        svc = _cell(r[i_svc])
        if not enp or not svc:
            continue
        fio_doc = _cell(r[i_doc]) if i_doc is not None else ""
        code_doc = _cell(r[i_doc_code]) if i_doc_code is not None else ""
        if not code_doc:
            code_doc = doc_codes.get(fio_doc, "")
        dt = parse_dt(r[i_created])
        rows.append(
            {
                "enp": enp,
                "patient": _cell(r[i_fio]) if i_fio is not None else "",
                "created": dt,
                "day": (dt.date() if dt else None),
                "research": _cell(r[i_research]) if i_research is not None else "",
                "service": svc,
                "service_name": _cell(r[i_svc_name]) if i_svc_name is not None else "",
                "doctor_fio": fio_doc,
                "doctor": code_doc,
                "diagnosis": _cell(r[i_diag]) if i_diag is not None else "",
            }
        )
    return rows, doc_codes


def build_talons(
    source: Path,
    *,
    out_talons: Path,
    out_services: Path,
    out_bundle: Path,
    default_building: str = "",
) -> tuple[int, int]:
    raw, doc_codes = load_source_rows(source)
    if not raw:
        raise SystemExit(f"Пустой источник: {source}")

    groups: dict[tuple[str, date], list[dict]] = defaultdict(list)
    skipped_no_day = 0
    for r in raw:
        if r["day"] is None:
            skipped_no_day += 1
            continue
        groups[(r["enp"], r["day"])].append(r)

    # сортировка групп по дате / енп
    keys = sorted(groups.keys(), key=lambda k: (k[1], k[0]))

    talon_rows: list[list] = [TALON_COLUMNS]
    service_rows: list[list] = [SERVICE_COLUMNS]

    for n, key in enumerate(keys, start=1):
        items = sorted(
            groups[key],
            key=lambda x: x["created"] or datetime.min,
        )
        enp, day = key
        oms_date = day.strftime("%d-%m-%y")

        # диагнозы в порядке появления, unique
        diags: list[str] = []
        for it in items:
            d = (it["diagnosis"] or "").strip()
            if d and d not in diags:
                diags.append(d)
        main_diag = diags[0] if diags else ""
        diag2 = diags[1] if len(diags) > 1 else ""
        diags_all = "; ".join(diags)

        # врач — самый частый код; иначе из справочника по ФИО
        doc_counter: Counter[str] = Counter()
        fio_counter: Counter[str] = Counter()
        for it in items:
            if it["doctor"]:
                doc_counter[it["doctor"]] += 1
            if it["doctor_fio"]:
                fio_counter[it["doctor_fio"]] += 1
        doctor = doc_counter.most_common(1)[0][0] if doc_counter else ""
        if not doctor and fio_counter:
            top_fio = fio_counter.most_common(1)[0][0]
            doctor = doc_codes.get(top_fio, "")

        patient = next((it["patient"] for it in items if it["patient"]), "")

        # услуги: unique code, amount = count, name = first
        svc_order: list[str] = []
        svc_meta: dict[str, dict] = {}
        for it in items:
            code = it["service"]
            if code not in svc_meta:
                svc_order.append(code)
                svc_meta[code] = {
                    "name": it["service_name"] or it["research"],
                    "amount": 0,
                    "doctor": it["doctor"] or doctor,
                }
            svc_meta[code]["amount"] += 1

        tid = str(n)
        talon_rows.append(
            [
                tid,
                oms_date,
                oms_date,
                enp,
                patient,
                doctor,
                default_building,
                DEFAULTS_541["цель"],
                DEFAULTS_541["случай"],
                DEFAULTS_541["случай_2"],
                DEFAULTS_541["посещений в МО"],
                main_diag,
                diag2,
                diags_all,
                DEFAULTS_541["ДИСП. НАБЛ"],
                DEFAULTS_541["тип"],
                DEFAULTS_541["результат"],
                DEFAULTS_541["исход"],
                DEFAULTS_541["характер"],
                len(svc_order),
                len(items),
            ]
        )
        for code in svc_order:
            meta = svc_meta[code]
            service_rows.append(
                [
                    tid,
                    code,
                    meta["name"],
                    oms_date,
                    oms_date,
                    str(meta["amount"]),
                    meta["doctor"] or doctor,
                ]
            )

    # отдельные файлы (как cel_3)
    def write_simple(path: Path, rows: list[list]) -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = "данные"
        for row in rows:
            ws.append(row)
        _style_header(ws, len(rows[0]))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)

    write_simple(out_talons, talon_rows)
    write_simple(out_services, service_rows)

    # бандл: один xlsx с двумя листами + сводка
    wb = Workbook()
    ws_i = wb.active
    ws_i.title = "инструкция"
    info = [
        ["Цель 541 — КТ (амбулаторный талон Web.ОМС)"],
        [""],
        ["Источник", str(source)],
        ["Сырых строк", len(raw)],
        ["Талонов (ЕНП+дата)", len(keys)],
        ["Строк услуг", len(service_rows) - 1],
        ["Без даты (пропуск)", skipped_no_day],
        ["Явка", "1"],
        ["Формат дат", "дд-мм-гг"],
        [""],
        ["Правила объединения"],
        ["Ключ", "ЕНП + календарная дата исследования"],
        ["Диагноз", "первый уникальный МКБ по времени"],
        ["Диагноз 2", "второй уникальный МКБ (сопутствующий)"],
        ["Диагнозы_все", "все МКБ через «; »"],
        ["Услуги", "уникальные коды; Кол-во = число повторов кода"],
        ["Врач", "самый частый код в группе (из справочника/примечания)"],
        [""],
        ["Листы"],
        ["талоны", "карта талона для ввода"],
        ["услуги", "join по колонке «Талон»"],
        [""],
        ["Запуск"],
        ["", r".\.venv\Scripts\python run_541.py --limit 3 --building \"ГП №3\""],
    ]
    for row in info:
        ws_i.append(row)
    ws_i["A1"].font = Font(bold=True, size=13)

    ws_t = wb.create_sheet("талоны")
    for row in talon_rows:
        ws_t.append(row)
    _style_header(ws_t, len(TALON_COLUMNS))
    ws_t.freeze_panes = "A2"
    ws_t.auto_filter.ref = ws_t.dimensions

    ws_s = wb.create_sheet("услуги")
    for row in service_rows:
        ws_s.append(row)
    _style_header(ws_s, len(SERVICE_COLUMNS))
    ws_s.freeze_panes = "A2"
    ws_s.auto_filter.ref = ws_s.dimensions

    out_bundle.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_bundle)

    print(f"source rows: {len(raw)} → talons: {len(keys)}, service lines: {len(service_rows)-1}")
    if skipped_no_day:
        print(f"WARN skipped (no date): {skipped_no_day}")
    multi = sum(1 for k in keys if len(groups[k]) > 1)
    multi_svc = sum(1 for r in talon_rows[1:] if int(r[19]) > 1)
    multi_diag = sum(1 for r in talon_rows[1:] if r[12])
    no_doc = sum(1 for r in talon_rows[1:] if not r[5])
    print(f"groups with >1 source row: {multi}")
    print(f"talons with >1 unique service: {multi_svc}")
    print(f"talons with accompanying diag: {multi_diag}")
    print(f"talons without doctor code: {no_doc}")
    print(f"talons → {out_talons}")
    print(f"services → {out_services}")
    print(f"bundle → {out_bundle}")
    return len(keys), len(service_rows) - 1


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Сборка Excel цели 541")
    p.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    p.add_argument("--building", default="", help="Корпус по умолчанию")
    p.add_argument("--out-dir", type=Path, default=DATA_DIR)
    return p.parse_args()


def main() -> int:
    args = parse_args()
    src = args.source
    if not src.is_file():
        alt = DATA_DIR / src.name
        if alt.is_file():
            src = alt
        else:
            raise SystemExit(f"Нет источника: {args.source}")
    out_dir = args.out_dir
    build_talons(
        src,
        out_talons=out_dir / "talons_541.xlsx",
        out_services=out_dir / "services_541.xlsx",
        out_bundle=out_dir / "talons_541_bundle.xlsx",
        default_building=args.building,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
