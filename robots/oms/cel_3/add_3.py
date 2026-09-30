#!/usr/bin/env python
"""
Ввод талонов цели 3 (БСК / ДН) в Web.ОМС — амбулаторный талон.

Как oms_307, но:
  • два Excel: талоны + услуги (join по «Талон»), как в талоны2_0.ipynb
  • несколько услуг на талон (цикл + #add-service)

Пример:
  python run_3.py --zip data/talon_bundle_….zip --category БСК --limit 3 \\
      --building \"ГП №3\" --doctor 30097144
  python run_3.py --talons data/talons_3.xlsx --services data/services_3.xlsx --limit 1
"""

from __future__ import annotations

import argparse
import csv
import sys
import zipfile
from dataclasses import asdict, dataclass, field
from datetime import datetime
from pathlib import Path
from time import sleep, time

from loguru import logger
from openpyxl import load_workbook
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.remote.webdriver import WebDriver
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait

WORK_DIR = Path(__file__).resolve().parent
OMS_DIR = WORK_DIR.parent  # robots/oms
PROJECT_ROOT = WORK_DIR.parents[2]  # BrowserAuto
for p in (PROJECT_ROOT, WORK_DIR, OMS_DIR / "dv_opv", OMS_DIR / "cel_307"):
    if str(p) not in sys.path:
        sys.path.insert(0, str(p))

from browser_auto.config import load_credentials
from browser_auto.auth import login_oms
from browser_auto.driver import GeckoBrowser

import add_307 as base  # noqa: E402 — переиспользуем хелперы ambulatory
import add_medical_exam as dv  # noqa: E402

load_credentials(WORK_DIR / "credentials.env")

RESULTS_DIR = WORK_DIR / "data"
DEFAULT_ZIP = WORK_DIR / "data" / "talon_bundle_iszl_2026-09-29_033845.zip"

DEFAULTS_3 = {
    "цель": "3",
    "случай": "Первичный",
    "случай_2": "Законченный",
    "ДИСП. НАБЛ": "1",
    "тип": "3",
    "результат": "301",
    "исход": "304",
    "характер": "3",
}


@dataclass
class ServiceLine:
    code: str
    begin: str
    end: str
    amount: str = "1"
    doctor: str = ""


@dataclass
class Talon3:
    talon_id: str
    enp: str
    begin: str
    end: str
    doctor: str
    building: str
    goal: str = "3"
    occurrence: str = DEFAULTS_3["случай"]
    occurrence2: str = DEFAULTS_3["случай_2"]
    mo_visits: str = "3"
    diagnosis: str = ""
    diagnosis2: str = ""
    dn: str = DEFAULTS_3["ДИСП. НАБЛ"]
    claim_type: str = DEFAULTS_3["тип"]
    result: str = DEFAULTS_3["результат"]
    outcome: str = DEFAULTS_3["исход"]
    character: str = DEFAULTS_3["характер"]
    category: str = ""
    social_status: str = ""
    occupation: str = ""
    place: str = "1"
    home: str = base.DEFAULT_HOME
    ref_number: str = base.DEFAULT_REF_NUMBER
    ref_org: str = base.DEFAULT_REF_ORG
    ref_cond: str = base.DEFAULT_REF_COND
    # онкология (БСК часто Повод=3)
    onko_povod: str = ""
    onko_consilium: str = ""
    dop_napr_org: str = ""
    dop_napr_date: str = ""
    dop_napr_vid: str = ""
    services: list[ServiceLine] = field(default_factory=list)
    raw_index: int = 0

    def as_307(self) -> base.Talon307:
        """Карта полей для fill_main / fill_referral / patient search."""
        first = self.services[0] if self.services else None
        return base.Talon307(
            enp=self.enp,
            begin=self.begin,
            end=self.end,
            doctor=self.doctor,
            building=self.building,
            goal=self.goal or "3",
            occurrence=self.occurrence,
            occurrence2=self.occurrence2,
            mo_visits=self.mo_visits,
            diagnosis=self.diagnosis,
            diagnosis2=self.diagnosis2,
            dn=self.dn,
            claim_type=self.claim_type,
            result=self.result,
            outcome=self.outcome,
            character=self.character,
            service_code=(first.code if first else ""),
            service_amount=(first.amount if first else "1"),
            social_status=self.social_status,
            occupation=self.occupation,
            place=self.place,
            home=self.home,
            ref_number=self.ref_number,
            ref_org=self.ref_org,
            ref_cond=self.ref_cond,
            raw_index=self.raw_index,
        )


@dataclass
class RowResult:
    index: int
    talon_id: str
    enp: str
    services: int
    status: str
    message: str
    elapsed_sec: float = 0.0


def _xlsx_rows(path: Path) -> tuple[list[str], list[tuple]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    it = ws.iter_rows(values_only=True)
    header = [base._cell(h) for h in next(it)]
    # keep original header names (not lower) for display; map lower→idx
    raw_header = list(next(load_workbook(path, read_only=True, data_only=True).active.iter_rows(max_row=1, values_only=True)))
    wb.close()
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    header_raw = [str(h).strip() if h is not None else "" for h in next(rows_iter)]
    data = [tuple(r) for r in rows_iter if r and any(v is not None and str(v).strip() != "" for v in r)]
    wb.close()
    return header_raw, data


def _col(header: list[str], *names: str) -> int | None:
    idx = {h.lower().strip(): i for i, h in enumerate(header) if h}
    for n in names:
        if n.lower() in idx:
            return idx[n.lower()]
    return None


def extract_bundle(zip_path: Path, dest: Path) -> tuple[Path, Path]:
    dest.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(zip_path) as zf:
        zf.extractall(dest)
        names = zf.namelist()
    tal = next((dest / n for n in names if "талон" in n.lower() and n.lower().endswith(".xlsx")), None)
    svc = next((dest / n for n in names if "услуг" in n.lower() and n.lower().endswith(".xlsx")), None)
    if not tal or not svc or not tal.is_file() or not svc.is_file():
        raise SystemExit(f"В архиве нужны талоны*.xlsx и услуги*.xlsx, есть: {names}")
    return tal, svc


def load_talons_bundle(
    *,
    talons_path: Path,
    services_path: Path,
    default_building: str = "",
    default_doctor: str = "",
    category: str = "",
) -> list[Talon3]:
    th, trows = _xlsx_rows(talons_path)
    sh, srows = _xlsx_rows(services_path)

    i_id = _col(th, "талон")
    i_enp = _col(th, "енп", "enp")
    i_beg = _col(th, "начало")
    i_end = _col(th, "окончание")
    i_doc = _col(th, "врач", "doctor")
    i_bldg = _col(th, "корпус", "building", "subdivision")
    i_goal = _col(th, "цель", "goal")
    i_sl1 = _col(th, "случай")
    i_sl2 = _col(th, "случай_2", "случай 2")
    i_mo = _col(th, "посещений в мо", "посещений")
    i_dx = _col(th, "диагноз", "мкб")
    i_dx2 = _col(th, "диагноз 2", "диагноз2")
    i_dn = _col(th, "дисп. набл", "дн")
    i_type = _col(th, "тип")
    i_res = _col(th, "результат")
    i_out = _col(th, "исход")
    i_har = _col(th, "характер")
    i_cat = _col(th, "категория")
    i_ref_n = _col(th, "напр номер")
    i_ref_o = _col(th, "напр орган")
    i_ref_c = _col(th, "напр дано")
    i_povod = _col(th, "повод")
    i_consil = _col(th, "цель консил")
    i_dop_org = _col(th, "дополнительное направление организация")
    i_dop_date = _col(th, "дополнительное направление дата")
    i_dop_vid = _col(th, "дополнительное направление вид")

    if None in (i_id, i_enp, i_beg, i_end):
        raise SystemExit("В талонах нужны колонки: Талон, ЕНП, начало, окончание")

    si_id = _col(sh, "талон")
    si_code = _col(sh, "код услуги", "услуга")
    si_beg = _col(sh, "дата начала", "начало")
    si_end = _col(sh, "дата окончания", "окончание")
    si_amt = _col(sh, "кол-во", "количество")
    si_doc = _col(sh, "врач")
    if None in (si_id, si_code):
        raise SystemExit("В услугах нужны колонки: Талон, Код услуги")

    by_talon: dict[str, list[ServiceLine]] = {}
    for raw in srows:
        tid = base._cell(raw[si_id])
        code = base._cell(raw[si_code]).upper()
        if not tid or not code:
            continue
        by_talon.setdefault(tid, []).append(
            ServiceLine(
                code=code,
                begin=base._cell(raw[si_beg]) if si_beg is not None else "",
                end=base._cell(raw[si_end]) if si_end is not None else "",
                amount=(base._cell(raw[si_amt]) if si_amt is not None else "") or "1",
                doctor=base._cell(raw[si_doc]) if si_doc is not None else "",
            )
        )

    cat_filter = (category or "").strip().lower()
    out: list[Talon3] = []
    for n, raw in enumerate(trows, start=2):
        tid = base._cell(raw[i_id])
        enp = base._cell(raw[i_enp])
        if not tid or not enp:
            continue
        cat = base._cell(raw[i_cat]) if i_cat is not None else ""
        if cat_filter and cat.lower() != cat_filter:
            continue
        doctor = base._cell(raw[i_doc]) if i_doc is not None else ""
        building = base._cell(raw[i_bldg]) if i_bldg is not None else ""
        doctor = doctor or default_doctor
        building = building or default_building
        services = by_talon.get(tid, [])
        # проставить врача услуг по умолчанию
        for s in services:
            if not s.doctor:
                s.doctor = doctor
            if not s.begin:
                s.begin = base._cell(raw[i_beg])
            if not s.end:
                s.end = base._cell(raw[i_end])
            if not s.amount:
                s.amount = "1"

        out.append(
            Talon3(
                talon_id=tid,
                enp=enp,
                begin=base._cell(raw[i_beg]),
                end=base._cell(raw[i_end]),
                doctor=doctor,
                building=building,
                goal=(base._cell(raw[i_goal]) if i_goal is not None else "") or "3",
                occurrence=(base._cell(raw[i_sl1]) if i_sl1 is not None else "") or DEFAULTS_3["случай"],
                occurrence2=(base._cell(raw[i_sl2]) if i_sl2 is not None else "") or DEFAULTS_3["случай_2"],
                mo_visits=(base._cell(raw[i_mo]) if i_mo is not None else "") or "3",
                diagnosis=base._cell(raw[i_dx]) if i_dx is not None else "",
                diagnosis2=base._cell(raw[i_dx2]) if i_dx2 is not None else "",
                dn=(base._cell(raw[i_dn]) if i_dn is not None else "") or DEFAULTS_3["ДИСП. НАБЛ"],
                claim_type=(base._cell(raw[i_type]) if i_type is not None else "") or DEFAULTS_3["тип"],
                result=(base._cell(raw[i_res]) if i_res is not None else "") or DEFAULTS_3["результат"],
                outcome=(base._cell(raw[i_out]) if i_out is not None else "") or DEFAULTS_3["исход"],
                character=(base._cell(raw[i_har]) if i_har is not None else "") or DEFAULTS_3["характер"],
                category=cat,
                ref_number=(base._cell(raw[i_ref_n]) if i_ref_n is not None else "") or base.DEFAULT_REF_NUMBER,
                ref_org=(base._cell(raw[i_ref_o]) if i_ref_o is not None else "") or base.DEFAULT_REF_ORG,
                ref_cond=(base._cell(raw[i_ref_c]) if i_ref_c is not None else "") or base.DEFAULT_REF_COND,
                onko_povod=base._cell(raw[i_povod]) if i_povod is not None else "",
                onko_consilium=base._cell(raw[i_consil]) if i_consil is not None else "",
                dop_napr_org=base._cell(raw[i_dop_org]) if i_dop_org is not None else "",
                dop_napr_date=base._cell(raw[i_dop_date]) if i_dop_date is not None else "",
                dop_napr_vid=base._cell(raw[i_dop_vid]) if i_dop_vid is not None else "",
                services=services,
                raw_index=n,
            )
        )
    return out


def _click_xpath(driver: WebDriver, xpath: str) -> None:
    """Как click_xpath в талоны2_0."""
    el = driver.find_element(By.XPATH, xpath)
    el.click()


def _input_enter_xpath(driver: WebDriver, xpath: str, text: str) -> None:
    """Как input_enter_xpath в талоны2_0: send_keys(text, ENTER) без лишней магии."""
    if text is None:
        return
    el = driver.find_element(By.XPATH, xpath)
    el.send_keys(str(text), Keys.ENTER)


def _input_enter_id(driver: WebDriver, element_id: str, text: str) -> None:
    if text is None:
        return
    el = driver.find_element(By.ID, element_id)
    el.send_keys(str(text), Keys.ENTER)


def service_notebook(
    driver: WebDriver,
    service_code: str,
    start_date: str,
    end_date: str,
    amount: str,
    doctor_id: str,
    index: int,
) -> None:
    """Точная копия service() из талоны2.ipynb / талоны2_0.ipynb.

    Важно: после Add всегда бьём в *-service-0 (не -1/-2…).
    Даты и врач — только для index > 0.
    """
    if index > 0:
        _click_xpath(driver, '//*[@id="add-service"]')
        sleep(0.5)

    _input_enter_xpath(driver, '//*[@id="combobox-service-0"]', service_code)

    if index > 0:
        _input_enter_xpath(driver, '//*[@id="begin-date-service-0"]', start_date)
        _input_enter_xpath(driver, '//*[@id="end-date-service-0"]', end_date)

    _input_enter_xpath(driver, '//*[@id="amount-service-0"]', amount)
    perehod = driver.find_element(By.XPATH, '//*[@id="amount-service-0"]')
    perehod.clear()
    perehod.send_keys(str(amount), Keys.ENTER)

    if index > 0:
        _input_enter_xpath(driver, '//*[@id="doctor-service-0"]', doctor_id)
    sleep(0.5)
    logger.info(
        "service[{}] {} amt={} dates={}/{} doctor={}",
        index,
        service_code,
        amount,
        start_date if index > 0 else "(из талона)",
        end_date if index > 0 else "(из талона)",
        doctor_id if index > 0 else "(главный)",
    )


def fill_services(driver: WebDriver, row: Talon3) -> int:
    """Цикл услуг как в талоны2_0: for idx, service(...)."""
    if not row.services:
        raise RuntimeError(f"Нет услуг для талона {row.talon_id}")

    filled = 0
    for idx, svc in enumerate(row.services):
        code = (svc.code or "").strip()
        if not code:
            continue
        begin = dv.normalize_service_date(svc.begin or row.begin)
        end = dv.normalize_service_date(svc.end or row.end)
        amount = (svc.amount or "1").strip() or "1"
        doctor = (svc.doctor or row.doctor or "").strip()
        service_notebook(driver, code, begin, end, amount, doctor, idx)
        filled += 1
    logger.info("Услуг введено: {}/{}", filled, len(row.services))
    return filled


def process_row(driver: WebDriver, row: Talon3, *, index: int, dry_run: bool) -> RowResult:
    t0 = time()
    logger.info(
        "=== [{}] цель3 talon={} enp={} services={} cat={} ===",
        index,
        row.talon_id,
        row.enp,
        len(row.services),
        row.category or "—",
    )
    r307 = row.as_307()
    try:
        if not (row.building or "").strip():
            raise RuntimeError("Пустой Корпус — укажите --building")
        if not (row.doctor or "").strip():
            raise RuntimeError("Пустой врач — укажите --doctor")

        # Порядок как в талоны2_0 / oms_307: поиск → соцстатус → направ → main → услуги → save
        base.open_ambulatory(driver)
        base.fill_patient_search(driver, r307)
        ensure_social_status_oms3(driver, r307)
        row.social_status, row.occupation = r307.social_status, r307.occupation

        base.fill_referral(driver, r307)
        try:
            # на ambulatory соцстатус уже в карте пациента — не дёргать #socialStatus в fill_main
            soc_keep, occ_keep = r307.social_status, r307.occupation
            r307.social_status, r307.occupation = "", ""
            base.fill_main(driver, r307)
            r307.social_status, r307.occupation = soc_keep, occ_keep
        except Exception as exc:
            msg = f"Пациент не найден / main: {exc}"
            logger.error("[{}] {}", index, msg)
            try:
                driver.refresh()
                dv.wait_busy_gone(driver)
            except Exception:
                pass
            return RowResult(index, row.talon_id, row.enp, 0, "error", msg, round(time() - t0, 2))

        # повтор цели + пауза под шаблон (как input_enter_ID('requestPurpose'); sleep(2))
        _input_enter_id(driver, "requestPurpose", row.goal or "3")
        sleep(2.0)
        dv.wait_busy_gone(driver, timeout=20)

        n_svc = fill_services(driver, row)

        if dry_run:
            msg = f"dry-run: форма+{n_svc} услуг, без Save"
            logger.success("[{}] {}", index, msg)
            return RowResult(index, row.talon_id, row.enp, n_svc, "dry_run", msg, round(time() - t0, 2))

        ok, msg = base.save_ambulatory(driver)
        logger.info("[{}] Save: {}", index, msg)
        # как oms_307: при ошибке соцстатуса — ещё раз карта пациента и Save
        if (not ok) and dv._is_social_status_error(msg):
            logger.warning("[{}] ошибка соцстатуса — повторяем через карту пациента", index)
            ensure_social_status_oms3(driver, r307)
            base.fill_main(driver, r307)
            _input_enter_id(driver, "requestPurpose", row.goal or "3")
            sleep(1.0)
            ok, msg = base.save_ambulatory(driver)
            logger.info("[{}] Save (повтор): {}", index, msg)
        if (not ok) and (msg or "").strip().lower() in ("error", ""):
            soft = base.dump_claim_errors(driver, click_button=True)
            if soft:
                msg = soft
                logger.warning("[{}] детали ошибок: {}", index, soft[:400])
            try:
                (RESULTS_DIR / "last_save_fail.html").write_text(driver.page_source or "", encoding="utf-8")
            except Exception:
                pass
        status = "ok" if ok else "fail"
        log = logger.success if ok else logger.error
        log("[{}] {} → {} ({:.1f}s)", index, status.upper(), msg, time() - t0)
        return RowResult(index, row.talon_id, row.enp, n_svc, status, msg, round(time() - t0, 2))
    except Exception as exc:
        msg = str(exc).split("Stacktrace:")[0].strip() or str(exc)
        logger.error("[{}] error ({:.1f}s): {}", index, time() - t0, msg)
        try:
            driver.refresh()
            dv.wait_busy_gone(driver)
        except Exception:
            pass
        return RowResult(index, row.talon_id, row.enp, 0, "error", msg, round(time() - t0, 2))


def ensure_social_status_oms3(driver: WebDriver, row: base.Talon307) -> tuple[str, str]:
    """Соцстатус/занятость как в oms_307; на ambulatory — через карту пациента.

    oms_307.ensure_social_status:
      Excel → ДР → social_occupation_by_age (18<a<60 → 11/1, иначе 22/3)
      → #socialStatus / #occupation на форме талона, если есть.

    На /claim/ambulatory этих полей нет («Нет поля #socialStatus») —
    тогда как oms_dv_opv: main-tab → карта пациента → save-button.
    """
    excel_soc = (row.social_status or "").strip()
    excel_occ = (row.occupation or "").strip()
    if excel_soc and excel_occ:
        logger.info("Соцстатус/занятость из Excel: {} / {}", excel_soc, excel_occ)

    # расчёт + попытка на форме талона (как add_307)
    soc, occ = base.ensure_social_status(driver, row)

    on_claim = bool(driver.find_elements(By.ID, "socialStatus")) and bool(
        driver.find_elements(By.ID, "occupation")
    )
    if on_claim:
        for sid, val in (("socialStatus", soc), ("occupation", occ)):
            el = driver.find_element(By.ID, sid)
            wrap = dv._closest_react_select(el)
            shown = (wrap.text if wrap is not None else "") or ""
            if str(val) not in shown:
                dv.try_input_enter_id(driver, sid, val)
        logger.info("Соцстатус на форме талона: {} / {}", soc, occ)
        return soc, occ

    # поля только в карте пациента (как ДВ4)
    logger.info("На форме талона нет #socialStatus — открываем карту пациента")
    dv.dismiss_overlays(driver)
    if driver.find_elements(By.ID, "main-tab"):
        dv.click_id(driver, "main-tab")
        dv.wait_busy_gone(driver)
        sleep(0.4)

    btn = None
    for xp in (
        "//button[contains(.,'Редактировать')]",
        dv.PATIENT_CARD_BTN_XPATH,
        "//table//tbody/tr[5]/td//button",
        "//table//tbody/tr[6]/td//button",
        "//table//tbody/tr[.//button][last()]//button",
        "//button[contains(.,'пациент') or contains(.,'Пациент')]",
    ):
        els = [e for e in driver.find_elements(By.XPATH, xp) if e.is_displayed()]
        if els:
            btn = els[0]
            break
    if btn is None:
        logger.warning("Кнопка карты пациента не найдена — оставляем {} / {}", soc, occ)
        return soc, occ

    try:
        btn.click()
    except Exception:
        driver.execute_script("arguments[0].click();", btn)
    dv.wait_busy_gone(driver, timeout=30)
    sleep(0.6)

    try:
        WebDriverWait(driver, 15).until(EC.presence_of_element_located((By.ID, "socialStatus")))
    except Exception:
        logger.warning("Карта пациента не открылась (#socialStatus) — {} / {}", soc, occ)
        return soc, occ

    # ДР в карте — только если Excel не задал оба поля
    if (not excel_soc or not excel_occ) and driver.find_elements(By.ID, "birth-date"):
        birth_raw = driver.find_element(By.ID, "birth-date").get_attribute("value") or ""
        year = dv._birth_year(birth_raw)
        if year is not None:
            age = datetime.now().year - year
            soc, occ = dv.social_occupation_by_age(age)
            logger.info(
                "Карта ДР={} → возраст≈{} → socialStatus={}, occupation={}",
                birth_raw,
                age,
                soc,
                occ,
            )

    def _fill_select(element_id: str, value: str) -> None:
        el = driver.find_element(By.ID, element_id)
        dv.fill_react_select_input(driver, el, value)
        sleep(0.25)
        dv.wait_busy_gone(driver)
        wrap = dv._closest_react_select(el)
        shown = (wrap.text if wrap is not None else "") or ""
        if str(value) not in shown:
            el.send_keys(Keys.CONTROL, "a")
            el.send_keys(Keys.BACKSPACE)
            el.send_keys(str(value), Keys.ENTER)
            sleep(0.5)
            wrap = dv._closest_react_select(el)
            shown = (wrap.text if wrap is not None else "") or ""
        if str(value) not in shown:
            raise RuntimeError(f"{element_id}: не выбрано «{value}», сейчас «{shown[:80]}»")
        logger.info("{} → {}", element_id, shown.split("\n")[0][:80])

    _fill_select("socialStatus", soc)
    _fill_select("occupation", occ)

    saves = [e for e in driver.find_elements(By.ID, "save-button") if e.is_displayed()]
    if not saves:
        raise RuntimeError("save-button карты пациента не найден")
    try:
        saves[0].click()
    except Exception:
        driver.execute_script("arguments[0].click();", saves[0])
    dv.wait_busy_gone(driver, timeout=60)
    sleep(0.5)
    dlg = dv.warning_dialog_text(driver)
    if dlg or dv.find_visible_dialogs(driver):
        ok, info = dv.handle_post_save_prompts(driver, timeout=20)
        if not ok:
            raise RuntimeError(f"Карта пациента не сохранена: {info}")
        if info:
            logger.info("После save карты: {}", info[:200])
    else:
        msg = dv.snackbar_text(driver) if hasattr(dv, "snackbar_text") else ""
        if msg:
            logger.info("После save карты: {}", msg[:200])
        low = (msg or "").lower()
        if any(x in low for x in ("ошибка", "обязательн", "не указан")):
            raise RuntimeError(f"Карта пациента не сохранена: {msg}")

    if driver.find_elements(By.ID, "main-tab"):
        dv.click_id(driver, "main-tab")
        dv.wait_busy_gone(driver)
        sleep(0.3)

    row.social_status, row.occupation = soc, occ
    logger.info("Соцстатус через карту пациента: {} / {}", soc, occ)
    return soc, occ


def write_results(path: Path, results: list[RowResult]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = list(asdict(results[0]).keys()) if results else [
        "index", "talon_id", "enp", "services", "status", "message", "elapsed_sec"
    ]
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        for r in results:
            w.writerow(asdict(r))


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Ввод талонов цели 3 (Web.ОМС ambulatory)")
    p.add_argument("--zip", type=Path, default=None, help="Архив talon_bundle_*.zip")
    p.add_argument("--talons", type=Path, default=None)
    p.add_argument("--services", type=Path, default=None)
    p.add_argument("--category", default="БСК", help="Фильтр Категория (пусто = все)")
    p.add_argument("--doctor", default="", help="Врач по умолчанию, если в Excel пусто")
    p.add_argument("--building", default="", help="Корпус по умолчанию")
    p.add_argument("--limit", type=int, default=0)
    p.add_argument("--offset", type=int, default=0)
    p.add_argument("--dry-run", action="store_true")
    p.add_argument("--keep-open", type=int, default=0)
    p.add_argument("--worker", default="")
    p.add_argument("--window-x", type=int, default=None)
    p.add_argument("--window-y", type=int, default=None)
    return p.parse_args()


def main() -> int:
    args = parse_args()
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)

    if args.zip:
        zpath = args.zip if args.zip.is_file() else WORK_DIR / "data" / args.zip.name
        talons_path, services_path = extract_bundle(zpath, RESULTS_DIR / "_extract")
    else:
        talons_path = args.talons or (WORK_DIR / "data" / "talons_3_bsk_test.xlsx")
        services_path = args.services or (WORK_DIR / "data" / "services_3_bsk_test.xlsx")
        if not talons_path.is_file() or not services_path.is_file():
            if DEFAULT_ZIP.is_file():
                talons_path, services_path = extract_bundle(DEFAULT_ZIP, RESULTS_DIR / "_extract")
            else:
                logger.error("Нужны --zip или --talons/--services")
                return 1

    rows = load_talons_bundle(
        talons_path=talons_path,
        services_path=services_path,
        default_building=args.building,
        default_doctor=args.doctor,
        category=args.category,
    )
    if args.offset:
        rows = rows[args.offset :]
    if args.limit and args.limit > 0:
        rows = rows[: args.limit]

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    worker_tag = (args.worker or "").strip() or f"o{args.offset}"
    result_path = RESULTS_DIR / f"oms_3_result_{stamp}_{worker_tag}.csv"
    logger.info(
        "[{}] талонов={} cat={!r} doctor={!r} building={!r} dry_run={}",
        worker_tag,
        len(rows),
        args.category,
        args.doctor,
        args.building,
        args.dry_run,
    )
    if not rows:
        logger.error("Нет строк после фильтра")
        return 1

    results: list[RowResult] = []
    t_all = time()
    win_kwargs: dict = {"implicit_wait": 0, "maximize": True}
    if args.window_x is not None or args.window_y is not None:
        win_kwargs["maximize"] = False
        win_kwargs["window_position"] = (args.window_x or 20, args.window_y or 20)
        win_kwargs["window_size"] = (1100, 850)

    with GeckoBrowser(**win_kwargs) as session:
        driver = session.driver
        driver.implicitly_wait(0)
        login_oms(driver)
        sleep(0.5)
        dv.wait_busy_gone(driver, timeout=30)

        for i, row in enumerate(rows, start=1 + args.offset):
            res = process_row(driver, row, index=i, dry_run=args.dry_run)
            results.append(res)
            write_results(result_path, results)

        if args.keep_open > 0:
            sleep(args.keep_open)

    write_results(result_path, results)
    ok_n = sum(1 for r in results if r.status in ("ok", "dry_run"))
    fail_n = len(results) - ok_n
    logger.info(
        "——— [{}] ИТОГ ({:.0f}s) OK={}, FAIL={} → {} ———",
        worker_tag,
        time() - t_all,
        ok_n,
        fail_n,
        result_path,
    )
    for r in results:
        logger.info("  #{} talon={} {} | {}", r.index, r.talon_id, r.status, (r.message or "")[:160])
    return 0 if fail_n == 0 and results else 1


if __name__ == "__main__":
    sys.exit(main())
