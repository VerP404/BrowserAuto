# -*- coding: utf-8 -*-
"""Count talon rows for cel_3 run_parallel (UTF-8 safe argv)."""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook


def pick_sheet(wb):
    if "талоны" in wb.sheetnames:
        return wb["талоны"]
    for name in wb.sheetnames:
        ws = wb[name]
        row1 = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not row1:
            continue
        headers = [str(h or "").strip().lower() for h in row1]
        if "енп" in headers or "enp" in headers:
            return ws
    return wb.active


def main() -> int:
    if len(sys.argv) < 2:
        print("usage: _count_talons.py <talons.xlsx> [category]", file=sys.stderr)
        return 2
    path = Path(sys.argv[1])
    cat = (sys.argv[2] if len(sys.argv) > 2 else "").strip().lower()
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = pick_sheet(wb)
    rows = ws.iter_rows(values_only=True)
    headers = [str(h or "").strip().lower() for h in next(rows)]
    i_enp = next((i for i, h in enumerate(headers) if h in ("енп", "enp")), None)
    i_cat = next((i for i, h in enumerate(headers) if h == "категория"), None)
    n = 0
    for raw in rows:
        if not raw or i_enp is None:
            continue
        enp = raw[i_enp]
        if enp is None or str(enp).strip() == "":
            continue
        if cat and i_cat is not None:
            c = str(raw[i_cat] or "").strip().lower()
            if c != cat:
                continue
        n += 1
    wb.close()
    print(n)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
