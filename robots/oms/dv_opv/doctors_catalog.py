"""
Справочник ДВ4/ОПВ/УД1: врачи услуг + подписи групп ДН.

Источник (по приоритету):
  1) лист «справочник_врачей» / «ДН» в файле талонов
  2) --catalog / «Справочник ДВ4 и ОПВ.xlsx» рядом с роботом
  3) data/doctors_by_building.xlsx

Лист справочник_врачей: корпус | код_услуги | роль | врач
Лист ДН: код_дн | ДН

Логика врача услуги:
  - код услуги + корпус талона → врач из справочника;
  - иначе основной врач карты.
"""

from __future__ import annotations

import calendar
import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook, load_workbook

_MODULE_DIR = Path(__file__).resolve().parent
DEFAULT_REFERENCE = _MODULE_DIR / "Справочник ДВ4 и ОПВ.xlsx"
DEFAULT_CATALOG_FALLBACK = _MODULE_DIR / "data" / "doctors_by_building.xlsx"

# Запасной словарь ДН, если лист «ДН» не прочитался.
BUILTIN_DN_LABELS: dict[str, str] = {
    "1": "Состоит",
    "2": "Взят",
    "3": "Не подлежит",
    "7": "Состоит, проведено диспансерное наблюдение",
    "8": "Взят, проведено диспансерное наблюдение",
}

# ДВ4: флюорография / ЭКГ / маммография → «Перезачет», дата = окончание − N мес.
# (если лист «перезачет» в справочнике пуст — используем это)
# Перезачёт: ДВ4 и УД1 (одинаковая логика исследований)
_RECREDIT_TYPES = ("dv4", "ud1")

BUILTIN_RECREDIT_CODES: dict[str, dict[str, object]] = {
    "A06.09.006": {"name": "Флюорография", "done": "Перезачет", "months_back": 1, "exam_types": _RECREDIT_TYPES},
    "A05.10.006": {"name": "ЭКГ", "done": "Перезачет", "months_back": 1, "exam_types": _RECREDIT_TYPES},
    "A06.20.004": {"name": "Маммография", "done": "Перезачет", "months_back": 1, "exam_types": _RECREDIT_TYPES},
}


def _norm_exam_type_token(raw: str) -> str:
    t = (raw or "").strip().lower().replace("ё", "е")
    t = t.replace("дв-4", "dv4").replace("дв4", "dv4")
    t = t.replace("опв", "opv")
    t = t.replace("уд-1", "ud1").replace("уд 1", "ud1").replace("уд1", "ud1")
    return t



@dataclass
class RecreditRule:
    service_code: str
    name: str = ""
    done: str = "Перезачет"
    months_back: int = 1
    exam_types: tuple[str, ...] = ("dv4", "ud1")  # пусто = любой тип


@dataclass
class DoctorRule:
    building: str
    service_pattern: str  # точный код, префикс%, или *
    role: str  # therapist | specialist | lab | …
    doctor: str  # код врача; пусто → main_doctor из карты
    recredit: str = ""  # колонка «перезачет»: Перезачет / пусто


@dataclass
class DoctorCatalog:
    rules: list[DoctorRule] = field(default_factory=list)
    source: str = ""

    @classmethod
    def load(cls, path: Path, sheet: str | None = None) -> DoctorCatalog:
        if not path.is_file():
            return cls()
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            if sheet and sheet in wb.sheetnames:
                ws = wb[sheet]
            elif "справочник_врачей" in wb.sheetnames:
                ws = wb["справочник_врачей"]
            else:
                ws = wb.active
            rows = ws.iter_rows(values_only=True)
            header = [str(h or "").strip().lower() for h in next(rows)]
            idx = {h: i for i, h in enumerate(header) if h}

            def col(*names: str) -> int | None:
                for n in names:
                    if n in idx:
                        return idx[n]
                return None

            i_b = col("корпус", "building")
            i_s = col("код_услуги", "услуга", "service", "service_code")
            i_r = col("роль", "role")
            i_d = col("врач", "doctor", "код врача")
            i_re = col("перезачет", "recredit", "перезачёт")
            if i_s is None:
                return cls()

            rules: list[DoctorRule] = []
            for raw in rows:
                if not raw:
                    continue
                svc = str(raw[i_s] or "").strip()
                if not svc:
                    continue
                rules.append(
                    DoctorRule(
                        building=str(raw[i_b] or "").strip() if i_b is not None else "",
                        service_pattern=svc,
                        role=str(raw[i_r] or "").strip().lower() if i_r is not None else "",
                        doctor=str(raw[i_d] or "").strip() if i_d is not None else "",
                        recredit=str(raw[i_re] or "").strip() if i_re is not None else "",
                    )
                )
            return cls(rules=rules, source=str(path))
        finally:
            wb.close()

    @classmethod
    def load_for_talons_file(cls, talons_path: Path, fallback: Path | None = None) -> DoctorCatalog:
        """Сначала лист в файле талонов, иначе --catalog / Справочник ДВ4 и ОПВ.xlsx."""
        if talons_path.is_file():
            cat = cls.load(talons_path, sheet="справочник_врачей")
            if cat.rules:
                return cat

        candidates: list[Path] = []
        if fallback is not None:
            candidates.append(Path(fallback))
        candidates.extend([DEFAULT_REFERENCE, DEFAULT_CATALOG_FALLBACK])

        seen: set[str] = set()
        for path in candidates:
            key = str(path.resolve()) if path.exists() else str(path)
            if key in seen:
                continue
            seen.add(key)
            if not path.is_file():
                continue
            cat = cls.load(path, sheet="справочник_врачей")
            if cat.rules:
                return cat
            cat = cls.load(path)
            if cat.rules:
                return cat
        return cls()

    @staticmethod
    def ensure_example(path: Path) -> None:
        if path.is_file():
            return
        path.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = "справочник_врачей"
        ws.append(["корпус", "код_услуги", "роль", "врач"])
        ws.append(["ГП №3", "B01.001.001", "", "30136001"])
        ws.append(["ГП №3", "A09.05.%", "lab", ""])
        ws_dn = wb.create_sheet("ДН")
        ws_dn.append(["код_дн", "ДН"])
        for code, label in BUILTIN_DN_LABELS.items():
            ws_dn.append([int(code) if code.isdigit() else code, label])
        wb.save(path)


def load_dn_labels(path: Path | None = None) -> dict[str, str]:
    """код_дн → текст из листа «ДН» (для react-select ОМС)."""
    candidates: list[Path] = []
    if path is not None:
        candidates.append(Path(path))
    candidates.extend([DEFAULT_REFERENCE, DEFAULT_CATALOG_FALLBACK])

    seen: set[str] = set()
    for p in candidates:
        key = str(p.resolve()) if p.exists() else str(p)
        if key in seen:
            continue
        seen.add(key)
        if not p.is_file():
            continue
        labels = _read_dn_sheet(p)
        if labels:
            return labels
    return dict(BUILTIN_DN_LABELS)


def _read_dn_sheet(path: Path) -> dict[str, str]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if "ДН" not in wb.sheetnames:
            return {}
        ws = wb["ДН"]
        rows = ws.iter_rows(values_only=True)
        header = [str(h or "").strip().lower() for h in next(rows)]
        idx = {h: i for i, h in enumerate(header) if h}
        i_code = next((idx[n] for n in ("код_дн", "код", "dn", "код дн") if n in idx), None)
        i_label = next((idx[n] for n in ("дн", "название", "label", "текст") if n in idx), None)
        if i_code is None:
            return {}
        out: dict[str, str] = {}
        for raw in rows:
            if not raw:
                continue
            code = str(raw[i_code] or "").strip()
            if isinstance(raw[i_code], float) and float(raw[i_code]).is_integer():
                code = str(int(raw[i_code]))
            if not code:
                continue
            label = str(raw[i_label] or "").strip() if i_label is not None else ""
            out[code] = label
        return out
    finally:
        wb.close()


def dn_search_texts(code_or_text: str, labels: dict[str, str] | None = None) -> list[str]:
    """Варианты для react-select ДН: только подпись из справочника (не голый код)."""
    labels = labels or BUILTIN_DN_LABELS
    text = (code_or_text or "").strip()
    if not text:
        return []

    code = text if text.isdigit() else ""
    if not code:
        low = text.lower().replace("ё", "е")
        for c, lab in labels.items():
            if not lab:
                continue
            lab_l = lab.lower().replace("ё", "е")
            if lab_l == low or lab_l in low or low in lab_l:
                code = c
                break
        # уже пришёл готовый текст без кода
        if not code:
            return [text]

    label = (labels.get(code) or BUILTIN_DN_LABELS.get(code) or "").strip()
    if not label:
        return [text]

    # в ОМС в поле ДН — текст справочника, не цифра
    variants = [label, f"{code} {label}"]
    seen: set[str] = set()
    out: list[str] = []
    for v in variants:
        if v and v not in seen:
            seen.add(v)
            out.append(v)
    return out


def _match_pattern(pattern: str, service_code: str) -> bool:
    p = (pattern or "").strip()
    code = (service_code or "").strip()
    if not p or p == "*":
        return True
    if p.endswith("%"):
        return code.startswith(p[:-1])
    if "*" in p or "?" in p:
        rx = re.escape(p).replace(r"\*", ".*").replace(r"\?", ".")
        return re.fullmatch(rx, code, flags=re.IGNORECASE) is not None
    return p.replace(" ", "").lower() == code.replace(" ", "").lower()


def _norm_building(value: str) -> str:
    """ГП №3 / ГП 3 / ГП#3 → одинаковый ключ."""
    text = (value or "").strip().lower().replace("ё", "е")
    text = text.replace("№", "").replace("#", "").replace("no", "")
    text = re.sub(r"[^0-9a-zа-я]+", "", text)
    return text


def resolve_service_doctor(
    catalog: DoctorCatalog,
    *,
    building: str,
    service_code: str,
    main_doctor: str,
) -> str:
    """Врач услуги: справочник (корпус+код) → иначе врач карты талона."""
    b = _norm_building(building)
    code = (service_code or "").strip()

    # 1) точное совпадение корпус + код_услуги
    # 2) код без корпуса / wildcard
    scored: list[tuple[int, DoctorRule]] = []
    for rule in catalog.rules:
        if not _match_pattern(rule.service_pattern, code):
            continue
        rb = _norm_building(rule.building)
        if rb and b and rb != b:
            continue
        exact = 0 if rule.service_pattern not in ("*", "") and "%" not in rule.service_pattern and "*" not in rule.service_pattern else 1
        building_rank = 0 if rb and b and rb == b else (1 if not rb else 2)
        scored.append((building_rank * 10 + exact, rule))

    scored.sort(key=lambda x: x[0])
    for _, rule in scored:
        if rule.doctor:
            return rule.doctor
        return main_doctor
    return main_doctor


def load_recredit_rules(path: Path | None = None) -> dict[str, RecreditRule]:
    """Правила перезачёта: колонка «перезачет» в справочник_врачей (или лист «перезачет»)."""
    candidates: list[Path] = []
    if path is not None:
        candidates.append(Path(path))
    candidates.extend([DEFAULT_REFERENCE, DEFAULT_CATALOG_FALLBACK])

    seen: set[str] = set()
    for p in candidates:
        key = str(p.resolve()) if p.exists() else str(p)
        if key in seen:
            continue
        seen.add(key)
        if not p.is_file():
            continue
        # 1) отдельный лист (если есть)
        rules = _read_recredit_sheet(p)
        if rules:
            return rules
        # 2) колонка «перезачет» в справочник_врачей — как в актуальном Excel
        rules = _read_recredit_from_doctors_sheet(p)
        if rules:
            return rules
    # builtin
    out: dict[str, RecreditRule] = {}
    for code, meta in BUILTIN_RECREDIT_CODES.items():
        out[code.upper()] = RecreditRule(
            service_code=code.upper(),
            name=str(meta.get("name") or ""),
            done=str(meta.get("done") or "Перезачет"),
            months_back=int(meta.get("months_back") or 1),
            exam_types=tuple(meta.get("exam_types") or ("dv4",)),
        )
    return out


def _read_recredit_from_doctors_sheet(path: Path) -> dict[str, RecreditRule]:
    """Из колонки «перезачет» листа справочник_врачей → уникальные коды услуг."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if "справочник_врачей" in wb.sheetnames:
            ws = wb["справочник_врачей"]
        else:
            ws = wb.active
        rows = ws.iter_rows(values_only=True)
        header = [str(h or "").strip().lower() for h in next(rows)]
        idx = {h: i for i, h in enumerate(header) if h}
        i_code = next((idx[n] for n in ("код_услуги", "услуга", "service", "service_code") if n in idx), None)
        i_re = next((idx[n] for n in ("перезачет", "перезачёт", "recredit") if n in idx), None)
        i_comm = next((idx[n] for n in ("комментарий", "название", "name") if n in idx), None)
        if i_code is None or i_re is None:
            return {}
        out: dict[str, RecreditRule] = {}
        for raw in rows:
            if not raw:
                continue
            code = str(raw[i_code] or "").strip().upper()
            flag = str(raw[i_re] or "").strip()
            if not code or not flag:
                continue
            low = flag.lower().replace("ё", "е")
            if "перезач" in low or low in ("1", "да", "true", "yes", "+", "x"):
                done = "Перезачет"
            else:
                done = flag
            name = str(raw[i_comm] or "").strip() if i_comm is not None else ""
            out[code] = RecreditRule(
                service_code=code,
                name=name,
                done=done,
                months_back=1,
                exam_types=_RECREDIT_TYPES,
            )
        return out
    finally:
        wb.close()


def _read_recredit_sheet(path: Path) -> dict[str, RecreditRule]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet_name = None
        for sn in wb.sheetnames:
            low = sn.strip().lower()
            if low in ("перезачет", "правила_услуг", "recredit"):
                sheet_name = sn
                break
        if not sheet_name:
            return {}
        ws = wb[sheet_name]
        rows = ws.iter_rows(values_only=True)
        header = [str(h or "").strip().lower() for h in next(rows)]
        idx = {h: i for i, h in enumerate(header) if h}

        def col(*names: str) -> int | None:
            for n in names:
                if n in idx:
                    return idx[n]
            return None

        i_code = col("код_услуги", "услуга", "service", "код")
        i_name = col("название", "name", "наименование")
        i_done = col("выполнено", "done", "статус")
        i_months = col("месяцев_назад", "месяцев", "months_back", "сдвиг")
        i_type = col("тип", "exam_type", "для_типа")
        if i_code is None:
            return {}

        out: dict[str, RecreditRule] = {}
        for raw in rows:
            if not raw:
                continue
            code = str(raw[i_code] or "").strip().upper()
            if not code:
                continue
            months = 1
            if i_months is not None and raw[i_months] is not None and str(raw[i_months]).strip() != "":
                try:
                    months = int(float(str(raw[i_months]).replace(",", ".")))
                except ValueError:
                    months = 1
            types_raw = str(raw[i_type] or "").strip().lower() if i_type is not None else "dv4,ud1"
            if not types_raw:
                types = _RECREDIT_TYPES
            else:
                types = tuple(
                    _norm_exam_type_token(t)
                    for t in re.split(r"[,;/|\s]+", types_raw)
                    if t.strip()
                )
            out[code] = RecreditRule(
                service_code=code,
                name=str(raw[i_name] or "").strip() if i_name is not None else "",
                done=(str(raw[i_done] or "").strip() if i_done is not None else "") or "Перезачет",
                months_back=max(0, months),
                exam_types=types or _RECREDIT_TYPES,
            )
        return out
    finally:
        wb.close()


def resolve_service_done_and_date(
    rules: dict[str, RecreditRule] | None,
    *,
    exam_type: str,
    service_code: str,
    end_date: str,
    default_done: str = "Да",
) -> tuple[str, str, RecreditRule | None]:
    """
    Для услуги: (выполнено, дата, правило|None).
    Если код в справочнике перезачёта и тип талона подходит —
    «Перезачет» и дата = окончание − months_back.
    """
    code = (service_code or "").strip().upper()
    et = _norm_exam_type_token(exam_type)
    rule = (rules or {}).get(code)
    if rule is None:
        return default_done, end_date, None
    allowed = tuple(t.lower() for t in (rule.exam_types or ()))
    if allowed and et not in allowed and "*" not in allowed:
        return default_done, end_date, None
    shifted = shift_date_months(end_date, -int(rule.months_back or 0)) or end_date
    return (rule.done or "Перезачет"), shifted, rule


def shift_date_months(value: str, months: int) -> str:
    """Сдвиг даты DD-MM-YY / DD.MM.YYYY на N месяцев. Возврат DD-MM-YY."""
    text = (value or "").strip()
    if not text:
        return ""
    m = re.match(r"^(\d{1,2})[-./](\d{1,2})[-./](\d{2,4})$", text)
    if not m:
        return text
    d, mo, y = m.groups()
    yi = int(y)
    if yi < 100:
        yi += 2000
    mi = int(mo)
    di = int(d)
    # сдвиг месяцев
    mi2 = mi - 1 + months
    yi2 = yi + mi2 // 12
    mi2 = mi2 % 12 + 1
    di2 = min(di, calendar.monthrange(yi2, mi2)[1])
    return f"{di2:02d}-{mi2:02d}-{str(yi2)[-2:]}"
