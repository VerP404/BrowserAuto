#!/usr/bin/env python
"""
Ввод талонов ДВ4 / ОПВ в Web.ОМС (medicalExamination).

URL:
  ДВ4 → /claim/medicalExamination/dv4
  ОПВ → /claim/medicalExamination/opv

Страница 1: даты начала/окончания (по ним подбираются услуги), ЕНП, врач, результат, период, диагноз, ДН.
#medicalExaminationPlace не трогаем.
Вкладка услуг: для каждой строки — врач из «справочник_врачей» (если код есть),
иначе основной врач карты; «Выполнено», дата.

Пример:
  python add_medical_exam.py --file data/talons.xlsx --limit 1
  python add_medical_exam.py --type opv --offset 0 --limit 5 --dry-run
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
from dataclasses import asdict, dataclass, field
from datetime import datetime
from pathlib import Path
from time import sleep, time

from loguru import logger
from openpyxl import Workbook, load_workbook
from selenium.common.exceptions import NoSuchElementException, StaleElementReferenceException, TimeoutException
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.remote.webdriver import WebDriver
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait

WORK_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = WORK_DIR.parents[2]  # BrowserAuto (robots/oms/dv_opv)
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))
if str(WORK_DIR) not in sys.path:
    sys.path.insert(0, str(WORK_DIR))

from browser_auto.config import OMS_BASE_URL, load_credentials
from browser_auto.auth import login_oms
from browser_auto.driver import GeckoBrowser

from doctors_catalog import (
    DEFAULT_REFERENCE,
    DoctorCatalog,
    dn_search_texts,
    load_dn_labels,
    load_recredit_rules,
    resolve_service_doctor,
    resolve_service_done_and_date,
)

load_credentials(WORK_DIR / "credentials.env")

DEFAULT_FILE = WORK_DIR / "data" / "talons_dv_opv.xlsx"
RESULTS_DIR = WORK_DIR / "data"
URL_BY_TYPE = {
    "dv4": "/claim/medicalExamination/dv4",
    "опв": "/claim/medicalExamination/opv",
    "opv": "/claim/medicalExamination/opv",
    "дв4": "/claim/medicalExamination/dv4",
}

# (legacy) колонка «место» в Excel читается, но #medicalExaminationPlace не заполняем
DEFAULT_PLACE = ""
DEFAULT_PERIOD = "январь"
DEFAULT_DONE = "Да"  # combobox-failure-services-* («Выполнено») — только текст, не код


def normalize_service_done(value: object) -> str:
    """Выполнено в ОМС: только «Да» или «Перезачет» (цифры/коды из Excel не годим)."""
    text = str(value or "").strip()
    if not text:
        return DEFAULT_DONE
    low = text.lower().replace("ё", "е")
    if "перезач" in low:
        return "Перезачет"
    if low in ("да", "yes", "true", "1", "1.0", "+", "выполнено"):
        return "Да"
    if low in ("нет", "no", "0", "false"):
        return "Да"  # для услуг ДВ4/ОПВ всегда отмечаем выполнение текстом
    # любой другой мусор (код врача и т.п.) — не пускаем
    if re.fullmatch(r"\d+([.,]\d+)?", text):
        return DEFAULT_DONE
    if text in ("Да", "Перезачет"):
        return text
    return DEFAULT_DONE

# соцстатус/занятость по возрасту (см. ensure_patient_social_status)
# 18<a<60 → 11/1; иначе (в т.ч. ≥60) → 22/3
DEFAULT_SOCIAL_STATUS = "22"
DEFAULT_OCCUPATION = "3"
# Назначения (purposes): при результате III группы ОМС требует вкладку
DEFAULT_PURPOSE_TYPE = "1"
DEFAULT_PURPOSE_FROM = "95"
DEFAULT_PURPOSE_TO = "95"

# Кнопка открытия карты пациента на вкладке main-tab
PATIENT_CARD_BTN_XPATH = (
    '//*[@id="root"]/div/div[2]/div[3]/div/div[1]/div/div/div[3]'
    "/div/div/div/div/table/tbody/tr[6]/td/button"
)


@dataclass
class TalonRow:
    enp: str
    begin: str
    end: str
    exam_type: str  # dv4 | opv
    doctor: str
    result: str
    period: str
    diagnosis: str
    dn: str
    building: str = ""
    place: str = DEFAULT_PLACE
    done_value: str = DEFAULT_DONE
    purpose_type: str = ""  # наз_вид → purposeType-0
    purpose_from: str = ""  # наз_от_кого → specialityId
    purpose_to: str = ""  # наз_к_кому → referenceSpecialityId
    social_status: str = ""  # пусто → посчитать по ДР в карте пациента
    occupation: str = ""
    # Сопутствующие: (диагноз, ДН), заполняются в diagnosis-1… / is-dispensary-observation-1…
    extra_diagnoses: list[tuple[str, str]] = field(default_factory=list)
    raw_index: int = 0


@dataclass
class RowResult:
    index: int
    enp: str
    exam_type: str
    status: str
    message: str
    services_filled: int = 0
    elapsed_sec: float = 0.0


def _wait(driver: WebDriver, timeout: float = 12) -> WebDriverWait:
    return WebDriverWait(driver, timeout, poll_frequency=0.15)


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d-%m-%y")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    # Excel часто даёт даты как datetime уже обработали; dd.mm.yyyy → dd-mm-yy
    m = re.match(r"^(\d{1,2})[./](\d{1,2})[./](\d{2,4})$", text)
    if m:
        d, mo, y = m.groups()
        if len(y) == 4:
            y = y[-2:]
        return f"{int(d):02d}-{int(mo):02d}-{y}"
    return text


def normalize_oms_date(value: str) -> str:
    """Единый формат дат Web.ОМС: DD-MM-YY (маска гг)."""
    return normalize_service_date(value)


def normalize_service_date(value: object) -> str:
    """Дата в ОМС: DD-MM-YY (гг, не гггг — иначе 2026 → 20)."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d-%m-%y")
    text = str(value).strip()
    if not text:
        return ""
    m = re.match(r"^(\d{1,2})[-./](\d{1,2})[-./](\d{2,4})$", text)
    if not m:
        # уже цифры без разделителей DDMMYY / DDMMYYYY
        digits = re.sub(r"\D", "", text)
        if len(digits) == 6:
            return f"{digits[0:2]}-{digits[2:4]}-{digits[4:6]}"
        if len(digits) == 8:
            return f"{digits[0:2]}-{digits[2:4]}-{digits[6:8]}"
        return text
    d, mo, y = m.groups()
    if len(y) == 4:
        y = y[-2:]
    return f"{int(d):02d}-{int(mo):02d}-{y}"


def normalize_mkb(value: str) -> str:
    """МКБ: кириллица → латиница (Е→E, К→K, …), верхний регистр буквы."""
    text = (value or "").strip().upper()
    if not text:
        return ""
    table = str.maketrans(
        {
            "А": "A",
            "В": "B",
            "Е": "E",
            "К": "K",
            "М": "M",
            "Н": "H",
            "О": "O",
            "Р": "P",
            "С": "C",
            "Т": "T",
            "Х": "X",
        }
    )
    return text.translate(table)


def _set_input_value(driver: WebDriver, el, text: str) -> None:
    """Надёжная установка value (masked / not reachable by keyboard)."""
    driver.execute_script(
        """
        const el = arguments[0], val = arguments[1];
        el.removeAttribute('readonly');
        el.removeAttribute('disabled');
        el.disabled = false;
        el.readOnly = false;
        const setter = Object.getOwnPropertyDescriptor(
            window.HTMLInputElement.prototype, 'value'
        ).set;
        el.focus();
        setter.call(el, '');
        el.dispatchEvent(new Event('input', {bubbles:true}));
        setter.call(el, val);
        el.dispatchEvent(new Event('input', {bubbles:true}));
        el.dispatchEvent(new Event('change', {bubbles:true}));
        el.dispatchEvent(new KeyboardEvent('keydown', {key:'Tab', bubbles:true}));
        el.dispatchEvent(new Event('blur', {bubbles:true}));
        """,
        el,
        text,
    )


def _norm_type(raw: str) -> str:
    t = (raw or "").strip().lower().replace(" ", "")
    if t in ("dv4", "дв4", "дв-4"):
        return "dv4"
    if t in ("opv", "опв"):
        return "opv"
    raise ValueError(f"Неизвестный тип талона: {raw!r} (нужно ДВ4 или ОПВ)")


def exam_url(exam_type: str) -> str:
    base = (OMS_BASE_URL or "http://10.36.0.142:9000").rstrip("/")
    path = URL_BY_TYPE[_norm_type(exam_type)]
    return f"{base}{path}"


def wait_busy_gone(driver: WebDriver, timeout: float = 30) -> None:
    """Ждём исчезновения «Пожалуйста, подождите...»."""
    end = time() + timeout
    while time() < end:
        try:
            texts = driver.find_elements(By.CSS_SELECTOR, "h2")
            busy = any("подождите" in (el.text or "").lower() for el in texts)
            if not busy:
                return
        except Exception:
            return
        sleep(0.25)


def input_id(driver: WebDriver, element_id: str, text: str, *, clear: bool = True) -> None:
    el = driver.find_element(By.ID, element_id)
    if clear:
        el.clear()
    el.send_keys(text)


def input_enter_id(driver: WebDriver, element_id: str, text: str) -> None:
    el = driver.find_element(By.ID, element_id)
    el.clear()
    el.send_keys(text, Keys.ENTER)


def try_input_enter_id(driver: WebDriver, element_id: str, text: str) -> bool:
    if not text:
        return False
    els = driver.find_elements(By.ID, element_id)
    if not els:
        return False
    el = els[0]
    el.clear()
    el.send_keys(text, Keys.ENTER)
    return True


def click_id(driver: WebDriver, element_id: str) -> None:
    el = driver.find_element(By.ID, element_id)
    try:
        el.click()
    except Exception:
        driver.execute_script("arguments[0].click();", el)


def snackbar_text(driver: WebDriver) -> str:
    """Текст из #client-snackbar / MUI Snackbar / Alert."""
    chunks: list[str] = []
    try:
        el = driver.find_element(By.ID, "client-snackbar")
        try:
            p = el.find_element(By.TAG_NAME, "p")
            text = (p.text or "").strip()
        except Exception:
            text = (el.text or "").strip()
        if text:
            chunks.append(text)
    except NoSuchElementException:
        pass
    for sel in (
        ".MuiSnackbar-root .MuiAlert-message",
        ".MuiSnackbarContent-message",
        "[class*='Snackbar'] [class*='message']",
        ".notistack-Snackbar",
    ):
        for el in driver.find_elements(By.CSS_SELECTOR, sel):
            try:
                if not el.is_displayed():
                    continue
                t = (el.text or "").strip()
                if t and t not in chunks:
                    chunks.append(t)
            except Exception:
                continue
    return " | ".join(chunks)


def collect_form_errors(driver: WebDriver) -> str:
    """Собрать подсказки валидации / helper-text / aria-invalid рядом с полями."""
    parts: list[str] = []
    for sel in (
        ".MuiFormHelperText-root.Mui-error",
        ".Mui-error .MuiFormHelperText-root",
        "[class*='error'][class*='HelperText']",
        ".field-error",
        ".error-message",
        ".MuiAlert-message",
    ):
        for el in driver.find_elements(By.CSS_SELECTOR, sel):
            try:
                if not el.is_displayed():
                    continue
                t = (el.text or "").strip()
                if t and t not in parts and t.upper() != "ОШИБКИ":
                    parts.append(t)
            except Exception:
                continue
    # панель ошибок талона (после клика claim-error)
    panel = read_claim_error_panel(driver, click=False)
    if panel and panel not in parts:
        parts.append(panel)
    # aria-invalid поля
    try:
        bad = driver.execute_script(
            """
            const out = [];
            document.querySelectorAll('[aria-invalid=\"true\"]').forEach(el => {
              const id = el.id || el.getAttribute('name') || '';
              const lab = (el.getAttribute('aria-label') || '').slice(0, 80);
              out.push((id || lab || el.tagName) + '');
            });
            return out.slice(0, 20);
            """
        )
        if bad:
            parts.append("aria-invalid: " + ", ".join(str(x) for x in bad))
    except Exception:
        pass
    return " | ".join(parts)[:800]


def read_claim_error_panel(driver: WebDriver, *, click: bool = True) -> str:
    """Текст панели ошибок талона (#claim-error) — короткие пункты, не вся страница."""
    try:
        btns = driver.find_elements(By.ID, "claim-error")
        if not btns:
            return ""
        btn = btns[0]
        if click and btn.is_displayed():
            try:
                _click_el(driver, btn)
            except Exception:
                driver.execute_script("arguments[0].click();", btn)
            sleep(0.45)
        text = driver.execute_script(
            """
            const pick = (root) => {
              if (!root) return '';
              const items = [...root.querySelectorAll('li, .MuiListItem-root, .MuiAlert-message, [class*=\"error\"]') ]
                .map(el => (el.innerText||'').trim())
                .filter(t => t && t.length > 3 && t.toUpperCase() !== 'ОШИБКИ'
                       && !t.includes('СОХРАНИТЬ') && !t.includes('КАРТА'));
              if (items.length) return [...new Set(items)].join(' | ').slice(0, 500);
              return '';
            };
            // drawer / popover / dialog рядом
            for (const sel of ['.MuiDrawer-root','.MuiPopover-root','.MuiModal-root','.MuiDialog-root','[role=\"presentation\"]']) {
              for (const el of document.querySelectorAll(sel)) {
                if (el.offsetParent === null && getComputedStyle(el).visibility === 'hidden') continue;
                const t = pick(el);
                if (t) return t;
              }
            }
            const btn = document.getElementById('claim-error');
            if (btn) {
              let p = btn.parentElement;
              for (let i = 0; i < 4 && p; i++, p = p.parentElement) {
                const t = pick(p);
                if (t) return t;
              }
            }
            return '';
            """
        )
        return (text or "").strip()
    except Exception:
        return ""


def normalize_dn(value: str) -> str:
    """ДН: текст → код (как на листе «ДН»)."""
    text = (value or "").strip()
    if not text:
        return ""
    if text.isdigit():
        return text
    low = text.lower().replace("ё", "е")
    # длинные ключи первыми («не состоит» / «не подлежит» раньше «состоит»)
    mapping = (
        ("не подлежит", "3"),
        ("не состоит", "3"),
        ("состоит, проведено", "7"),
        ("взят, проведено", "8"),
        ("проведено диспансерное наблюдение", "7"),
        ("достижение возрастного ограничения", "7"),
        ("снятие", "7"),
        ("снят", "7"),
        ("состоит", "1"),
        ("взят", "2"),
    )
    for key, code in mapping:
        if key in low:
            return code
    return text


def _visible_enabled(el) -> bool:
    try:
        if not el.is_displayed():
            return False
        if el.get_attribute("disabled") is not None:
            return False
        aria = (el.get_attribute("aria-disabled") or "").lower()
        if aria in ("true", "1"):
            return False
        return el.is_enabled()
    except Exception:
        return False


def _click_el(driver: WebDriver, el) -> None:
    driver.execute_script("arguments[0].scrollIntoView({block:'center'});", el)
    sleep(0.15)
    try:
        el.click()
    except Exception:
        driver.execute_script("arguments[0].click();", el)


def _btn_label(driver: WebDriver, btn) -> str:
    try:
        return (
            driver.execute_script(
                "return (arguments[0].innerText || arguments[0].textContent || '').trim();",
                btn,
            )
            or ""
        ).strip()
    except Exception:
        return (btn.text or "").strip()


def find_visible_dialogs(driver: WebDriver) -> list:
    """Видимые модалки ОМС (MUI Dialog / role=dialog)."""
    found = []
    seen = set()
    for sel in ("[role='dialog']", ".MuiDialog-paper", ".MuiPaper-root[role='dialog']"):
        for el in driver.find_elements(By.CSS_SELECTOR, sel):
            try:
                if not el.is_displayed():
                    continue
                key = (el.get_attribute("outerHTML") or "")[:160]
                if key in seen:
                    continue
                seen.add(key)
                found.append(el)
            except Exception:
                continue
    return found


def warning_dialog_text(driver: WebDriver) -> str:
    """Полный текст видимого диалога (заголовок + тело)."""
    parts: list[str] = []
    for dlg in find_visible_dialogs(driver):
        try:
            t = (dlg.text or "").strip()
        except Exception:
            t = ""
        if t:
            lines = [ln for ln in t.splitlines() if "подождите" not in ln.lower()]
            parts.append("\n".join(lines).strip())
    if parts:
        return "\n---\n".join(parts)
    # fallback: модалка ОМС без role=dialog (классы jss*)
    for h2 in driver.find_elements(By.CSS_SELECTOR, "h2"):
        try:
            text = (h2.text or "").strip()
        except Exception:
            continue
        low = text.lower()
        if not text or "подождите" in low:
            continue
        if "предупрежд" in low or "внимание" in low or "проверк" in low:
            # взять текст родителя (заголовок + тело + кнопки)
            try:
                parent = h2.find_element(By.XPATH, "./ancestor::div[4]")
                blob = (parent.text or text).strip()
                lines = [ln for ln in blob.splitlines() if "подождите" not in ln.lower()]
                return "\n".join(lines).strip() or text
            except Exception:
                return text
    return ""


def _is_intersection_warning(text: str) -> bool:
    """Пересечение с КС/СТАЦ/другим талоном — сохранять нельзя."""
    low = (text or "").lower()
    return "пересечен" in low  # пересечение / пересечения / пересечением


def _is_blocking_dialog(text: str) -> bool:
    """Жёсткий отказ: дубликат, пересечение, недопустимое ДН — не жать подтверждение."""
    low = (text or "").lower().replace("ё", "е")
    if "дубликат" in low:
        return True
    if _is_intersection_warning(low):
        return True
    # нельзя править ДН из диалога «для пациента» — только НАЗАД и правка на форме
    if "недопустим" in low and ("диспансерн" in low or "наблюден" in low):
        return True
    if "диспансерн" in low and "диагноз" in low and any(
        x in low for x in ("недопустим", "неверн", "нельзя", "исправить")
    ):
        return True
    if any(
        x in low
        for x in (
            "имеются предупрежд",
            "проверка врачей",
            "данные пациента были изменены",
            "сохранить изменения",
            "для талона",
        )
    ):
        return False
    blockers = (
        "результат поиска в црп",
        "отсутствует пациент",
    )
    return any(x in low for x in blockers)


def _is_confirmable_warning(text: str) -> bool:
    """Диалоги после Save, которые нужно подтвердить кнопкой."""
    low = (text or "").lower()
    if not low.strip():
        return False
    # пересечение / жёсткий дубликат — не подтверждаем
    if _is_blocking_dialog(low):
        return False
    markers = (
        "предупрежд",
        "проверка врачей",
        "рекоменд",
        "внимание",
        "имеются",
        "данные пациента были изменены",
        "сохранить изменения",
        "для талона",
    )
    return any(x in low for x in markers)


def dismiss_warning_dialog(driver: WebDriver) -> tuple[bool, str]:
    """Закрыть диалог без сохранения (НАЗАД / Отмена).

    Если есть «Сохранить с предупреждением» — НЕ гасим (нужно подтвердить).
    """
    text = warning_dialog_text(driver)
    low_text = (text or "").lower().replace("ё", "е")
    buttons: list[tuple[str, object]] = []
    xpaths = (
        "//button[contains(.,'НАЗАД') or contains(.,'Назад')]",
        "//button[contains(.,'Отмена') or contains(.,'ОТМЕНА')]",
        "//div[@role='dialog']//button",
        "//div[contains(@class,'MuiDialog')]//button",
        "//div[contains(@class,'MuiDialogActions')]//button",
        "/html/body/div[3]//button",
        "/html/body/div[4]//button",
        "/html/body/div[5]//button",
    )
    seen = set()
    for xp in xpaths:
        for btn in driver.find_elements(By.XPATH, xp):
            try:
                if not btn.is_displayed():
                    continue
                key = btn.id
                if key in seen:
                    continue
                seen.add(key)
            except Exception:
                continue
            buttons.append((_btn_label(driver, btn), btn))

    btn_blob = " | ".join((l or "").lower().replace("\n", " ") for l, _ in buttons)
    if (
        "сохранить с предупреждением" in btn_blob
        or "с предупреждением" in btn_blob
        or "имеются предупрежд" in low_text
    ):
        logger.info(
            "dismiss пропущен — диалог предупреждений (нужно «Сохранить с предупреждением»)"
        )
        return False, text or "нужно подтвердить предупреждение"

    if not buttons:
        try:
            driver.switch_to.active_element.send_keys(Keys.ESCAPE)
            wait_busy_gone(driver, timeout=10)
            sleep(0.3)
            if not find_visible_dialogs(driver):
                return True, text or "dismiss via Escape"
        except Exception:
            pass
        return False, text or "нет кнопок для отмены диалога"

    prefer = (
        "назад",
        "отмена",
        "отменить",
        "нет",
        "закрыть",
        "cancel",
        "close",
        "не сохр",
        "вернуться",
    )
    labels = [l or "(пусто)" for l, _ in buttons]
    logger.info("Отмена диалога — кнопки ({}): {}", len(buttons), labels[:12])

    chosen = None
    chosen_label = ""
    for want in prefer:
        for label, btn in buttons:
            low = (label or "").lower().replace("\n", " ").strip()
            if want in low:
                chosen, chosen_label = btn, label
                break
        if chosen is not None:
            break
    if chosen is None:
        chosen_label, chosen = buttons[-1]

    try:
        driver.execute_script(
            "arguments[0].scrollIntoView({block:'center'}); arguments[0].click();",
            chosen,
        )
    except Exception:
        try:
            chosen.click()
        except Exception as exc:
            logger.warning("Не удалось нажать отмену диалога: {}", exc)
            return False, text or str(exc)

    logger.info("Нажали отмену диалога «{}»", (chosen_label or "").replace("\n", " ")[:80])
    wait_busy_gone(driver, timeout=30)
    sleep(0.4)
    return True, text or f"отменено [{chosen_label}]"


def confirm_warning_dialog(driver: WebDriver) -> tuple[bool, str]:
    """Прочитать предупреждение и нажать Да (сохранить несмотря на предупреждение)."""
    text = warning_dialog_text(driver)
    if not text and not find_visible_dialogs(driver):
        return False, ""
    if not text:
        text = "(диалог без текста)"
    if _is_blocking_dialog(text):
        # пересечение / дубликат — только отмена, без сохранения
        logger.error("Блокирующее предупреждение (отмена Save): {}", text[:400])
        dismissed, _ = dismiss_warning_dialog(driver)
        if not dismissed:
            logger.warning("Не удалось закрыть блокирующий диалог")
        return False, text

    logger.info("Предупреждение после сохранения:\n{}", text[:800])

    buttons: list[tuple[str, object]] = []
    # широкий поиск кнопок модалки (MUI portal)
    xpaths = (
        "//button[contains(.,'Сохранить с предупреждением') or contains(.,'сохранить с предупреждением')]",
        "//button[contains(.,'ДЛЯ ТАЛОНА И ПАЦИЕНТА') or contains(.,'Для талона и пациента')]",
        "//button[contains(.,'ДЛЯ ТАЛОНА')]",
        "//button[contains(.,'Да') or contains(.,'ДА')]",
        "//div[@role='dialog']//button",
        "//div[contains(@class,'MuiDialog')]//button",
        "//div[contains(@class,'MuiDialogActions')]//button",
        "//div[contains(@class,'MuiPaper-root')][@role='dialog']//button",
        "/html/body/div[3]//button",
        "/html/body/div[4]//button",
        "/html/body/div[5]//button",
        "/html/body/div[3]/div[2]/div[3]/div/div[1]/button",
        "/html/body/div[4]/div[2]/div[3]/div/div[1]/button",
    )
    seen = set()
    for xp in xpaths:
        for btn in driver.find_elements(By.XPATH, xp):
            try:
                if not btn.is_displayed():
                    continue
                key = btn.id
                if key in seen:
                    continue
                seen.add(key)
            except Exception:
                continue
            buttons.append((_btn_label(driver, btn), btn))

    if not buttons:
        # dump для отладки
        try:
            dump = RESULTS_DIR / "last_warning_dialog.html"
            html = driver.execute_script(
                """
                const roots = [
                  ...document.querySelectorAll('[role=dialog], .MuiDialog-root, .MuiDialog-paper')
                ];
                return roots.map(r => r.outerHTML).join('\\n\\n---\\n\\n').slice(0, 50000);
                """
            )
            dump.write_text(html or "", encoding="utf-8")
            logger.warning("Диалог без кнопок — HTML → {}", dump)
        except Exception as exc:
            logger.warning("Диалог есть, кнопок нет; dump failed: {}", exc)
        return False, text

    labels = [l or "(пусто)" for l, _ in buttons]
    logger.info("Кнопки диалога ({}): {}", len(buttons), labels[:12])

    low_text = text.lower().replace("ё", "е")
    deny = ("назад", "отмена", "отменить", "нет", "закрыть", "cancel", "close", "не сохр", "вернуться")

    button_blob = " | ".join(l.lower().replace("\n", " ") for l, _ in buttons)
    # ВАЖНО: если есть «Сохранить с предупреждением» — всегда её (даже если текст ещё про пациента)
    if (
        "сохранить с предупреждением" in button_blob
        or "с предупреждением" in button_blob
        or "имеются предупрежд" in low_text
        or "несоответств" in low_text
    ):
        prefer = (
            "сохранить с предупреждением",
            "с предупреждением",
            "сохранить с",
            "да",
            "ok",
            "ок",
        )
    elif "пациент" in low_text and ("изменены" in low_text or "сохранить изменения" in low_text):
        prefer = (
            "для талона",
            "для талона и пациента",
            "талона и пациента",
            "и пациента",
        )
    else:
        prefer = (
            "сохранить с предупреждением",
            "с предупреждением",
            "для талона и пациента",
            "да",
            "ok",
            "ок",
            "сохранить",
            "подтверд",
            "продолж",
            "yes",
        )

    chosen = None
    chosen_label = ""
    for want in prefer:
        for label, btn in buttons:
            low = label.lower().replace("\n", " ").strip()
            if any(d in low for d in deny):
                continue
            # «для талона» ≠ «для талона и пациента»
            if want == "для талона":
                if low == "для талона" or (low.startswith("для талона") and "пациент" not in low):
                    chosen, chosen_label = btn, label
                    break
                continue
            if want == low or want in low:
                chosen, chosen_label = btn, label
                break
        if chosen is not None:
            break
    if chosen is None:
        for label, btn in buttons:
            low = label.lower().replace("\n", " ").strip()
            if low in ("да", "ok", "ок") or low.startswith("да "):
                chosen, chosen_label = btn, label
                break
    if chosen is None:
        neutrals = [
            (lab, b)
            for lab, b in buttons
            if not any(d in lab.lower() for d in deny)
        ]
        if neutrals:
            chosen_label, chosen = neutrals[-1]
        elif buttons:
            chosen_label, chosen = buttons[0]
        else:
            return False, text

    try:
        driver.execute_script(
            "arguments[0].scrollIntoView({block:'center'}); arguments[0].click();",
            chosen,
        )
        logger.info("Нажали кнопку диалога «{}»", (chosen_label or "").replace("\n", " ")[:80])
    except Exception as exc:
        try:
            chosen.click()
            logger.info("Нажали кнопку диалога (native) «{}»", (chosen_label or "")[:80])
        except Exception as exc2:
            logger.warning("Не удалось нажать кнопку диалога: {} / {}", exc, exc2)
            return False, text

    wait_busy_gone(driver, timeout=60)
    sleep(0.6)
    return True, f"{text[:200]} → [{(chosen_label or '').replace(chr(10), ' ')[:60]}]"


def handle_post_save_prompts(driver: WebDriver, *, timeout: float = 40) -> tuple[bool, str]:
    """После Save: диалоги по очереди — пациент → предупреждения (Сохранить с предупреждением).

    Если всплыло недопустимое ДН (часто после «ДЛЯ ТАЛОНА И ПАЦИЕНТА») —
    жмём НАЗАД и возвращаем ошибку, чтобы на форме поставить «Состоит» и Save снова.
    """
    deadline = time() + timeout
    confirmed_texts: list[str] = []
    last_snack = ""
    snack_before = snackbar_text(driver)
    saved_with_warning = False
    patient_choice_done = False
    patient_clicks = 0

    def _dn_error_blob(dialog_text: str = "") -> str:
        bits = [dialog_text or ""]
        try:
            panel = read_claim_error_panel(driver, click=True)
            if panel:
                bits.append(panel)
        except Exception:
            pass
        try:
            snack = snackbar_text(driver)
            if snack:
                bits.append(snack)
        except Exception:
            pass
        return " | ".join(b for b in bits if b)

    while time() < deadline:
        wait_busy_gone(driver, timeout=15)

        text = warning_dialog_text(driver)
        dialogs = find_visible_dialogs(driver)

        # до/во время диалога — недопустимое ДН → НАЗАД, правим на форме
        # НО не путать с «Имеются предупреждения» по услугам
        try:
            btn_blob0 = " | ".join(
                (_btn_label(driver, b) or "").lower().replace("\n", " ")
                for b in driver.find_elements(By.XPATH, "//button")
                if b.is_displayed()
            )
        except Exception:
            btn_blob0 = ""
        low0 = (text or "").lower().replace("ё", "е")
        service_warn0 = (
            "имеются предупрежд" in low0
            or "несоответств" in low0
            or "специальност" in low0
            or "сохранить с предупреждением" in btn_blob0
        )
        blob = _dn_error_blob(text)
        if (
            not service_warn0
            and (_is_invalid_dn_error(blob) or _is_invalid_dn_error(text or ""))
        ):
            logger.warning(
                "Недопустимое ДН в диалоге/ошибках — НАЗАД, чтобы исправить на форме: {}",
                blob[:300],
            )
            dismiss_warning_dialog(driver)
            sleep(0.3)
            t_after = (warning_dialog_text(driver) or "").lower().replace("ё", "е")
            if "имеются предупрежд" in t_after or "несоответств" in t_after:
                logger.info("После НАЗАД по ДН — предупреждения по услугам, подтвердим")
            elif find_visible_dialogs(driver) or warning_dialog_text(driver):
                dismiss_warning_dialog(driver)
                return False, blob[:500]
            else:
                return False, blob[:500]

        if dialogs or (text and (_is_confirmable_warning(text) or _is_blocking_dialog(text))):
            if text and _is_blocking_dialog(text):
                # не считать блокирующим, если уже есть «Сохранить с предупреждением»
                try:
                    bb0 = " | ".join(
                        (_btn_label(driver, b) or "").lower().replace("\n", " ")
                        for b in driver.find_elements(By.XPATH, "//button")
                        if b.is_displayed()
                    )
                except Exception:
                    bb0 = ""
                if "сохранить с предупреждением" in bb0 or "с предупреждением" in bb0:
                    logger.info("Есть «Сохранить с предупреждением» — подтверждаем, не НАЗАД")
                else:
                    logger.error("Блокирующий диалог после Save — отмена (НАЗАД): {}", text[:400])
                    dismiss_warning_dialog(driver)
                    return False, text

            # Порядок как сказал пользователь:
            # 1) «ДЛЯ ТАЛОНА» (диалог пациента)
            # 2) в следующем окне «Сохранить с предупреждением»
            # Между ними НЕ жмём НАЗАД.
            ok, wtext = confirm_warning_dialog(driver)
            if not ok:
                return False, wtext or text or "Диалог не подтверждён"
            low_w = (wtext or "").lower().replace("ё", "е")
            clicked = ""
            if "→ [" in low_w:
                clicked = low_w.rsplit("→ [", 1)[-1].rstrip("]").strip()
            if "сохранить с предупреждением" in clicked or "с предупреждением" in clicked:
                saved_with_warning = True
                patient_choice_done = True
            # После «ДЛЯ ТАЛОНА» (не «и пациента») ждём следующее окно и жмём
            # «Сохранить с предупреждением». Между ними НЕ жмём НАЗАД.
            if clicked.startswith("для талона"):
                patient_choice_done = True
                patient_clicks += 1
                logger.info(
                    "Нажали «{}» — ждём окно «Имеются предупреждения» / «Сохранить с предупреждением»…",
                    clicked[:60],
                )
                wait_until = time() + 10.0
                while time() < wait_until:
                    try:
                        bb = " | ".join(
                            (_btn_label(driver, b) or "").lower().replace("\n", " ")
                            for b in driver.find_elements(By.XPATH, "//button")
                            if b.is_displayed()
                        )
                    except Exception:
                        bb = ""
                    t_w = (warning_dialog_text(driver) or "").lower().replace("ё", "е")
                    if (
                        "сохранить с предупреждением" in bb
                        or "с предупреждением" in bb
                        or "имеются предупрежд" in t_w
                        or "несоответств" in t_w
                    ):
                        logger.info("Появилось окно предупреждений — жмём «Сохранить с предупреждением»")
                        break
                    if not find_visible_dialogs(driver) and not (warning_dialog_text(driver) or "").strip():
                        logger.info("После «ДЛЯ ТАЛОНА» диалогов нет")
                        break
                    # реальная ошибка ДН (не предупреждения по услугам) → НАЗАД
                    try:
                        panel = read_claim_error_panel(driver, click=True)
                    except Exception:
                        panel = ""
                    if (
                        panel
                        and _is_invalid_dn_error(panel)
                        and "несоответств" not in panel.lower()
                        and "имеются предупрежд" not in panel.lower()
                    ):
                        logger.warning("После «ДЛЯ ТАЛОНА» ошибка ДН — НАЗАД: {}", panel[:200])
                        dismiss_warning_dialog(driver)
                        return False, panel[:500]
                    sleep(0.3)
                try:
                    bb2 = " | ".join(
                        (_btn_label(driver, b) or "").lower().replace("\n", " ")
                        for b in driver.find_elements(By.XPATH, "//button")
                        if b.is_displayed()
                    )
                except Exception:
                    bb2 = ""
                if "сохранить с предупреждением" in bb2 or "с предупреждением" in bb2:
                    ok2, w2 = confirm_warning_dialog(driver)
                    if ok2:
                        saved_with_warning = True
                        if w2:
                            confirmed_texts.append(w2[:240])
                        sleep(0.4)
                        continue
            if wtext:
                confirmed_texts.append(wtext[:240])
            sleep(0.4)
            continue

        # диалогов больше нет
        msg = snackbar_text(driver)
        if msg and msg != snack_before:
            last_snack = msg
            if "дубликат" in msg.lower():
                return False, msg
            if _is_invalid_dn_error(msg):
                return False, msg
            if _is_save_fail_message(msg) and not saved_with_warning and not confirmed_texts:
                wait_dlg_until = time() + 3
                while time() < wait_dlg_until:
                    if find_visible_dialogs(driver) or _is_confirmable_warning(warning_dialog_text(driver)):
                        break
                    sleep(0.25)
                else:
                    # claim-error может держать текст про ДН
                    blob4 = _dn_error_blob(msg)
                    return False, blob4[:500] or msg
                continue
            if _is_save_ok_message(msg) or confirmed_texts:
                summary = " | ".join(confirmed_texts)
                if msg:
                    summary = f"{summary} | {msg}" if summary else msg
                return True, summary or msg

        if confirmed_texts:
            # успех только если диалоги закрылись
            sleep(0.8)
            if find_visible_dialogs(driver) or _is_confirmable_warning(warning_dialog_text(driver)):
                continue
            summary = " | ".join(confirmed_texts)
            if last_snack and last_snack != snack_before:
                summary = f"{summary} | {last_snack}"
            return True, summary

        sleep(0.35)

    # таймаут: если диалог предупреждений ещё висит — FAIL (нельзя считать OK)
    still = warning_dialog_text(driver)
    if find_visible_dialogs(driver) or _is_confirmable_warning(still):
        blob5 = _dn_error_blob(still or "")
        if _is_invalid_dn_error(blob5) or find_visible_dialogs(driver):
            dismiss_warning_dialog(driver)
        return False, (blob5 or still or "Диалог остался открытым после Save")[:400]
    if confirmed_texts:
        return True, " | ".join(confirmed_texts)
    if last_snack:
        if _is_save_fail_message(last_snack):
            return False, last_snack
        if _is_save_ok_message(last_snack):
            return True, last_snack
        return False, last_snack
    return False, "После Save нет snackbar/диалога — талон скорее всего не сохранён"


def _date_field_empty(driver: WebDriver, element_id: str) -> bool:
    """True только если поля нет или value пустой. Уже введённую дату НЕ считаем битой."""
    els = driver.find_elements(By.ID, element_id)
    if not els:
        return True
    actual = (els[0].get_attribute("value") or "").strip()
    # маски вида __-__-__ / .. считаем пустыми
    digits = re.sub(r"\D", "", actual)
    return not digits


def ensure_claim_dates_present(driver: WebDriver, begin: str, end: str, *, why: str = "") -> bool:
    """Если begin/end пустые — ввести. Если уже есть — НЕ трогать. True если что-то вводили."""
    need_begin = bool((begin or "").strip()) and _date_field_empty(driver, "begin-date")
    need_end = bool((end or "").strip()) and _date_field_empty(driver, "end-date")
    if not need_begin and not need_end:
        try:
            bv = (driver.find_element(By.ID, "begin-date").get_attribute("value") or "").strip()
            ev = (driver.find_element(By.ID, "end-date").get_attribute("value") or "").strip()
        except Exception:
            bv = ev = "?"
        logger.info("Даты на месте ({}), не трогаем begin={!r} end={!r}", why or "check", bv, ev)
        return False
    logger.warning(
        "Даты пустые ({}) begin_empty={} end_empty={} — вводим только пустые",
        why or "check",
        need_begin,
        need_end,
    )
    if need_begin:
        fill_date_field(driver, "begin-date", begin)
    if need_end:
        fill_date_field(driver, "end-date", end)
    return True


def ensure_claim_dates(driver: WebDriver, begin: str, end: str) -> bool:
    """Совместимость: только дозаполнение пустых."""
    return ensure_claim_dates_present(driver, begin, end, why="ensure_claim_dates")


def ensure_service_dates(driver: WebDriver, date_text: str) -> int:
    """Заполнить date-picker-services-* только там, где дата пустая. Уже заполненные не трогать."""
    if not (date_text or "").strip():
        return 0
    open_services_tab(driver)
    indices = list_service_indices(driver)
    if not indices:
        logger.warning("Нет строк услуг для проверки дат")
        return 0
    filled = 0
    skipped = 0
    for i in indices:
        try:
            els = driver.find_elements(By.ID, f"date-picker-services-{i}")
            if not els:
                continue
            actual = (els[0].get_attribute("value") or "").strip()
            if re.sub(r"\D", "", actual):
                skipped += 1
                continue
            try:
                fill_service_done(driver, i, "Да")
            except Exception:
                pass
            fill_service_date(driver, i, date_text)
            filled += 1
        except Exception as exc:
            logger.warning("Услуга[{}] ensure даты: {}", i, exc)
    logger.info(
        "Даты услуг: дозаполнено={}, уже были={}, эталон={}",
        filled,
        skipped,
        normalize_service_date(date_text),
    )
    return filled


# совместимость со старым именем
def refill_service_dates(driver: WebDriver, date_text: str) -> int:
    return ensure_service_dates(driver, date_text)


def click_claim_save(
    driver: WebDriver,
    *,
    begin: str = "",
    end: str = "",
) -> str:
    """Нажать Сохранить талона на текущей вкладке (услуги/назначения). На main-tab НЕ уходим."""
    dismiss_overlays(driver)
    # Даты НЕ трогаем. На карту (main-tab) НЕ переключаемся — Save есть на услугах.
    _ = (begin, end)

    def _collect_save_btns() -> list[tuple[str, object]]:
        found: list[tuple[str, object]] = []
        # id-save — кнопка сохранения талона (в т.ч. на вкладке Услуги)
        for el in driver.find_elements(By.ID, "id-save"):
            if _visible_enabled(el):
                found.append(("id-save", el))
        for el in driver.find_elements(By.ID, "save-button"):
            if _visible_enabled(el):
                # save-button на карте пациента — пропускаем по подписи
                label = _btn_label(driver, el).lower()
                if any(x in label for x in ("пациент", "карт", "соц")):
                    continue
                found.append(("save-button", el))
        for el in driver.find_elements(
            By.XPATH,
            "//button[contains(.,'Сохранить') or contains(.,'СОХРАНИТЬ')]",
        ):
            if _visible_enabled(el):
                label = _btn_label(driver, el).lower()
                if any(x in label for x in ("пациент", "карт", "соц")):
                    continue
                found.append(("text-Сохранить", el))
        return found

    candidates = _collect_save_btns()
    if not candidates:
        # если были на назначениях и кнопки нет — вернуться на Услуги (не на карту)
        if driver.find_elements(By.ID, "services-tab"):
            try:
                click_id(driver, "services-tab")
                wait_busy_gone(driver, timeout=15)
                sleep(0.25)
            except Exception as exc:
                logger.debug("services-tab перед Save: {}", exc)
        candidates = _collect_save_btns()

    if not candidates:
        try:
            _wait(driver, 8).until(
                lambda d: any(_visible_enabled(e) for e in d.find_elements(By.ID, "id-save"))
                or any(_visible_enabled(e) for e in d.find_elements(By.ID, "save-button"))
            )
        except TimeoutException:
            pass
        candidates = _collect_save_btns()

    if candidates:
        how, el = candidates[0]
        logger.info("Сохранение талона: клик {} (вариантов {}), без перехода на main-tab", how, len(candidates))
        _click_el(driver, el)
        return how

    errs = collect_form_errors(driver)
    logger.info("Сохранение талона: F2 (кнопка не найдена). errors={}", errs)
    try:
        body = driver.find_element(By.TAG_NAME, "body")
        body.send_keys(Keys.F2)
        return "F2"
    except Exception as exc:
        raise RuntimeError(f"Кнопка сохранения не найдена и F2 не сработал: {exc}; errors={errs!r}") from exc


def save_claim(
    driver: WebDriver,
    *,
    begin: str = "",
    end: str = "",
) -> tuple[bool, str]:
    """Сохранить талон: Save → прочитать предупреждение → нажать Да (как в талоны2)."""
    enp_before = ""
    try:
        enp_before = (driver.find_element(By.ID, "enp").get_attribute("value") or "").strip()
    except Exception:
        pass

    pre_errs = collect_form_errors(driver)
    if pre_errs:
        logger.warning("Перед Save ошибки формы: {}", pre_errs)

    try:
        click_claim_save(driver, begin=begin, end=end)
    except RuntimeError as exc:
        return False, str(exc)

    wait_busy_gone(driver, timeout=60)
    sleep(0.4)

    ok, msg = handle_post_save_prompts(driver, timeout=30)
    if ok:
        return True, msg

    # форма сбросилась — часто так выглядит успех после подтверждения
    try:
        enp_now = ""
        if driver.find_elements(By.ID, "enp"):
            enp_now = (driver.find_element(By.ID, "enp").get_attribute("value") or "").strip()
        if enp_before and enp_now == "":
            return True, msg or "форма сброшена после сохранения"
    except Exception:
        pass

    post_errs = collect_form_errors(driver)
    panel = read_claim_error_panel(driver, click=True)
    if panel and panel not in (post_errs or ""):
        post_errs = f"{post_errs} | {panel}".strip(" |") if post_errs else panel
    snack = snackbar_text(driver)
    dlg = warning_dialog_text(driver)
    extra_bits = [b for b in (msg, snack, dlg, post_errs) if b]
    # dump кнопок на странице для отладки
    try:
        btns = driver.execute_script(
            """
            return [...document.querySelectorAll('button')]
              .filter(b => b.offsetParent !== null)
              .map(b => ((b.id||'') + ':' + (b.innerText||'').trim()).slice(0,80))
              .slice(0, 25);
            """
        )
        logger.warning("После Save visible buttons: {}", btns)
    except Exception:
        pass
    if post_errs:
        logger.warning("После Save ошибки формы: {}", post_errs)
    if extra_bits:
        return False, " | ".join(extra_bits)[:500]
    return False, "После Save нет snackbar/диалога — талон скорее всего не сохранён"


def _is_save_fail_message(msg: str) -> bool:
    low = (msg or "").lower()
    if not low:
        return False
    fail_markers = (
        "ошибка",
        "обязательн",
        "необходимо",
        "не найден",
        "заполнить",
        "заполните",
        "по маршруту",
        "не должна",
        "не должен",
        "больше текущ",
        "дубликат",
        "не сохран",
        "отсутствует",
        "изменены",
        "онкология",
    )
    return any(x in low for x in fail_markers)


def _is_save_ok_message(msg: str) -> bool:
    low = (msg or "").lower()
    return any(x in low for x in ("сохран", "успеш", "создан", "записан", "добавлен"))


def load_talons(path: Path) -> list[TalonRow]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["талоны"] if "талоны" in wb.sheetnames else wb.active
    rows_iter = ws.iter_rows(values_only=True)
    header = [_cell(h).lower() for h in next(rows_iter)]
    idx = {h: i for i, h in enumerate(header) if h}

    def col(*names: str) -> int | None:
        for n in names:
            if n.lower() in idx:
                return idx[n.lower()]
        return None

    i_enp = col("енп", "enp")
    i_beg = col("начало", "begin", "дата начала", "begin_date")
    i_end = col("окончание", "end", "дата окончания", "end_date")
    i_type = col("тип", "type", "exam_type", "цель")
    i_doc = col("врач", "doctor", "код врача")
    i_res = col("результат", "result", "request_result")
    i_per = col("период", "period", "next_month", "nextmonth")
    i_dx = col("диагноз", "diagnosis", "ds")
    i_dn = col("дн", "дисп. набл", "dispensary", "is_dn")
    i_bldg = col("корпус", "building", "subdivision")
    i_place = col("место", "место обращения", "place", "service_place")
    i_done = col("выполнено", "done", "failure")
    i_purp = col("наз_вид", "вид назначения", "purpose", "purpose_type", "purposetype")
    i_from = col("наз_от_кого", "от кого", "speciality", "specialityid", "purpose_from")
    i_to = col("наз_к_кому", "к кому", "reference_speciality", "referencespecialityid", "purpose_to")
    i_soc = col("соцстатус", "социальный статус", "social_status", "socialstatus", "статус")
    i_occ = col("занятость", "вид занятости", "occupation", "занят")

    # врач необязателен — можно дозаполнить вручную в Excel
    if i_enp is None or i_beg is None or i_end is None or i_type is None:
        raise ValueError(
            "В Excel нужны колонки: ЕНП, начало, окончание, тип. "
            f"Сейчас: {header}"
        )

    out: list[TalonRow] = []
    for n, raw in enumerate(rows_iter, start=2):
        if not raw or all(v is None or str(v).strip() == "" for v in raw):
            continue
        enp = _cell(raw[i_enp]).replace("`", "").replace(" ", "")
        if not enp or not re.sub(r"\D", "", enp):
            continue
        # пропускаем строки-подсказки / мусор без цифрового ЕНП
        if len(re.sub(r"\D", "", enp)) < 11:
            continue
        extras: list[tuple[str, str]] = []
        for n_extra in range(2, 8):
            i_dx_n = col(f"диагноз{n_extra}", f"diagnosis{n_extra}", f"ds{n_extra}")
            i_dn_n = col(f"дн{n_extra}", f"dn{n_extra}", f"is_dn{n_extra}")
            if i_dx_n is None:
                continue
            dx_n = normalize_mkb(_cell(raw[i_dx_n]))
            if not dx_n:
                continue
            dn_n = normalize_dn(_cell(raw[i_dn_n])) if i_dn_n is not None else ""
            extras.append((dx_n, dn_n))
        out.append(
            TalonRow(
                enp=enp,
                begin=normalize_service_date(_cell(raw[i_beg])),
                end=normalize_service_date(_cell(raw[i_end])),
                exam_type=_norm_type(_cell(raw[i_type])),
                doctor=_cell(raw[i_doc]) if i_doc is not None else "",
                result=_cell(raw[i_res]) if i_res is not None else "",
                period=_cell(raw[i_per]) if i_per is not None else DEFAULT_PERIOD,
                diagnosis=normalize_mkb(_cell(raw[i_dx])) if i_dx is not None else "",
                dn=normalize_dn(_cell(raw[i_dn])) if i_dn is not None else "",
                building=_cell(raw[i_bldg]) if i_bldg is not None else "",
                place=_cell(raw[i_place]) if i_place is not None else DEFAULT_PLACE,
                done_value=normalize_service_done(_cell(raw[i_done]) if i_done is not None else DEFAULT_DONE),
                purpose_type=_cell(raw[i_purp]) if i_purp is not None else "",
                purpose_from=_cell(raw[i_from]) if i_from is not None else "",
                purpose_to=_cell(raw[i_to]) if i_to is not None else "",
                social_status=_cell(raw[i_soc]) if i_soc is not None else "",
                occupation=_cell(raw[i_occ]) if i_occ is not None else "",
                extra_diagnoses=extras,
                raw_index=n,
            )
        )
    return out


def open_exam_form(driver: WebDriver, exam_type: str) -> None:
    url = exam_url(exam_type)
    logger.info("Открываем {}", url)
    driver.get(url)
    wait_busy_gone(driver)
    dismiss_overlays(driver)
    _wait(driver, 15).until(EC.presence_of_element_located((By.ID, "enp")))
    sleep(0.4)


def dismiss_overlays(driver: WebDriver) -> None:
    """Закрыть модалки / snackbar, мешающие клику по полям."""
    for xp in (
        "//div[@role='dialog']//button[contains(.,'OK') or contains(.,'Ок') or contains(.,'Закрыть') or contains(.,'Да')]",
        "//button[@aria-label='Close' or @aria-label='close']",
    ):
        for btn in driver.find_elements(By.XPATH, xp):
            try:
                btn.click()
                sleep(0.2)
            except Exception:
                continue
    wait_busy_gone(driver, timeout=10)


def fill_date_field(driver: WebDriver, element_id: str, raw: str) -> None:
    """Поля begin-date / end-date: маска DD-MM-YY (как у услуг)."""
    text = normalize_service_date(raw)  # DD-MM-YY
    if not text:
        return
    digits = re.sub(r"\D", "", text)
    want_yy = text[-2:]
    dismiss_overlays(driver)
    try:
        _wait(driver, 8).until(EC.element_to_be_clickable((By.ID, element_id)))
    except TimeoutException:
        pass
    el = driver.find_element(By.ID, element_id)
    driver.execute_script("arguments[0].scrollIntoView({block:'center'});", el)
    try:
        el.click()
    except Exception:
        dismiss_overlays(driver)
        try:
            el.click()
        except Exception:
            driver.execute_script("arguments[0].focus();", el)
    sleep(0.1)
    el.send_keys(Keys.CONTROL, "a")
    el.send_keys(Keys.BACKSPACE)
    # посимвольно — маска начала/окончания иначе глотает ввод
    for ch in digits:
        el.send_keys(ch)
        sleep(0.04)
    _set_input_value(driver, el, text)
    el.send_keys(Keys.TAB)
    sleep(0.2)
    actual = (el.get_attribute("value") or "").strip().replace(".", "-")
    if not actual.endswith(want_yy):
        logger.warning("{} дата «{}» (нужно {}) → повтор", element_id, actual or "пусто", text)
        try:
            el.click()
        except Exception:
            driver.execute_script("arguments[0].focus();", el)
        el.send_keys(Keys.CONTROL, "a")
        el.send_keys(Keys.BACKSPACE)
        for ch in digits:
            el.send_keys(ch)
            sleep(0.04)
        _set_input_value(driver, el, text)
        el.send_keys(Keys.TAB)
        sleep(0.2)
        actual = (el.get_attribute("value") or "").strip()
    logger.info("{} = {}", element_id, actual)


def fill_patient_and_search(driver: WebDriver, row: TalonRow) -> None:
    """Даты начала/окончания → ЕНП → поиск. По датам ОМС подбирает услуги."""
    dismiss_overlays(driver)
    fill_date_field(driver, "begin-date", row.begin)
    fill_date_field(driver, "end-date", row.end)
    # один раз проверяем, что даты реально в полях (если пусто — добить сейчас, не на Save)
    for eid, raw in (("begin-date", row.begin), ("end-date", row.end)):
        if not (raw or "").strip():
            continue
        if _date_field_empty(driver, eid):
            logger.warning("{} пусто после первого ввода — повторяем СЕЙЧАС (не на Save)", eid)
            fill_date_field(driver, eid, raw)
        else:
            val = (driver.find_element(By.ID, eid).get_attribute("value") or "").strip()
            logger.info("{} на месте: {!r} — дальше не трогаем", eid, val)
    click_id(driver, "enp")
    input_id(driver, "enp", row.enp)
    # кнопка поиска рядом с ЕНП (как в ambulatory)
    clicked = False
    for xpath in (
        '//*[@id="enp"]/ancestor::div[contains(@class,"jss")][1]/following-sibling::div//button',
        '//button[.//span[contains(.,"Поиск")] or contains(.,"Найти")]',
        '//*[@id="enp"]/../../..//button',
    ):
        btns = driver.find_elements(By.XPATH, xpath)
        if btns:
            try:
                btns[0].click()
                clicked = True
                break
            except Exception:
                driver.execute_script("arguments[0].click();", btns[0])
                clicked = True
                break
    if not clicked:
        # fallback: Enter в поле ЕНП
        driver.find_element(By.ID, "enp").send_keys(Keys.ENTER)
    wait_busy_gone(driver, timeout=60)
    sleep(0.8)
    err = snackbar_text(driver)
    low = (err or "").lower()
    # «Маршруты не найдены» — предупреждение, не блокирует талон
    if err and "маршрут" not in low and any(
        x in low for x in ("пациент не найден", "некорректн", "ошибка поиска")
    ):
        raise RuntimeError(f"Поиск пациента: {err}")
    if err:
        logger.info("После поиска snackbar: {}", err)
    # если выпадающий список пациентов — выбрать первого
    for xp in (
        "//div[contains(@class,'Select-menu')]//div[contains(@class,'option')]",
        "//div[@role='option']",
    ):
        opts = driver.find_elements(By.XPATH, xp)
        if opts:
            try:
                opts[0].click()
                wait_busy_gone(driver)
                sleep(0.4)
                break
            except Exception:
                continue
    # после успешного поиска поле врача обычно становится доступно
    try:
        _wait(driver, 20).until(EC.element_to_be_clickable((By.ID, "doctor")))
    except TimeoutException as exc:
        raise RuntimeError(
            f"После поиска ЕНП поле doctor не готово. snackbar={err!r}"
        ) from exc


def fill_main_page(
    driver: WebDriver,
    row: TalonRow,
    dn_labels: dict[str, str] | None = None,
) -> None:
    """Первая вкладка medicalExamination."""
    if (row.doctor or "").strip():
        input_enter_id(driver, "doctor", row.doctor)
        sleep(0.5)
        wait_busy_gone(driver)
    else:
        logger.info("Код врача пуст — пропускаем #doctor (заполните вручную при необходимости)")

    # если поля соцстатуса есть на карте талона — дублируем
    if (row.social_status or "").strip():
        try_input_enter_id(driver, "socialStatus", row.social_status.strip())
    if (row.occupation or "").strip():
        try_input_enter_id(driver, "occupation", row.occupation.strip())

    # Результат обращения
    if row.result:
        if not try_input_enter_id(driver, "request-result", row.result):
            logger.warning("request-result не найден / не заполнен ({})", row.result)
        sleep(0.35)
        wait_busy_gone(driver)

    # #medicalExaminationPlace («Место прохождения диспансеризации») — НИКОГДА не заполняем.
    # В Excel колонки под это поле нет; вмешательство ломает форму.

    period = row.period or DEFAULT_PERIOD
    if not try_input_enter_id(driver, "nextMonth", period):
        logger.warning("Поле nextMonth не найдено")

    # Диагнозы + ДН: основной (0). Сопутствующие — только если строки уже есть в форме
    # (кнопку #add-diagnosis не жмём — лишние строки не создаём).
    dn_map = dn_labels if dn_labels is not None else load_dn_labels()
    dx_pairs: list[tuple[str, str]] = []
    if row.diagnosis:
        dx_pairs.append((row.diagnosis, dn_for_exam_type(row.exam_type, row.dn or "")))
    for dx, dn in row.extra_diagnoses:
        if dx:
            dx_pairs.append((dx, dn_for_exam_type(row.exam_type, dn or "")))
    for i, (dx, dn) in enumerate(dx_pairs):
        if i > 0 and not driver.find_elements(By.ID, f"diagnosis-{i}"):
            logger.info(
                "Пропуск сопутствующего диагноза[{}] {} — строки нет (add-diagnosis не используем)",
                i,
                dx,
            )
            continue
        try_input_enter_id(driver, f"diagnosis-{i}", dx)
        sleep(0.2)
        if dn:
            if not fill_dn_field(driver, i, dn, dn_map):
                logger.warning("ДН[{}] не заполнен ({})", i, dn)
            sleep(0.15)


def ensure_diagnosis_slot(driver: WebDriver, index: int) -> None:
    """Устарело: строки диагнозов через #add-diagnosis не создаём."""
    return



def _birth_year(value: str) -> int | None:
    text = (value or "").strip()
    if not text:
        return None
    m = re.search(r"(\d{1,2})[-./](\d{1,2})[-./](\d{2,4})", text)
    if not m:
        m4 = re.search(r"(19|20)\d{2}", text)
        return int(m4.group(0)) if m4 else None
    y = m.group(3)
    if len(y) == 2:
        y = f"20{y}" if int(y) <= 30 else f"19{y}"
    return int(y)


def social_occupation_by_age(age: int) -> tuple[str, str]:
    """18 < age < 60 → 11/1; иначе (в т.ч. ≥60) → 22/3."""
    if 18 < age < 60:
        return "11", "1"
    return "22", "3"


def _read_birth_from_main_tab(driver: WebDriver) -> str:
    """Дата рождения из таблицы на main-tab (без открытия карты)."""
    for tr in driver.find_elements(By.CSS_SELECTOR, "table tbody tr"):
        cells = [((td.text or "").strip()) for td in tr.find_elements(By.TAG_NAME, "td")]
        blob = " ".join(cells)
        if "рожден" in blob.lower():
            m = re.search(r"(\d{1,2}[-./]\d{1,2}[-./]\d{2,4})", blob)
            if m:
                return m.group(1)
    return ""


def ensure_patient_social_status(driver: WebDriver, row: TalonRow | None = None) -> tuple[str, str]:
    """
    Кнопка карты пациента → socialStatus / occupation по году рождения.
    Сохранение пациента: #save-button.
    НЕ кликаем main-tab «на всякий случай» — это сбрасывает begin/end даты лечения.
    """
    if row and (row.social_status or "").strip() and (row.occupation or "").strip():
        soc, occ = row.social_status.strip(), row.occupation.strip()
        logger.info("Соцстатус/занятость из Excel: {} / {} — всё равно пишем в карту пациента", soc, occ)
    else:
        soc = occ = ""

    dismiss_overlays(driver)
    # main-tab только если кнопки пациента ещё не видно (без лишнего remount формы)
    birth_raw = _read_birth_from_main_tab(driver)
    patient_btns = [
        e
        for e in driver.find_elements(By.XPATH, PATIENT_CARD_BTN_XPATH)
        if e.is_displayed()
    ]
    if not patient_btns and driver.find_elements(By.ID, "main-tab"):
        logger.info("Кнопки пациента нет на экране — один клик main-tab")
        click_id(driver, "main-tab")
        wait_busy_gone(driver)
        sleep(0.4)
        birth_raw = birth_raw or _read_birth_from_main_tab(driver)
        try:
            _wait(driver, 8).until(
                lambda d: any(
                    e.is_displayed()
                    for e in d.find_elements(By.XPATH, PATIENT_CARD_BTN_XPATH)
                )
                or any(
                    e.is_displayed()
                    for e in d.find_elements(By.XPATH, "//button[contains(.,'Редактировать')]")
                )
                or any(
                    e.is_displayed()
                    for e in d.find_elements(By.XPATH, "//table//tbody//button")
                )
            )
        except TimeoutException:
            pass
        sleep(0.3)

    if (not soc or not occ) and birth_raw:
        year = _birth_year(birth_raw)
        if year is not None:
            age = datetime.now().year - year
            soc, occ = social_occupation_by_age(age)
            logger.info(
                "ДР на main-tab={} → год={}, возраст≈{} → socialStatus={}, occupation={}",
                birth_raw,
                year,
                age,
                soc,
                occ,
            )

    btn = None
    for xp in (
        "//button[contains(.,'Редактировать')]",
        PATIENT_CARD_BTN_XPATH,
        "//table//tbody/tr[5]/td//button",
        "//table//tbody/tr[6]/td//button",
        "//table//tbody/tr[.//button][last()]//button",
        "//table//tbody//button",
        "//button[contains(.,'пациент') or contains(.,'Пациент')]",
    ):
        els = [e for e in driver.find_elements(By.XPATH, xp) if e.is_displayed()]
        if els:
            btn = els[0]
            break
    if btn is None:
        # если ДР уже дала соцстатус — заполним на карте талона без окна пациента
        if soc and occ and row is not None:
            row.social_status = soc
            row.occupation = occ
            logger.warning(
                "Кнопка карты пациента не найдена — ставим social/occupation на форме талона ({}/{})",
                soc,
                occ,
            )
            return soc, occ
        raise RuntimeError("Кнопка карты пациента на main-tab не найдена")

    try:
        btn.click()
    except Exception:
        driver.execute_script("arguments[0].click();", btn)
    wait_busy_gone(driver, timeout=30)
    sleep(0.6)

    try:
        _wait(driver, 15).until(EC.presence_of_element_located((By.ID, "socialStatus")))
    except TimeoutException as exc:
        raise RuntimeError("Окно пациента не открылось (нет #socialStatus)") from exc

    if not soc or not occ:
        if driver.find_elements(By.ID, "birth-date"):
            birth_raw = (
                driver.find_element(By.ID, "birth-date").get_attribute("value") or birth_raw
            )
        year = _birth_year(birth_raw)
        if year is None:
            raise RuntimeError(f"Не разобрать дату рождения: {birth_raw!r}")
        age = datetime.now().year - year
        soc, occ = social_occupation_by_age(age)
        logger.info(
            "Пациент ДР={} → год={}, возраст≈{} → socialStatus={}, occupation={}",
            birth_raw,
            year,
            age,
            soc,
            occ,
        )

    def _fill_select(element_id: str, value: str) -> None:
        el = driver.find_element(By.ID, element_id)
        fill_react_select_input(driver, el, value)
        sleep(0.25)
        wait_busy_gone(driver)
        wrap = _closest_react_select(el)
        shown = (wrap.text if wrap is not None else "") or ""
        if str(value) not in shown:
            # повтор: только код + Enter
            el.send_keys(Keys.CONTROL, "a")
            el.send_keys(Keys.BACKSPACE)
            el.send_keys(str(value), Keys.ENTER)
            sleep(0.5)
            wrap = _closest_react_select(el)
            shown = (wrap.text if wrap is not None else "") or ""
        if str(value) not in shown:
            raise RuntimeError(f"{element_id}: не выбрано «{value}», сейчас «{shown[:80]}»")
        logger.info("{} → {}", element_id, shown.split("\n")[0][:80])

    _fill_select("socialStatus", soc)
    _fill_select("occupation", occ)
    logger.info("В карте пациента выставлено: socialStatus={}, occupation={}", soc, occ)

    # только #save-button карты пациента (не #id-save талона)
    saves = [e for e in driver.find_elements(By.ID, "save-button") if e.is_displayed()]
    if not saves:
        raise RuntimeError("save-button карты пациента не найден")
    try:
        saves[0].click()
    except Exception:
        driver.execute_script("arguments[0].click();", saves[0])
    wait_busy_gone(driver, timeout=60)
    sleep(0.5)
    # если ОМС показал предупреждение — читаем и подтверждаем (Да)
    dlg = warning_dialog_text(driver)
    if dlg or find_visible_dialogs(driver):
        ok, info = handle_post_save_prompts(driver, timeout=20)
        if not ok and _is_blocking_dialog(info):
            raise RuntimeError(f"Карта пациента не сохранена: {info}")
        if info:
            logger.info("После сохранения пациента (диалог): {}", info[:240])
    else:
        msg = snackbar_text(driver)
        if msg:
            logger.info("После сохранения пациента: {}", msg)
        low = (msg or "").lower()
        if any(x in low for x in ("ошибка", "обязательн", "не указан")):
            raise RuntimeError(f"Карта пациента не сохранена: {msg}")

    url = driver.current_url or ""
    if "medicalExamination" not in url:
        raise RuntimeError(f"После save карты ушли со страницы талона: {url}")

    # НЕ жмём main-tab после save — remount сбрасывает begin/end.
    # Остаёмся на форме талона как есть после закрытия карты пациента.

    if row is not None:
        row.social_status = soc
        row.occupation = occ
    return soc, occ


def _is_social_status_error(msg: str) -> bool:
    low = (msg or "").lower()
    return any(
        x in low
        for x in (
            "социальн",
            "соцстатус",
            "занятост",
            "вид занятости",
            "не указан социальн",
            "отсутствует вид занятости",
        )
    )


def _is_invalid_dn_error(msg: str) -> bool:
    """«Недопустимое значение диспансерного наблюдения при диагнозе …» → ставим Состоит."""
    low = (msg or "").lower().replace("ё", "е")
    if not low:
        return False
    if "недопустим" in low and ("диспансерн" in low or "наблюден" in low):
        return True
    if "значен" in low and "диспансерн" in low:
        return True
    if "диспансерн" in low and "диагноз" in low and any(
        x in low for x in ("недопустим", "неверн", "ошибк", "не соответ", "правим дн")
    ):
        return True
    # диалог пациента завис / НАЗАД — тоже чиним ДН на «Состоит»
    if "правим дн на форме" in low or "диалог пациента" in low and "назад" in low:
        return True
    return False


def _needs_dn_consists_fix(msg: str) -> bool:
    """После Save нужно НАЗАД→Состоит→Save: явная ошибка ДН или зависший диалог пациента."""
    if _is_invalid_dn_error(msg):
        return True
    low = (msg or "").lower().replace("ё", "е")
    if "данные пациента были изменены" in low and any(
        x in low for x in ("error", "ошибк", "назад", "не закрыл", "claim-error")
    ):
        return True
    return False


def apply_dn_code(
    driver: WebDriver,
    row: TalonRow,
    dn_code: str,
    dn_labels: dict[str, str] | None,
    *,
    why: str = "",
) -> None:
    """Выставить ДН=dn_code на основной (+ сопутствующие с тем же кодом, если были 1/3/пусто)."""
    code = normalize_dn(dn_code) or str(dn_code).strip()
    row.dn = code
    new_extras: list[tuple[str, str]] = []
    for dx, dn in row.extra_diagnoses:
        prev = normalize_dn(dn) or (dn or "").strip()
        if dx and prev in ("", "1", "3", "7"):
            new_extras.append((dx, code))
        else:
            new_extras.append((dx, dn))
    row.extra_diagnoses = new_extras
    if not driver.find_elements(By.ID, "is-dispensary-observation-0"):
        if driver.find_elements(By.ID, "main-tab"):
            try:
                click_id(driver, "main-tab")
                wait_busy_gone(driver, timeout=15)
                sleep(0.25)
            except Exception:
                pass
        ensure_claim_dates_present(driver, row.begin, row.end, why=f"после main для ДН={code}")
    dn_map = dn_labels if dn_labels is not None else load_dn_labels()
    logger.warning("Ставим ДН={} ({}) на диагнозы", code, why or "fix")
    if not fill_dn_field(driver, 0, code, dn_map):
        logger.warning("ДН[0]={} не заполнен", code)
    for i, (dx, dn) in enumerate(row.extra_diagnoses, start=1):
        if not dx or not driver.find_elements(By.ID, f"is-dispensary-observation-{i}"):
            continue
        if (normalize_dn(dn) or dn) == code:
            if not fill_dn_field(driver, i, code, dn_map):
                logger.warning("ДН[{}]={} не заполнен", i, code)


def dn_for_exam_type(exam_type: str, dn: str) -> str:
    """ОПВ: «Состоит»(1) → «Состоит, проведено»(7). ДВ4 без изменений."""
    code = normalize_dn(dn) or (dn or "").strip()
    if (exam_type or "").strip().lower() == "opv" and code == "1":
        return "7"
    return code or (dn or "").strip()


def _is_opv_dn_consists_error(msg: str) -> bool:
    """ОМС: пациент уже состоит на ДН по диагнозу — для ОПВ нужен код 7."""
    low = (msg or "").lower()
    if not low:
        return False
    has_consist = any(
        x in low
        for x in (
            "состоит",
            "состоит на",
            "диспансерн",
            "уже взят",
            "взят на дн",
            "находится на дн",
            "группа дн",
        )
    )
    has_dx = any(x in low for x in ("диагноз", "мкб", "дн", "наблюден"))
    # частые формулировки ОМС / claim-error
    if "состоит" in low and ("диагноз" in low or "дн" in low or "наблюден" in low):
        return True
    if has_consist and has_dx:
        return True
    if "не подлежит" in low and "состоит" in low:
        return True
    return False


def apply_opv_dn7(
    driver: WebDriver,
    row: TalonRow,
    dn_labels: dict[str, str] | None,
) -> None:
    """Поставить ДН=7 на основной и сопутствующие (где был 1/3/пусто с диагнозом)."""
    row.dn = "7"
    new_extras: list[tuple[str, str]] = []
    for dx, dn in row.extra_diagnoses:
        code = normalize_dn(dn) or (dn or "").strip()
        if dx and code in ("", "1", "3"):
            new_extras.append((dx, "7"))
        else:
            new_extras.append((dx, dn))
    row.extra_diagnoses = new_extras
    # на main только если поля ДН нет на текущем экране — main-tab сбрасывает даты
    if not driver.find_elements(By.ID, "is-dispensary-observation-0"):
        if driver.find_elements(By.ID, "main-tab"):
            try:
                click_id(driver, "main-tab")
                wait_busy_gone(driver, timeout=15)
                sleep(0.25)
            except Exception:
                pass
        ensure_claim_dates_present(driver, row.begin, row.end, why="после main для ДН=7")
    dn_map = dn_labels if dn_labels is not None else load_dn_labels()
    # основной
    if not fill_dn_field(driver, 0, "7", dn_map):
        logger.warning("ОПВ: не удалось выставить ДН[0]=7")
    for i, (dx, dn) in enumerate(row.extra_diagnoses, start=1):
        if not dx or not driver.find_elements(By.ID, f"is-dispensary-observation-{i}"):
            continue
        if (normalize_dn(dn) or dn) == "7":
            if not fill_dn_field(driver, i, "7", dn_map):
                logger.warning("ОПВ: не удалось выставить ДН[{}]=7", i)


def fill_dn_field(
    driver: WebDriver,
    index: int,
    dn: str,
    dn_labels: dict[str, str] | None = None,
) -> bool:
    """Заполнить is-dispensary-observation-{i} по справочнику ДН (код → подпись)."""
    variants = dn_search_texts(dn, dn_labels)
    if not variants:
        return False

    els = driver.find_elements(By.ID, f"is-dispensary-observation-{index}")
    targets = list(els)
    if index == 0 and not targets:
        for xp in (
            "//div[contains(.,'Диспансерное наблюдение')]"
            "//input[contains(@id,'react-select') or @type='text']",
            "//label[contains(.,'Диспансерное')]/following::input[1]",
        ):
            targets.extend(driver.find_elements(By.XPATH, xp))
    if not targets:
        logger.warning("Поле ДН[{}] не найдено", index)
        return False

    el = targets[0]
    last_exc: Exception | None = None
    for text in variants:
        try:
            fill_react_select_input(driver, el, text)
            logger.info("ДН[{}] подпись «{}» (из кода {})", index, text, normalize_dn(dn) or dn)
            return True
        except Exception as exc:
            last_exc = exc
            continue
    if last_exc:
        logger.debug("ДН[{}] ошибки вариантов: {}", index, last_exc)
    return False


def _try_fill_dn_fallback(driver: WebDriver, dn: str) -> None:
    """Совместимость: ДН через xpath, если id не найден."""
    fill_dn_field(driver, 0, dn)


def _tab_is_enabled(driver: WebDriver, tab_id: str) -> bool:
    els = driver.find_elements(By.ID, tab_id)
    if not els:
        return False
    el = els[0]
    if el.get_attribute("disabled") is not None:
        return False
    return el.is_enabled()


def wait_tab_enabled(driver: WebDriver, tab_id: str, timeout: float = 45) -> None:
    try:
        _wait(driver, timeout).until(lambda d: _tab_is_enabled(d, tab_id))
    except TimeoutException as exc:
        raise RuntimeError(
            f"Вкладка {tab_id} не активировалась за {timeout}s "
            "(обычно нужно успешно сохранить карту)."
        ) from exc


def ensure_default_purposes(row: TalonRow) -> bool:
    """Если наз_вид пуст, но результат обычно требует назначений — подставить 1/95/95."""
    if (row.purpose_type or "").strip():
        return False
    result = (row.result or "").strip()
    # 31/32/33/34 — III группа; 12 тоже часто тянет назначения на ДВ4
    if result not in {"12", "31", "32", "33", "34"}:
        return False
    row.purpose_type = DEFAULT_PURPOSE_TYPE
    row.purpose_from = row.purpose_from or DEFAULT_PURPOSE_FROM
    row.purpose_to = row.purpose_to or DEFAULT_PURPOSE_TO
    logger.info(
        "Авто-назначение (результат {}): вид={}, от={}, к={}",
        result,
        row.purpose_type,
        row.purpose_from,
        row.purpose_to,
    )
    return True


def purposes_required_message(msg: str) -> bool:
    low = (msg or "").lower()
    return "назнач" in low and ("необходим" in low or "требуется" in low or "нужно" in low)


def fill_purposes(driver: WebDriver, row: TalonRow) -> None:
    """Блок назначений: purposeType-0 / specialityId / referenceSpecialityId.

    Вкладка purposes-tab активна, когда результат/правила ОМС требуют назначений
    (в Excel — колонка наз_вид).
    """
    if not (row.purpose_type or "").strip():
        return

    logger.info(
        "Назначение: вид={}, от={}, к={}",
        row.purpose_type,
        row.purpose_from or "—",
        row.purpose_to or "—",
    )
    if not _tab_is_enabled(driver, "purposes-tab"):
        sleep(1.0)
        wait_busy_gone(driver)
    wait_tab_enabled(driver, "purposes-tab", timeout=30)
    click_id(driver, "purposes-tab")
    wait_busy_gone(driver)
    sleep(0.5)

    # Сначала «+» / add-purposes — иначе purposeType-0 ещё нет в DOM
    if not driver.find_elements(By.ID, "purposeType-0"):
        if driver.find_elements(By.ID, "add-purposes"):
            click_id(driver, "add-purposes")
            wait_busy_gone(driver)
            sleep(0.6)
        else:
            btns = driver.find_elements(
                By.XPATH,
                "//button[@aria-label='add' or contains(.,'Добавить')]",
            )
            if btns:
                btns[0].click()
                wait_busy_gone(driver)
                sleep(0.6)

    try:
        _wait(driver, 10).until(EC.presence_of_element_located((By.ID, "purposeType-0")))
    except TimeoutException as exc:
        raise RuntimeError(
            "purposes-tab открыт, но поле purposeType-0 не найдено "
            "(после add-purposes)."
        ) from exc

    if not try_input_enter_id(driver, "purposeType-0", row.purpose_type):
        raise RuntimeError(f"Не удалось заполнить purposeType-0 = {row.purpose_type!r}")
    sleep(0.25)

    if row.purpose_from:
        if not try_input_enter_id(driver, "specialityId", row.purpose_from):
            raise RuntimeError(f"Не удалось заполнить specialityId (от кого) = {row.purpose_from!r}")
        sleep(0.2)

    if row.purpose_to:
        if not try_input_enter_id(driver, "referenceSpecialityId", row.purpose_to):
            raise RuntimeError(
                f"Не удалось заполнить referenceSpecialityId (к кому) = {row.purpose_to!r}"
            )
        sleep(0.2)

    # Дата направления * = окончание талона → #referenceDate
    purpose_date = (row.end or "").strip()
    if not purpose_date:
        logger.warning("Нет даты окончания — дата направления не заполнена")
    else:
        text = normalize_service_date(purpose_date)
        logger.info("Дата направления (из окончания) = {}", text)
        filled_date = False
        for pid in ("referenceDate", "reference-date"):
            els = driver.find_elements(By.ID, pid)
            if not els:
                continue
            try:
                # клик по контейнеру — иначе input «not reachable by keyboard»
                wrap = els[0]
                try:
                    driver.execute_script(
                        "arguments[0].scrollIntoView({block:'center'});"
                        "const p = arguments[0].closest('div');"
                        "if (p) p.click();",
                        wrap,
                    )
                    sleep(0.2)
                except Exception:
                    pass
                _set_input_value(driver, els[0], text)
                actual = (els[0].get_attribute("value") or "").strip()
                logger.info("{} = {} (force)", pid, actual or "пусто")
                filled_date = bool(actual)
                if filled_date:
                    break
            except Exception as exc:
                logger.debug("purpose date {}: {}", pid, exc)
        if not filled_date:
            logger.warning("Дата направления не заполнена (#referenceDate)")


def fill_date_field_el(driver: WebDriver, el, value: str) -> None:
    """Заполнить date input по элементу — всегда DD-MM-YY."""
    text = normalize_service_date(value)
    if not text:
        return
    digits = re.sub(r"\D", "", text)
    driver.execute_script("arguments[0].scrollIntoView({block:'center'});", el)
    sleep(0.1)
    try:
        el.click()
    except Exception:
        driver.execute_script("arguments[0].focus();", el)
    el.send_keys(Keys.CONTROL, "a")
    el.send_keys(Keys.BACKSPACE)
    for ch in digits:
        el.send_keys(ch)
        sleep(0.03)
    _set_input_value(driver, el, text)
    el.send_keys(Keys.TAB)
    sleep(0.15)


def open_services_tab(driver: WebDriver) -> None:
    # дождаться подгрузки услуг после карты
    try:
        _wait(driver, 25).until(
            lambda d: _tab_is_enabled(d, "services-tab")
            or "услуг" in (snackbar_text(d) or "").lower()
        )
    except TimeoutException:
        pass
    wait_tab_enabled(driver, "services-tab", timeout=45)
    click_id(driver, "services-tab")
    wait_busy_gone(driver, timeout=45)
    sleep(0.8)
    try:
        _wait(driver, 20).until(
            lambda d: bool(
                d.find_elements(By.CSS_SELECTOR, "[id^='date-picker-services-']")
                or d.find_elements(By.CSS_SELECTOR, "[id^='combobox-failure-services-']")
            )
        )
    except TimeoutException:
        logger.warning("Таймаут ожидания строк услуг на services-tab")


def list_service_indices(driver: WebDriver) -> list[int]:
    ids = set()
    for el in driver.find_elements(By.CSS_SELECTOR, "[id^='date-picker-services-']"):
        m = re.search(r"date-picker-services-(\d+)$", el.get_attribute("id") or "")
        if m:
            ids.add(int(m.group(1)))
    for el in driver.find_elements(By.CSS_SELECTOR, "[id^='combobox-failure-services-']"):
        m = re.search(r"combobox-failure-services-(\d+)$", el.get_attribute("id") or "")
        if m:
            ids.add(int(m.group(1)))
    return sorted(ids)


_SERVICE_CODE_RE = re.compile(
    r"[A-Za-z]\d{2}\.\d{2,3}\.\d{3}(?:\.\d+)*",
    re.IGNORECASE,
)


def service_code_label(driver: WebDriver, index: int) -> str:
    """Код услуги из колонки «Наименование» (первый combobox-doctor-N в ОМС)."""
    els = driver.find_elements(By.ID, f"combobox-doctor-{index}")
    if not els:
        return ""
    wrap = _closest_react_select(els[0])
    if not wrap:
        return ""
    chunks: list[str] = []
    for sel in wrap.find_elements(By.CSS_SELECTOR, ".Selected-item, .Select-value, .Select-value-label"):
        text = (sel.text or sel.get_attribute("title") or "").strip()
        if text:
            chunks.append(text)
    blob = " ".join(chunks)
    if not blob:
        return ""
    # UI иногда дублирует значение («B01.001.001B01.001.001») — берём первое вхождение
    m = _SERVICE_CODE_RE.search(blob.replace(" ", ""))
    if m:
        return m.group(0)
    return re.sub(r"\s+", "", blob)


def _closest_react_select(input_el):
    """Ближайший контейнер react-select (.Select), не .Select-input."""
    el = input_el
    for _ in range(10):
        try:
            el = el.find_element(By.XPATH, "..")
        except Exception:
            return None
        cls = (el.get_attribute("class") or "").split()
        if "Select" in cls:
            return el
    return None


def fill_react_select_input(driver: WebDriver, input_el, text: str) -> None:
    """Ввод в react-select / Select (меню часто в portal, опции не всегда сразу видны)."""
    wrap = _closest_react_select(input_el)

    driver.execute_script("arguments[0].scrollIntoView({block:'center'});", input_el)
    sleep(0.1)
    driver.execute_script("arguments[0].focus();", input_el)
    sleep(0.1)

    opened = False
    if wrap is not None:
        for sel in (".Select-arrow-zone", ".Select-control", ".Select-arrow"):
            zones = wrap.find_elements(By.CSS_SELECTOR, sel)
            if not zones:
                continue
            try:
                driver.execute_script(
                    "arguments[0].dispatchEvent(new MouseEvent('mousedown',{bubbles:true}));"
                    "arguments[0].dispatchEvent(new MouseEvent('mouseup',{bubbles:true}));"
                    "arguments[0].click();",
                    zones[0],
                )
                opened = True
                break
            except Exception:
                continue
    if not opened:
        input_el.send_keys(Keys.ARROW_DOWN)
    sleep(0.45)

    try:
        input_el.send_keys(Keys.CONTROL, "a")
        input_el.send_keys(Keys.BACKSPACE)
    except Exception:
        pass
    input_el.send_keys(str(text))
    sleep(0.7)

    options = driver.find_elements(
        By.CSS_SELECTOR,
        ".Select-option, .Select-menu .Select-option, .Select-menu-outer .Select-option, div[role='option']",
    )
    needle = str(text).strip().lower()
    for opt in options:
        label = (opt.text or "").strip().lower()
        if not label:
            continue
        if needle == label or label.startswith(needle) or needle in label.split()[0:1] or needle in label:
            try:
                driver.execute_script("arguments[0].click();", opt)
                sleep(0.25)
                return
            except Exception:
                continue
    if options:
        try:
            driver.execute_script("arguments[0].click();", options[0])
            sleep(0.25)
            return
        except Exception:
            pass
    # часто после ввода кода опция уже подставилась в Select (async)
    if wrap is not None:
        shown = (wrap.text or "").strip().lower()
        if needle and needle in shown:
            input_el.send_keys(Keys.TAB)
            sleep(0.2)
            return
    input_el.send_keys(Keys.ENTER)
    sleep(0.3)
    if wrap is not None and needle and needle not in (wrap.text or "").lower():
        input_el.send_keys(Keys.TAB)
        sleep(0.2)


def fill_service_doctor(driver: WebDriver, index: int, doctor: str) -> None:
    """
    В ОМС у строки услуги id combobox-doctor-N дублируется:
      [0] — код услуги (уже заполнен), [1] — врач.
    Берём последний input с этим id.
    """
    els = driver.find_elements(By.ID, f"combobox-doctor-{index}")
    if not els:
        raise RuntimeError(f"нет combobox-doctor-{index}")
    target = els[-1]
    fill_react_select_input(driver, target, doctor)


def fill_service_done(driver: WebDriver, index: int, value: str) -> None:
    el = driver.find_element(By.ID, f"combobox-failure-services-{index}")
    text = normalize_service_done(value)
    # ОМС может писать «Перезачёт» с ё
    variants = [text]
    if text == "Перезачет":
        variants.append("Перезачёт")
    logger.info("Услуга[{}] Выполнено ← «{}»", index, text)
    last_exc: Exception | None = None
    for v in variants:
        try:
            fill_react_select_input(driver, el, v)
            return
        except Exception as exc:
            last_exc = exc
    if last_exc:
        raise last_exc


def fill_service_date(driver: WebDriver, index: int, date_text: str) -> None:
    """Дата услуги: маска DD-MM-YY. Ввод цифр + native setter (React иначе не видит)."""
    el = driver.find_element(By.ID, f"date-picker-services-{index}")
    text = normalize_service_date(date_text)  # DD-MM-YY
    digits = re.sub(r"\D", "", text)
    want_yy = text[-2:] if text else ""

    driver.execute_script("arguments[0].scrollIntoView({block:'center'});", el)
    try:
        el.click()
    except Exception:
        driver.execute_script("arguments[0].focus();", el)
    sleep(0.08)

    # клавиатура: только цифры под маску
    el.send_keys(Keys.CONTROL, "a")
    el.send_keys(Keys.BACKSPACE)
    el.send_keys(digits)
    sleep(0.05)

    # React controlled input: native value setter + input event
    _set_input_value(driver, el, text)
    try:
        el.send_keys(Keys.TAB)
    except Exception:
        pass
    sleep(0.12)

    actual = (el.get_attribute("value") or "").strip().replace(".", "-")
    if want_yy and not actual.endswith(want_yy):
        logger.warning("Услуга[{}] дата «{}» ≠ …{} — повтор digits+JS", index, actual or "пусто", want_yy)
        el.click()
        el.send_keys(Keys.CONTROL, "a")
        el.send_keys(Keys.BACKSPACE)
        for ch in digits:
            el.send_keys(ch)
            sleep(0.03)
        _set_input_value(driver, el, text)
        el.send_keys(Keys.TAB)
        sleep(0.12)
        actual = (el.get_attribute("value") or "").strip()
    logger.debug("Услуга[{}] дата итого={}", index, actual)


def fill_services(
    driver: WebDriver,
    row: TalonRow,
    catalog: DoctorCatalog,
    recredit_rules: dict | None = None,
) -> int:
    open_services_tab(driver)
    indices = list_service_indices(driver)
    if not indices:
        logger.warning("Строк услуг не найдено")
        return 0

    date_for_services = row.end or row.begin
    done_default = normalize_service_done(row.done_value)
    rules = recredit_rules if recredit_rules is not None else load_recredit_rules()
    filled = 0
    for i in indices:
        last_exc: Exception | None = None
        for attempt in range(1, 4):
            try:
                code = service_code_label(driver, i)
                doctor = resolve_service_doctor(
                    catalog,
                    building=row.building,
                    service_code=code,
                    main_doctor=row.doctor,
                )
                done_val, svc_date, rr = resolve_service_done_and_date(
                    rules,
                    exam_type=row.exam_type,
                    service_code=code,
                    end_date=date_for_services,
                    default_done=done_default,
                )
                if attempt == 1:
                    src = "справочник" if doctor and doctor != (row.doctor or "").strip() else (
                        "карта" if doctor else "пусто"
                    )
                    extra = ""
                    if rr is not None:
                        extra = f" | перезачет={done_val!r} дата={svc_date} (−{rr.months_back}м, {rr.name or code})"
                    logger.info(
                        "Услуга[{}] code={} doctor={} [{}] корпус={}{}",
                        i,
                        code or "—",
                        doctor or "—",
                        src,
                        row.building or "—",
                        extra,
                    )
                # врач и «Выполнено» сначала — дата последней (иначе сбрасывается)
                fill_service_doctor(driver, i, doctor)
                fill_service_done(driver, i, done_val)
                fill_service_date(driver, i, svc_date)
                filled += 1
                last_exc = None
                break
            except (StaleElementReferenceException, NoSuchElementException) as exc:
                last_exc = exc
                logger.warning("Услуга[{}] stale/missing, retry {}/3", i, attempt)
                sleep(0.4)
                wait_busy_gone(driver, timeout=15)
        if last_exc is not None:
            raise last_exc
    return filled


def process_row(
    driver: WebDriver,
    row: TalonRow,
    *,
    index: int,
    catalog: DoctorCatalog,
    dry_run: bool,
    dn_labels: dict[str, str] | None = None,
    recredit_rules: dict | None = None,
) -> RowResult:
    t0 = time()
    logger.info("=== [{}] {} {} ===", index, row.exam_type.upper(), row.enp)
    try:
        open_exam_form(driver, row.exam_type)
        fill_patient_and_search(driver, row)
        ensure_patient_social_status(driver, row)
        # карта пациента / лишний main-tab могли стереть begin/end — вернуть ТОЛЬКО если пусто
        ensure_claim_dates_present(driver, row.begin, row.end, why="после карты пациента")
        fill_main_page(driver, row, dn_labels=dn_labels)
        # ещё раз перед услугами: fill_main не должен трогать даты, но ОМС бывает капризный
        ensure_claim_dates_present(driver, row.begin, row.end, why="перед услугами")

        # После заполнения карты вкладка Услуги обычно уже активна («Услуги получены»)
        n_svc = fill_services(driver, row, catalog, recredit_rules=recredit_rules)
        # Результат III группы / пустые наз_* в Excel → всё равно заполняем назначения
        ensure_default_purposes(row)
        if (row.purpose_type or "").strip() or _tab_is_enabled(driver, "purposes-tab"):
            if not (row.purpose_type or "").strip():
                row.purpose_type = DEFAULT_PURPOSE_TYPE
                row.purpose_from = row.purpose_from or DEFAULT_PURPOSE_FROM
                row.purpose_to = row.purpose_to or DEFAULT_PURPOSE_TO
            fill_purposes(driver, row)

        if dry_run:
            msg = f"dry-run: услуг заполнено {n_svc}, сохранение пропущено"
            if row.purpose_type:
                msg += f", наз_вид={row.purpose_type}"
            logger.success("[{}] {}", index, msg)
            return RowResult(index, row.enp, row.exam_type, "dry_run", msg, n_svc, round(time() - t0, 2))

        ok, msg = save_claim(driver, begin=row.begin, end=row.end)
        if (not ok) and purposes_required_message(msg):
            logger.warning("[{}] ОМС требует назначения — заполняем и повторяем Save", index)
            row.purpose_type = row.purpose_type or DEFAULT_PURPOSE_TYPE
            row.purpose_from = row.purpose_from or DEFAULT_PURPOSE_FROM
            row.purpose_to = row.purpose_to or DEFAULT_PURPOSE_TO
            fill_purposes(driver, row)
            ok, msg = save_claim(driver, begin=row.begin, end=row.end)
        if (not ok) and _is_social_status_error(msg):
            logger.warning("[{}] ошибка соцстатуса/занятости — открываем карту пациента и повторяем", index)
            # сбросить «из Excel» нельзя — ensure всё равно пишет в карту; force пересчёт если пусто
            ensure_patient_social_status(driver, row)
            ensure_claim_dates_present(driver, row.begin, row.end, why="после соцстатуса retry")
            fill_main_page(driver, row, dn_labels=dn_labels)
            ensure_claim_dates_present(driver, row.begin, row.end, why="перед услугами retry")
            n_svc = fill_services(driver, row, catalog, recredit_rules=recredit_rules)
            fill_purposes(driver, row)
            ok, msg = save_claim(driver, begin=row.begin, end=row.end)
        # Недопустимое ДН / зависший диалог пациента → НАЗАД уже нажали → «Состоит» → Save
        if (not ok) and _needs_dn_consists_fix(msg):
            et = (row.exam_type or "").lower()
            want = "7" if et == "opv" else "1"
            # если уже Состоит и снова та же ошибка — для ДВ4 всё равно ещё раз выставим подпись
            logger.warning(
                "[{}] ДН/диалог пациента — ставим {} ({}) на форме и Save снова ({})",
                index,
                want,
                "Состоит, проведено" if want == "7" else "Состоит",
                (msg or "")[:200],
            )
            # закрыть хвосты диалога, если остались
            try:
                if find_visible_dialogs(driver) or warning_dialog_text(driver):
                    dismiss_warning_dialog(driver)
            except Exception:
                pass
            apply_dn_code(driver, row, want, dn_labels, why="retry после НАЗАД/ошибки ДН")
            if _is_social_status_error(msg):
                ensure_patient_social_status(driver, row)
                ensure_claim_dates_present(driver, row.begin, row.end, why="ДН+соц retry")
            # услуги не перезаполняем — только Save с услуг
            if driver.find_elements(By.ID, "services-tab"):
                try:
                    click_id(driver, "services-tab")
                    wait_busy_gone(driver, timeout=15)
                    sleep(0.25)
                except Exception:
                    pass
            ok, msg = save_claim(driver, begin=row.begin, end=row.end)
        # ОПВ: «состоит по диагнозу» / неверный код ДН → ставим 7 и Save ещё раз
        if (
            (not ok)
            and (row.exam_type or "").lower() == "opv"
            and _is_opv_dn_consists_error(msg)
            and (normalize_dn(row.dn) or row.dn or "") != "7"
        ):
            logger.warning(
                "[{}] ОПВ: ошибка ДН/«состоит» — ставим ДН=7 и повторяем Save ({})",
                index,
                (msg or "")[:160],
            )
            apply_opv_dn7(driver, row, dn_labels)
            ok, msg = save_claim(driver, begin=row.begin, end=row.end)
        status = "ok" if ok else "fail"
        log = logger.success if ok else logger.error
        log("[{}] {} → {} ({:.1f}s)", index, status.upper(), msg, time() - t0)
        return RowResult(index, row.enp, row.exam_type, status, msg, n_svc, round(time() - t0, 2))
    except Exception as exc:
        msg = str(exc).split("Stacktrace:")[0].strip() or str(exc)
        logger.error("[{}] error ({:.1f}s): {}", index, time() - t0, msg)
        return RowResult(index, row.enp, row.exam_type, "error", msg, 0, round(time() - t0, 2))


def write_results(path: Path, results: list[RowResult]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(asdict(results[0]).keys()) if results else [
            "index", "enp", "exam_type", "status", "message", "services_filled", "elapsed_sec"
        ])
        w.writeheader()
        for r in results:
            w.writerow(asdict(r))


def ensure_template(path: Path) -> None:
    if path.is_file():
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "талоны"
    ws.append([
        "ЕНП", "начало", "окончание", "тип", "врач", "результат",
        "период", "диагноз", "ДН", "диагноз2", "ДН2", "диагноз3", "ДН3",
        "корпус", "место", "выполнено",
        "наз_вид", "наз_от_кого", "наз_к_кому",
    ])
    # врач можно оставить пустым и дописать вручную
    ws.append([
        "1234567890123456", "01-01-26", "15-01-26", "ДВ4", "",
        "1", "январь", "Z00.0", "3", "", "", "", "",
        "", "1", "Да",
        "", "", "",
    ])
    # ДВ4: ДН=1 (есть в ИСЗЛ) + сопутствующий с ДН=1
    ws.append([
        "1234567890123457", "01-01-26", "15-01-26", "ДВ4", "12345",
        "32", "январь", "K29.7", "1", "I11.9", "1", "", "",
        "", "1", "Да",
        "1", "95", "95",
    ])
    # ОПВ: ДН=7 если диагноз в ИСЗЛ (Состоит, проведено)
    ws.append([
        "1234567890123458", "01-01-26", "15-01-26", "ОПВ", "",
        "2", "январь", "I11.9", "7", "E11.9", "3", "", "",
        "", "1", "Да",
        "", "", "",
    ])
    ws_dn = wb.create_sheet("ДН")
    ws_dn.append(["код_дн", "ДН"])
    for code, label in (
        ("1", "Состоит"),
        ("2", "Взят"),
        ("3", "Не подлежит"),
        ("7", "Состоит, проведено диспансерное наблюдение"),
        ("8", "Взят, проведено диспансерное наблюдение"),
    ):
        ws_dn.append([code, label])
    wb.save(path)
    logger.info("Создан шаблон Excel: {}", path)


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Ввод талонов ДВ4/ОПВ (Web.ОМС medicalExamination)")
    p.add_argument("--file", type=Path, default=DEFAULT_FILE)
    p.add_argument("--limit", type=int, default=0)
    p.add_argument("--offset", type=int, default=0)
    p.add_argument("--type", type=str, default="", help="Принудительно ДВ4 или ОПВ для всех строк")
    p.add_argument(
        "--catalog",
        type=Path,
        default=DEFAULT_REFERENCE,
        help="Справочник ДВ4/ОПВ (листы «ДН» и «справочник_врачей»)",
    )
    p.add_argument("--dry-run", action="store_true", help="Заполнить форму без Save")
    p.add_argument("--keep-open", type=int, default=0)
    p.add_argument("--make-template", action="store_true")
    return p.parse_args()


def main() -> int:
    args = parse_args()
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)

    if args.make_template:
        ensure_template(args.file)
        DoctorCatalog.ensure_example(args.catalog)
        return 0

    if not args.file.is_file():
        ensure_template(args.file)
        logger.error("Нет файла {}. Заполните шаблон и запустите снова.", args.file)
        return 1

    rows = load_talons(args.file)
    if args.type:
        force = _norm_type(args.type)
        for r in rows:
            r.exam_type = force
    if args.offset:
        rows = rows[args.offset :]
    if args.limit and args.limit > 0:
        rows = rows[: args.limit]

    catalog = DoctorCatalog.load_for_talons_file(args.file, fallback=args.catalog)
    dn_labels = load_dn_labels(args.catalog)
    if not dn_labels:
        dn_labels = load_dn_labels(DEFAULT_REFERENCE)
    recredit_rules = load_recredit_rules(args.catalog)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    result_path = RESULTS_DIR / f"dv_opv_result_{stamp}_o{args.offset}.csv"

    catalog_src = Path(catalog.source).name if catalog.source else str(args.catalog)
    logger.info(
        "Строк: {} | dry_run={} | врачи: {} правил ({}) | ДН: {} | перезачет: {} кодов",
        len(rows),
        args.dry_run,
        len(catalog.rules),
        catalog_src,
        len(dn_labels),
        len(recredit_rules),
    )
    results: list[RowResult] = []
    t_all = time()

    with GeckoBrowser(implicit_wait=0) as session:
        driver = session.driver
        driver.implicitly_wait(0)
        last_login_exc: Exception | None = None
        for attempt in range(1, 4):
            try:
                login_oms(driver)
                sleep(1.2)
                wait_busy_gone(driver, timeout=45)
                # признак успешного входа — ушли с /login или есть меню
                url = (driver.current_url or "").lower()
                if "login" in url and attempt < 3:
                    raise RuntimeError(f"всё ещё login url={url}")
                last_login_exc = None
                logger.info("OMS login OK (attempt {})", attempt)
                break
            except Exception as exc:
                last_login_exc = exc
                logger.warning("OMS login fail attempt {}/3: {}", attempt, exc)
                sleep(2.5 * attempt)
        if last_login_exc is not None:
            raise RuntimeError(f"Не удалось авторизоваться в ОМС: {last_login_exc}") from last_login_exc

        for i, row in enumerate(rows, start=1 + args.offset):
            res = process_row(
                driver,
                row,
                index=i,
                catalog=catalog,
                dry_run=args.dry_run,
                dn_labels=dn_labels,
                recredit_rules=recredit_rules,
            )
            results.append(res)

        if args.keep_open > 0:
            sleep(args.keep_open)

    write_results(result_path, results)
    ok_n = sum(1 for r in results if r.status in ("ok", "dry_run"))
    fail_n = len(results) - ok_n
    logger.info("——— ИТОГ ({:.0f}s) OK={}, FAIL={} → {} ———", time() - t_all, ok_n, fail_n, result_path)
    return 0 if fail_n == 0 and results else 1


if __name__ == "__main__":
    sys.exit(main())
