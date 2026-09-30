#!/usr/bin/env python
"""
Проставление соцстатуса / вида занятости в карте пациента Web.ОМС.

Автономный проект: достаточно скопировать эту папку на другой ПК,
запустить install.bat и run.bat.

URL: /patient/edit
Логика возраста:
  18 < age < 60 → socialStatus=11, occupation=1
  иначе (в т.ч. ≥60) → 22 / 3

Пример:
  python add_socstatus.py --limit 3
  python add_socstatus.py --file data/list.xlsx --force-age
  python add_socstatus.py --dry-run --keep-open 30
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from time import sleep, time

from loguru import logger
from openpyxl import Workbook, load_workbook
from selenium.common.exceptions import TimeoutException
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.remote.webdriver import WebDriver
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait

ROOT = Path(__file__).resolve().parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from lib.browser import GeckoBrowser
from lib.config import DATA_DIR, load_credentials, oms_base_url, oms_login, oms_password
from lib import oms_ui as ui

load_credentials()

DEFAULT_FILE = DATA_DIR / "list.xlsx"
RESULTS_DIR = DATA_DIR
PATIENT_EDIT_PATH = "/patient/edit"

ENP_RADIO_ID = "2"
ENP_INPUT_XPATH = (
    "/html/body/div[1]/div/div[2]/div[1]/div[2]/div[1]/div/div[2]/div/div/div/div/input"
)
SEARCH_BTN_XPATH = (
    "/html/body/div[1]/div/div[2]/div[1]/div[2]/div[1]/div/div[3]/div/div[1]/button"
)


@dataclass
class SocRow:
    enp: str
    social_status: str = ""
    occupation: str = ""
    excel_row: int = 0
    sheet_index: int = 0


@dataclass
class RowResult:
    index: int
    enp: str
    social_status: str
    occupation: str
    status: str
    message: str
    elapsed_sec: float = 0.0


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]
    if text.lower() in {"nan", "none", "-"}:
        return ""
    return text


def _norm_header(name: object) -> str:
    return str(name or "").strip().casefold().replace("ё", "е")


def _col_map(headers: list[object]) -> dict[str, int]:
    out: dict[str, int] = {}
    for i, h in enumerate(headers):
        key = _norm_header(h)
        if key and key not in out:
            out[key] = i
    return out


def _find_col(cmap: dict[str, int], *aliases: str) -> int | None:
    for alias in aliases:
        idx = cmap.get(_norm_header(alias))
        if idx is not None:
            return idx
    return None


def patient_edit_url() -> str:
    return f"{oms_base_url()}{PATIENT_EDIT_PATH}"


def _wait(driver: WebDriver, timeout: float = 15) -> WebDriverWait:
    return WebDriverWait(driver, timeout)


def open_patient_edit(driver: WebDriver) -> None:
    driver.get(patient_edit_url())
    ui.wait_busy_gone(driver)
    sleep(0.8)
    logger.info("Открыта страница {}", driver.current_url)


def select_enp_search(driver: WebDriver) -> None:
    radios = driver.find_elements(By.ID, ENP_RADIO_ID)
    if radios:
        ui.js_click(driver, radios[0])
        sleep(0.2)
        return
    for xp in (
        "//input[@id='2']",
        "//label[contains(.,'ЕНП')]/preceding::input[@type='radio'][1]",
        "//input[@type='radio' and (@value='2' or @value='enp')]",
    ):
        els = driver.find_elements(By.XPATH, xp)
        if els:
            ui.js_click(driver, els[0])
            sleep(0.2)
            return
    logger.warning("Радиокнопка поиска по ЕНП не найдена — пробуем ввод как есть")


def _enp_input(driver: WebDriver):
    for by, sel in (
        (By.XPATH, ENP_INPUT_XPATH),
        (By.CSS_SELECTOR, "input[type='text']"),
    ):
        els = [e for e in driver.find_elements(by, sel) if e.is_displayed()]
        if els:
            return els[0]
    raise RuntimeError("Поле ввода ЕНП на /patient/edit не найдено")


def _search_button(driver: WebDriver):
    for by, sel in (
        (By.XPATH, SEARCH_BTN_XPATH),
        (By.XPATH, "//button[contains(.,'Найти') or contains(.,'Поиск')]"),
    ):
        els = [e for e in driver.find_elements(by, sel) if e.is_displayed()]
        if els:
            return els[0]
    raise RuntimeError("Кнопка поиска пациента не найдена")


def search_patient(driver: WebDriver, enp: str) -> None:
    ui.dismiss_overlays(driver)
    field = _enp_input(driver)
    driver.execute_script("arguments[0].scrollIntoView({block:'center'});", field)
    try:
        field.click()
    except Exception:
        ui.js_click(driver, field)
    field.send_keys(Keys.CONTROL, "a")
    field.send_keys(Keys.BACKSPACE)
    field.send_keys(enp)
    sleep(0.2)
    # Сначала Enter в поле — надёжнее, чем кнопка под overlay
    try:
        field.send_keys(Keys.ENTER)
    except Exception:
        pass
    ui.wait_busy_gone(driver, timeout=15)
    sleep(0.3)
    # Если карта не открылась — жмём поиск через JS
    if not driver.find_elements(By.ID, "socialStatus"):
        ui.dismiss_overlays(driver)
        ui.js_click(driver, _search_button(driver))
        ui.wait_busy_gone(driver, timeout=40)
        sleep(0.5)
    try:
        _wait(driver, 20).until(EC.presence_of_element_located((By.ID, "socialStatus")))
    except TimeoutException as exc:
        raise RuntimeError(f"Пациент ЕНП={enp} не открылся (нет #socialStatus)") from exc


def _read_birth(driver: WebDriver) -> str:
    for element_id in ("birth-date", "birthDate", "patient-birth-date"):
        els = driver.find_elements(By.ID, element_id)
        if els:
            val = (els[0].get_attribute("value") or "").strip()
            if val:
                return val
    for el in driver.find_elements(By.CSS_SELECTOR, "input[disabled], input[readonly]"):
        val = (el.get_attribute("value") or "").strip()
        if re.search(r"\d{1,2}[-./]\d{1,2}[-./]\d{2,4}", val):
            return val
    return ""


def resolve_soc_occ(driver: WebDriver, row: SocRow, *, force_age: bool) -> tuple[str, str]:
    soc = (row.social_status or "").strip()
    occ = (row.occupation or "").strip()
    if force_age:
        soc = occ = ""

    if soc and occ:
        logger.info("Из Excel: socialStatus={}, occupation={}", soc, occ)
        return soc, occ

    birth_raw = _read_birth(driver)
    try:
        soc, occ = ui.social_occupation_from_birth(birth_raw)
    except ValueError:
        if soc and not occ:
            return soc, "1" if soc == "11" else "3"
        if occ and not soc:
            return ("11" if occ == "1" else "22"), occ
        raise RuntimeError(f"Не разобрать дату рождения: {birth_raw!r}")
    year = ui.birth_year(birth_raw)
    age = datetime.now().year - (year or 0)
    logger.info(
        "ДР={} → год={}, возраст≈{} → socialStatus={}, occupation={}",
        birth_raw,
        year,
        age,
        soc,
        occ,
    )
    return soc, occ


def _fill_select(driver: WebDriver, element_id: str, value: str) -> None:
    el = driver.find_element(By.ID, element_id)
    ui.fill_react_select_input(driver, el, value)
    sleep(0.25)
    ui.wait_busy_gone(driver)
    wrap = ui.closest_react_select(el)
    shown = (wrap.text if wrap is not None else "") or ""
    if str(value) not in shown:
        el.send_keys(Keys.CONTROL, "a")
        el.send_keys(Keys.BACKSPACE)
        el.send_keys(str(value), Keys.ENTER)
        sleep(0.5)
        wrap = ui.closest_react_select(el)
        shown = (wrap.text if wrap is not None else "") or ""
    if str(value) not in shown:
        raise RuntimeError(f"{element_id}: не выбрано «{value}», сейчас «{shown[:80]}»")
    logger.info("{} → {}", element_id, shown.split("\n")[0][:80])


def save_patient(driver: WebDriver) -> None:
    saves = [e for e in driver.find_elements(By.ID, "save-button") if e.is_displayed()]
    if not saves:
        raise RuntimeError("#save-button не найден")
    ui.js_click(driver, saves[0])
    ui.wait_busy_gone(driver, timeout=60)
    sleep(0.4)
    dlg = ui.warning_dialog_text(driver)
    if dlg or ui.find_visible_dialogs(driver):
        ok, info = ui.handle_post_save_prompts(driver, timeout=15)
        if not ok:
            raise RuntimeError(f"Сохранение отклонено: {info}")
        if info:
            logger.info("Диалог после save: {}", info[:200])
    else:
        msg = ui.snackbar_text(driver)
        if msg:
            logger.info("После save: {}", msg[:200])
            low = msg.lower()
            if any(x in low for x in ("ошибка", "обязательн", "не указан")):
                raise RuntimeError(f"Карта не сохранена: {msg}")


def process_row(
    driver: WebDriver,
    row: SocRow,
    *,
    force_age: bool,
    dry_run: bool,
) -> RowResult:
    t0 = time()
    try:
        search_patient(driver, row.enp)
        soc, occ = resolve_soc_occ(driver, row, force_age=force_age)
        if dry_run:
            logger.info("[dry-run] ЕНП={} → {} / {} (без save)", row.enp, soc, occ)
            return RowResult(
                index=row.sheet_index + 1,
                enp=row.enp,
                social_status=soc,
                occupation=occ,
                status="dry-run",
                message=f"{soc}/{occ}",
                elapsed_sec=time() - t0,
            )
        _fill_select(driver, "socialStatus", soc)
        _fill_select(driver, "occupation", occ)
        save_patient(driver)
        return RowResult(
            index=row.sheet_index + 1,
            enp=row.enp,
            social_status=soc,
            occupation=occ,
            status="Сохранен",
            message=f"{soc}/{occ}",
            elapsed_sec=time() - t0,
        )
    except Exception as exc:
        logger.exception("ЕНП={}: {}", row.enp, exc)
        return RowResult(
            index=row.sheet_index + 1,
            enp=row.enp,
            social_status=row.social_status,
            occupation=row.occupation,
            status="Ошибка",
            message=str(exc)[:400],
            elapsed_sec=time() - t0,
        )


def load_rows(path: Path) -> tuple[list[SocRow], list[list[object]], list[object]]:
    wb = load_workbook(path, read_only=False, data_only=True)
    ws = wb.active
    rows_raw = list(ws.iter_rows(values_only=True))
    if not rows_raw:
        raise ValueError(f"Пустой файл: {path}")
    headers = list(rows_raw[0])
    cmap = _col_map(headers)
    i_enp = _find_col(cmap, "енп", "enp")
    if i_enp is None:
        raise KeyError(f"Нет колонки ЕНП. Колонки: {headers}")
    i_soc = _find_col(cmap, "соцстатус", "социальный статус", "social_status", "socialstatus")
    i_occ = _find_col(cmap, "вид занятости", "занятость", "occupation")
    data: list[SocRow] = []
    for idx, raw in enumerate(rows_raw[1:]):
        if raw is None:
            continue
        cells = list(raw)
        enp = _cell(cells[i_enp] if i_enp < len(cells) else "")
        if not enp or len(re.sub(r"\D", "", enp)) < 11:
            continue
        soc = _cell(cells[i_soc]) if i_soc is not None and i_soc < len(cells) else ""
        occ = _cell(cells[i_occ]) if i_occ is not None and i_occ < len(cells) else ""
        if soc.lower() in {"сохранен", "ошибка", "dry-run", "пропуск"}:
            soc = ""
        data.append(
            SocRow(
                enp=enp,
                social_status=soc,
                occupation=occ,
                excel_row=idx + 2,
                sheet_index=idx,
            )
        )
    return data, [list(r) if r else [] for r in rows_raw], headers


def ensure_progress_columns(headers: list[object], rows: list[list[object]]) -> tuple[int, int]:
    cmap = _col_map(headers)
    i_status = _find_col(cmap, "status")
    i_info = _find_col(cmap, "info")
    if i_status is None:
        headers.append("Status")
        i_status = len(headers) - 1
    if i_info is None:
        headers.append("info")
        i_info = len(headers) - 1
    for row in rows:
        while len(row) < len(headers):
            row.append("")
    rows[0] = list(headers)
    return i_status, i_info


def write_progress(
    path: Path,
    headers: list[object],
    rows: list[list[object]],
    results_by_sheet_index: dict[int, RowResult],
    i_status: int,
    i_info: int,
) -> None:
    for idx, result in results_by_sheet_index.items():
        excel_idx = idx + 1
        if excel_idx >= len(rows):
            continue
        row = rows[excel_idx]
        while len(row) < len(headers):
            row.append("")
        row[i_status] = result.status
        row[i_info] = result.message

    wb = Workbook()
    ws = wb.active
    ws.title = "socstatus"
    for row in rows:
        ws.append([("" if c is None else c) for c in row])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    logger.info("Прогресс записан → {}", path)


def write_results_csv(results: list[RowResult], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(
            f,
            fieldnames=[
                "index",
                "enp",
                "social_status",
                "occupation",
                "status",
                "message",
                "elapsed_sec",
            ],
        )
        w.writeheader()
        for r in results:
            w.writerow(
                {
                    "index": r.index,
                    "enp": r.enp,
                    "social_status": r.social_status,
                    "occupation": r.occupation,
                    "status": r.status,
                    "message": r.message,
                    "elapsed_sec": round(r.elapsed_sec, 2),
                }
            )
    logger.info("Результат CSV → {}", path)


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Проставление соцстатуса на /patient/edit")
    p.add_argument("--file", type=Path, default=DEFAULT_FILE, help="Excel с ЕНП")
    p.add_argument("--offset", type=int, default=0)
    p.add_argument("--limit", type=int, default=0, help="0 = все")
    p.add_argument("--force-age", action="store_true", help="Всегда считать по ДР")
    p.add_argument("--skip-saved", action="store_true", default=True)
    p.add_argument("--no-skip-saved", action="store_false", dest="skip_saved")
    p.add_argument("--dry-run", action="store_true")
    p.add_argument("--keep-open", type=int, default=0)
    p.add_argument("--save-every", type=int, default=5)
    return p.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    path = args.file.resolve()
    if not path.is_file():
        logger.error("Файл не найден: {}", path)
        logger.error("Положите Excel в data/list.xlsx или укажите --file")
        return 2

    rows, sheet_rows, headers = load_rows(path)
    i_status, i_info = ensure_progress_columns(headers, sheet_rows)

    if args.skip_saved:
        filtered: list[SocRow] = []
        for r in rows:
            st = ""
            if r.sheet_index + 1 < len(sheet_rows):
                raw = sheet_rows[r.sheet_index + 1]
                if i_status < len(raw):
                    st = _cell(raw[i_status])
            if st == "Сохранен":
                logger.info("Пропуск ЕНП={} (уже Сохранен)", r.enp)
                continue
            filtered.append(r)
        rows = filtered

    if args.offset:
        rows = rows[args.offset :]
    if args.limit and args.limit > 0:
        rows = rows[: args.limit]

    if not rows:
        logger.warning("Нет строк для обработки")
        return 0

    logger.info(
        "К обработке: {} из {} (force_age={}, dry_run={})",
        len(rows),
        path.name,
        args.force_age,
        args.dry_run,
    )

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    result_csv = RESULTS_DIR / f"oms_socstatus_result_{stamp}.csv"
    results: list[RowResult] = []
    by_idx: dict[int, RowResult] = {}

    with GeckoBrowser(maximize=True, window_size=(1200, 900)) as browser:
        driver = browser.driver
        ui.login_oms(
            driver,
            base_url=oms_base_url(),
            login=oms_login(),
            password=oms_password(),
        )
        sleep(1)
        open_patient_edit(driver)
        select_enp_search(driver)

        for n, row in enumerate(rows, start=1):
            logger.info("——— {}/{} ЕНП={} ———", n, len(rows), row.enp)
            res = process_row(driver, row, force_age=args.force_age, dry_run=args.dry_run)
            results.append(res)
            by_idx[row.sheet_index] = res
            logger.info(
                "[{}] {}  {}  {}/{}  ({:.1f}s)",
                res.index,
                res.status,
                res.enp,
                res.social_status,
                res.occupation,
                res.elapsed_sec,
            )
            if n % max(1, args.save_every) == 0:
                write_progress(path, headers, sheet_rows, by_idx, i_status, i_info)

        write_progress(path, headers, sheet_rows, by_idx, i_status, i_info)
        write_results_csv(results, result_csv)

        if args.keep_open > 0:
            logger.info("Окно открыто ещё {} с…", args.keep_open)
            sleep(args.keep_open)

    ok = sum(1 for r in results if r.status in {"Сохранен", "dry-run"})
    err = sum(1 for r in results if r.status == "Ошибка")
    logger.info("Готово: ок={}, ошибка={}, CSV={}", ok, err, result_csv)
    return 0 if err == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
